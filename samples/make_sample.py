"""Generate a sample integrated SRS (rev A and rev B) for a BLDC position-control ECU.

python samples/make_sample.py  ->  samples/SRS_BLDC_revA.docx, samples/SRS_BLDC_revB.docx
"""
from io import BytesIO
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

OUT = Path(__file__).parent


def state_machine_png() -> bytes:
    plt.rcParams["font.family"] = ["DejaVu Sans"]
    fig, ax = plt.subplots(figsize=(7, 2.6), dpi=130)
    ax.axis("off")
    states = {"INIT": (0.08, 0.5), "NORMAL": (0.38, 0.5), "FAILSAFE": (0.68, 0.8), "SLEEP": (0.68, 0.2)}
    for n, (x, y) in states.items():
        ax.add_patch(FancyBboxPatch((x - 0.08, y - 0.1), 0.16, 0.2, boxstyle="round,pad=0.02", fc="#e7effb", ec="#1f5fbf"))
        ax.text(x, y, n, ha="center", va="center", fontsize=10, weight="bold")
    def arrow(a, b, label):
        (x0, y0), (x1, y1) = states[a], states[b]
        ax.annotate("", xy=(x1 - 0.09, y1), xytext=(x0 + 0.09, y0), arrowprops=dict(arrowstyle="->", color="#333"))
        ax.text((x0 + x1) / 2, (y0 + y1) / 2 + 0.06, label, ha="center", fontsize=7.5)
    arrow("INIT", "NORMAL", "Init OK")
    arrow("NORMAL", "FAILSAFE", "Fault > debounce")
    arrow("NORMAL", "SLEEP", "IG OFF & idle 3s")
    ax.set_xlim(0, 0.8); ax.set_ylim(0, 1)
    buf = BytesIO(); fig.savefig(buf, format="png", bbox_inches="tight"); plt.close(fig)
    return buf.getvalue()


def block_diagram_png() -> bytes:
    fig, ax = plt.subplots(figsize=(7, 2.2), dpi=130)
    ax.axis("off")
    boxes = [("CAN\nTransceiver", 0.05), ("TC364\nMCU", 0.3), ("Gate Driver\nIC", 0.55), ("BLDC Motor\n+ Hall", 0.8)]
    for t, x in boxes:
        ax.add_patch(FancyBboxPatch((x, 0.3), 0.15, 0.4, boxstyle="round,pad=0.02", fc="#eef6ee", ec="#22803a"))
        ax.text(x + 0.075, 0.5, t, ha="center", va="center", fontsize=9)
    for i in range(3):
        ax.annotate("", xy=(boxes[i + 1][1] - 0.02, 0.5), xytext=(boxes[i][1] + 0.17, 0.5), arrowprops=dict(arrowstyle="<->"))
    ax.set_xlim(0, 1); ax.set_ylim(0, 1)
    buf = BytesIO(); fig.savefig(buf, format="png", bbox_inches="tight"); plt.close(fig)
    return buf.getvalue()


_CACHE = {}


def cached(fn):
    if fn.__name__ not in _CACHE:
        _CACHE[fn.__name__] = fn()
    return _CACHE[fn.__name__]


def merge_vertical(table, col, r0, r1):
    a = table.cell(r0, col)
    b = table.cell(r1, col)
    a.merge(b)


def add_table(doc, rows, widths=None, style="Table Grid"):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = style
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            if v is not None:
                t.cell(i, j).text = v
                if i == 0:
                    for run in t.cell(i, j).paragraphs[0].runs:
                        run.bold = True
    return t


def build(rev: str) -> Path:
    B = rev == "B"
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Malgun Gothic"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    st.font.size = Pt(10)

    sec = doc.sections[0]
    sec.header.paragraphs[0].text = "SRS-BLDC-001  |  사내 대외비 (Confidential)"
    sec.footer.paragraphs[0].text = f"ACME Mobility  -  BLDC ECU  -  Rev {rev}"
    doc.add_paragraph("BLDC 위치제어 ECU 통합 요구사항 명세서 (SRS)", style="Title")
    doc.add_paragraph(f"문서번호: SRS-BLDC-001  개정: {rev}")

    doc.add_heading("개요", 1)
    doc.add_paragraph("본 문서는 TC364 기반 BLDC 모터 위치제어 ECU의 시스템 및 소프트웨어 요구사항을 정의한다.")
    doc.add_paragraph("적용 규격: 현대자동차 ES 전원 규격, ISO 14229(UDS), AUTOSAR Classic 4.4.")

    doc.add_heading("시스템 요구사항", 1)
    doc.add_heading("전원", 2)
    doc.add_paragraph("제어기는 9V ~ 16V 작동 전압 범위에서 모든 기능을 정상 수행해야 한다.")
    doc.add_paragraph("IG OFF 상태에서 제어기의 암전류는 100uA 이하여야 한다.")
    if B:
        doc.add_paragraph("저전압(9V 미만) 상태가 300ms 이상 지속되면 모터 구동을 중지해야 한다. 이때 DTC B1601을 저장해야 한다.")
    else:
        doc.add_paragraph("저전압(9V 미만) 상태가 500ms 이상 지속되면 모터 구동을 중지해야 한다. 이때 DTC B1601을 저장해야 한다.")
    doc.add_heading("환경 및 기구", 2)
    doc.add_paragraph("제어기는 작동 온도 -40°C ~ 85°C 범위에서 동작해야 한다.")
    doc.add_paragraph("액추에이터 스트로크는 0 ~ 45mm 범위를 가져야 한다.")
    if not B:
        doc.add_paragraph("하우징은 IP6K9K 방수 등급을 만족해야 한다.")
    doc.add_paragraph("그림 1 시스템 블록 구성")  # caption placed ABOVE on purpose (caption detection both ways)
    doc.add_picture(BytesIO(cached(block_diagram_png)), width=Cm(15))

    doc.add_heading("소프트웨어 요구사항", 1)
    doc.add_heading("위치 제어", 2)
    if B:
        doc.add_paragraph("SW는 목표 위치 수신 시 1ms 주기 PID 제어로 위치 편차 0.5deg 이내로 수렴하도록 제어해야 한다.")
    else:
        doc.add_paragraph("SW는 목표 위치 수신 시 1ms 주기 PID 제어로 위치 편차 1deg 이내로 수렴하도록 제어해야 한다.")
    doc.add_paragraph("SW는 모터를 빠르게 안정적으로 구동해야 하며 소음을 최소화해야 한다.")
    doc.add_paragraph("영점 학습(Homing) 루틴은 UDS 0x31 RoutineControl 요청 시 수행되어야 한다. 학습 결과는 NvM에 저장해야 한다.")
    doc.add_paragraph("모터 구속(Stall) 전류가 15A 이상으로 500ms 이상 지속되면 SW는 Fail-safe 상태로 천이해야 한다.")
    if B:
        doc.add_paragraph("과온도(125°C 이상) 검출 시 SW는 PWM 출력을 차단해야 한다.")

    doc.add_heading("CAN 통신", 2)
    doc.add_paragraph("표 1 CAN 통신 요구사항")
    timeout = "500ms" if B else "300ms"
    add_table(doc, [
        ["ID", "요구사항", "검증방법", "ASIL"],
        ["REQ-COM-001", "SW는 목표 위치 신호 PosCmd_Tgt를 10ms 주기 메시지 BLDC_CMD로 수신해야 한다.", "SWE.6 CANoe", "QM"],
        ["REQ-COM-002", f"BLDC_CMD 메시지가 {timeout} 이상 수신되지 않으면 SW는 Timeout으로 판단하고 현재 위치를 유지해야 한다.", "SWE.6 고장주입", "ASIL-B"],
        ["REQ-COM-003", "SW는 Alive Counter 오류가 3회 연속 발생하면 해당 메시지를 무효로 처리해야 한다.", "SWE.6 고장주입", "ASIL-B"],
        ["REQ-COM-004", "SW는 BLDC_CMD 메시지의 CRC 오류 시 해당 프레임을 폐기해야 한다.", "SWE.6 고장주입", "ASIL-B"],
    ])

    doc.add_heading("진단", 2)
    doc.add_paragraph("SW는 UDS 0x10 세션 제어, 0x22 ReadDataByIdentifier, 0x2E WriteDataByIdentifier, 0x19 ReadDTCInformation 서비스를 지원해야 한다.")
    doc.add_paragraph("DID 0xF190 VIN 쓰기 조건은 TBD이며 확장 세션에서만 허용되어야 한다.")

    doc.add_heading("HW-SW 인터페이스", 1)
    doc.add_paragraph("표 2 TC364 핀맵")
    t = add_table(doc, [
        ["Signal", "Pin", "Port", "Peripheral", "비고"],
        ["PWM_U_HS", "45", "P02.0", "GTM TOM0", "3상 PWM 20kHz"],
        ["PWM_V_HS", "46", "P02.1", None, "3상 PWM 20kHz"],
        ["PWM_W_HS", "47", "P02.2", None, "3상 PWM 20kHz"],
        ["I_SHUNT", "88", "AN0", "EVADC G0", "전류 션트"],
        ["CAN_TX", "12", "P20.8", "MCMCAN0", "500kbps"],
    ])
    merge_vertical(t, 3, 1, 3)
    t.cell(1, 3).text = "GTM TOM0"
    doc.add_picture(BytesIO(cached(state_machine_png)), width=Cm(15))
    doc.add_paragraph("그림 2 전원 모드 상태 천이")

    p = OUT / f"SRS_BLDC_rev{rev}.docx"
    doc.save(p)
    return p


if __name__ == "__main__":
    for r in ("A", "B"):
        print(build(r))
