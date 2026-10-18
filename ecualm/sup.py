"""Supporting processes (ASPICE PAM 4.0 SUP.1 / SUP.8 / SUP.9 / SUP.10) on one record engine.

Record kinds
  PR   problem record (SUP.9)          new → analyzing → resolving → verifying → closed   (rejected · duplicate)
  CR   change request (SUP.10)         submitted → analyzing → review → approved → implementing → verifying → closed
                                       (rejected · deferred)   — impact analysis + CCB decision are recorded
  QA   quality assurance check (SUP.1) planned → done          — checklist, partly evaluated from ALM data
  QF   QA finding (SUP.1)              open → in_progress → resolved → verified    (escalation recorded)
  REL  release / configuration record (SUP.8)  draft → released — baseline + build/commit + consistency checks

Every record has fields (per kind schema), linked ALM items with a role, references to other records, an owner,
a due date and a full history (who changed what, every status transition with its comment). Transitions are
guarded: e.g. a problem cannot be closed before its verifying test cases PASS, a change request cannot go to the
CCB without a saved impact analysis, a finding cannot be verified by the person who resolved it.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import shutil
import sqlite3
import tempfile
import time
import zipfile
from collections import Counter, defaultdict, deque
from pathlib import Path

SCHEMA = """
CREATE TABLE IF NOT EXISTS sup_records (
    id       TEXT PRIMARY KEY,
    kind     TEXT NOT NULL,
    title    TEXT NOT NULL,
    status   TEXT NOT NULL,
    owner    TEXT NOT NULL DEFAULT '',
    due      TEXT NOT NULL DEFAULT '',
    fields   TEXT NOT NULL DEFAULT '{}',
    items    TEXT NOT NULL DEFAULT '[]',
    refs     TEXT NOT NULL DEFAULT '[]',
    created  REAL NOT NULL,
    created_by TEXT NOT NULL DEFAULT '',
    updated  REAL NOT NULL,
    updated_by TEXT NOT NULL DEFAULT ''
);
CREATE INDEX IF NOT EXISTS ix_sup_kind ON sup_records(kind, status);
CREATE TABLE IF NOT EXISTS sup_history (
    rec TEXT NOT NULL, at REAL NOT NULL, by TEXT NOT NULL, action TEXT NOT NULL,
    from_status TEXT NOT NULL DEFAULT '', to_status TEXT NOT NULL DEFAULT '', comment TEXT NOT NULL DEFAULT '',
    changes TEXT NOT NULL DEFAULT '{}'
);
CREATE INDEX IF NOT EXISTS ix_sup_hist ON sup_history(rec, at);
"""


class SupError(ValueError):
    pass


def F(key, label, type_="text", req=False, options=None, help_="", group=""):
    return {"key": key, "label": label, "type": type_, "required": req, "options": options or [], "help": help_, "group": group}


SEVERITY = ["Critical", "Major", "Minor"]
KINDS = {
    "PR": {
        "label": "문제", "process": "SUP.9", "title": "문제 해결 관리",
        "statuses": {"new": "접수", "analyzing": "분석", "resolving": "조치", "verifying": "검증", "closed": "종결",
                     "rejected": "반려", "duplicate": "중복"},
        "terminal": ["closed", "rejected", "duplicate"],
        "flow": {"new": ["analyzing", "rejected", "duplicate"], "analyzing": ["resolving", "rejected", "duplicate"],
                 "resolving": ["verifying", "analyzing"], "verifying": ["closed", "resolving"],
                 "closed": ["analyzing"], "rejected": ["new"], "duplicate": ["new"]},
        "item_roles": {"found_by": "발견 TC", "affects": "영향 항목 (요구사항·설계)", "verify": "확인 TC (재시험)"},
        "fields": [
            F("description", "문제 기술 (D2)", "long", True, help_="무엇이·언제·어디서·어떤 조건에서 — 재현 절차 포함"),
            F("source", "출처", "enum", True, ["시험 실패", "리뷰 지적", "고객 이슈", "필드 불량", "내부 발견", "공급사"]),
            F("severity", "심각도", "enum", True, SEVERITY),
            F("safety", "안전 관련", "enum", True, ["아니오", "예 (ASIL)", "확인 필요"]),
            F("reproducibility", "재현성", "enum", False, ["항상", "간헐", "1회"]),
            F("found_in", "발견 빌드·샘플", "text"), F("session", "테스트 세션", "text"),
            F("team", "담당 팀 (D1)", "text", group="8D"),
            F("containment", "임시 조치 (D3)", "long", group="분석"),
            F("root_cause_cat", "원인 분류", "enum", False, ["요구사항", "설계", "코드", "TC·시험 절차", "시험 환경", "HW", "도구", "공정", "외부·공급사"], group="분석"),
            F("root_cause", "근본 원인 (D4)", "long", group="분석", help_="5 Why·특성요인도 결과"),
            F("horizontal", "수평 전개 결과", "long", group="분석", help_="같은 원인이 다른 변형·모듈·프로젝트에 있는지 확인한 결과"),
            F("corrective", "영구 조치 (D5·D6)", "long", group="조치"),
            F("prevention", "재발 방지 (D7)", "long", group="조치"),
            F("verification", "조치 검증 (D8)", "long", group="검증", help_="재시험 세션·결과"),
            F("duplicate_of", "중복 대상", "text", group="종결"),
        ],
        "guards": {"resolving": ["root_cause_cat", "root_cause"], "verifying": ["corrective"], "closed": ["verification"],
                   "duplicate": ["duplicate_of"]},
        "create_role": "engineer",
    },
    "CR": {
        "label": "변경 요청", "process": "SUP.10", "title": "변경 요청 관리",
        "statuses": {"submitted": "접수", "analyzing": "영향 분석", "review": "CCB 검토", "approved": "승인",
                     "implementing": "반영", "verifying": "검증", "closed": "종결", "rejected": "반려", "deferred": "보류"},
        "terminal": ["closed", "rejected"],
        "flow": {"submitted": ["analyzing", "rejected"], "analyzing": ["review"], "review": ["approved", "rejected", "deferred", "analyzing"],
                 "approved": ["implementing"], "implementing": ["verifying"], "verifying": ["closed", "implementing"],
                 "deferred": ["analyzing"], "rejected": ["submitted"], "closed": []},
        "manager_to": ["approved", "rejected", "deferred"],
        "item_roles": {"scope": "변경 대상 항목"},
        "fields": [
            F("description", "변경 내용", "long", True), F("reason", "변경 사유", "enum", True,
              ["고객 요구 변경", "문제 조치 (SUP.9)", "내부 개선", "안전", "보안", "법규·인증"]),
            F("priority", "우선순위", "enum", True, ["긴급", "높음", "보통", "낮음"]),
            F("customer_ref", "고객 요청 번호", "text"), F("target_release", "적용 릴리스", "text", help_="예: B 샘플 SW 0.4.0"),
            F("effort_estimate", "추정 공수 (MD)", "num", group="영향 분석"),
            F("schedule_impact", "일정 영향", "long", group="영향 분석"),
            F("safety_impact", "안전 영향 (ISO 26262)", "long", group="영향 분석"),
            F("security_impact", "보안 영향 (ISO/SAE 21434)", "long", group="영향 분석"),
            F("interface_impact", "인터페이스·고객 영향", "long", group="영향 분석", help_="CAN·진단·다른 ECU·고객 합의"),
            F("decision", "CCB 결정 내용", "long", group="CCB"),
            F("verification", "반영 검증 결과", "long", group="검증"),
        ],
        "guards": {"review": ["effort_estimate", "safety_impact"], "approved": ["decision"], "rejected": ["decision"],
                   "closed": ["verification"]},
        "create_role": "engineer",
    },
    "QA": {
        "label": "QA 점검", "process": "SUP.1", "title": "품질 보증",
        "statuses": {"planned": "계획", "done": "완료"}, "terminal": ["done"],
        "flow": {"planned": ["done"], "done": ["planned"]},
        "item_roles": {},
        "fields": [
            F("check_type", "점검 유형", "enum", True, ["산출물 점검", "프로세스 점검"]),
            F("target", "점검 대상", "text", True, help_="단계 ID(SYS2, SWE1 …) 또는 프로세스(SUP.8, MAN.3 …)"),
            F("planned", "계획일", "date"), F("criteria", "점검 기준", "long", help_="적용 규격·체크리스트·샘플링 기준"),
            F("summary", "점검 결과 요약", "long", group="결과"),
        ],
        "guards": {"done": ["summary"]},
        "create_role": "manager",
    },
    "QF": {
        "label": "QA 지적", "process": "SUP.1", "title": "품질 보증",
        "statuses": {"open": "열림", "in_progress": "조치 중", "resolved": "조치 완료", "verified": "확인 완료"},
        "terminal": ["verified"],
        "flow": {"open": ["in_progress", "resolved"], "in_progress": ["resolved"], "resolved": ["verified", "in_progress"], "verified": []},
        "item_roles": {"concerns": "관련 항목"},
        "fields": [
            F("description", "지적 내용", "long", True), F("severity", "등급", "enum", True, SEVERITY),
            F("process", "해당 프로세스", "text", True), F("qa", "QA 점검", "text"),
            F("corrective", "시정 조치", "long", group="조치"), F("escalated_to", "에스컬레이션 대상", "text", group="에스컬레이션"),
        ],
        "guards": {"resolved": ["corrective"]},
        "create_role": "manager",
    },
    "REL": {
        "label": "릴리스", "process": "SUP.8", "title": "형상 관리",
        "statuses": {"draft": "준비", "released": "릴리스"}, "terminal": ["released"],
        "flow": {"draft": ["released"], "released": []}, "manager_to": ["released"],
        "item_roles": {},
        "fields": [
            F("sw_version", "SW 버전", "text", True), F("hw_sample", "HW 샘플", "text"),
            F("git_url", "소스 저장소", "text"), F("git_commit", "커밋 / 태그", "text", True),
            F("build_id", "빌드 ID·해시", "text"), F("calibration", "캘리브레이션 데이터 버전", "text"),
            F("arxml", "ARXML·A2L 버전", "text"), F("sessions", "포함 테스트 세션", "text", help_="쉼표 구분 (예: TS-003, TS-004)"),
            F("notes", "릴리스 노트", "long"), F("override", "검사 예외 사유", "long", group="릴리스"),
        ],
        "guards": {}, "create_role": "manager",
    },
}


def init(repo):
    repo.con.executescript(SCHEMA)
    repo.con.commit()


def _rank(user) -> int:
    from .auth import rank
    return rank(user["role"])


def _need(user, role):
    from .auth import ROLE_LABEL, rank
    if _rank(user) < rank(role):
        raise PermissionError(f"{ROLE_LABEL[role]} 이상만 할 수 있습니다")


def _row(r) -> dict:
    d = dict(r)
    for k in ("fields", "items", "refs"):
        d[k] = json.loads(d[k])
    return d


def get(repo, rid: str) -> dict:
    r = repo.con.execute("SELECT * FROM sup_records WHERE id=?", (rid,)).fetchone()
    if not r:
        raise KeyError(f"{rid} 없음")
    return _row(r)


def _next_id(repo, kind: str) -> str:
    n = 0
    for (i,) in repo.con.execute("SELECT id FROM sup_records WHERE kind=?", (kind,)):
        m = re.match(rf"{kind}-(\d+)$", i)
        if m:
            n = max(n, int(m.group(1)))
    return f"{kind}-{n + 1:03d}"


def _hist(repo, rid, by, action, frm="", to="", comment="", changes=None):
    repo.con.execute("INSERT INTO sup_history VALUES(?,?,?,?,?,?,?,?)",
                     (rid, time.time(), by, action, frm, to, comment or "", json.dumps(changes or {}, ensure_ascii=False)))


def _clean_fields(kind: str, data: dict) -> dict:
    out = {}
    for f in KINDS[kind]["fields"]:
        if f["key"] not in data:
            continue
        v = data[f["key"]]
        if f["type"] == "num":
            v = None if v in ("", None) else float(v)
        elif f["type"] == "enum" and v not in ("", None) and v not in f["options"]:
            raise SupError(f"{f['label']}: {', '.join(f['options'])} 중에서 고르세요")
        elif f["type"] == "date" and v and not re.match(r"^\d{4}-\d{2}-\d{2}$", str(v)):
            raise SupError(f"{f['label']}: 날짜 형식은 YYYY-MM-DD")
        else:
            v = "" if v is None else str(v).strip()
        out[f["key"]] = v
    return out


def _items(repo, kind, items) -> list[dict]:
    roles = KINDS[kind]["item_roles"]
    out, seen = [], set()
    for x in items or []:
        iid, role = (x.get("id"), x.get("role")) if isinstance(x, dict) else (x, next(iter(roles), ""))
        if role not in roles:
            raise SupError(f"링크 역할은 {', '.join(roles)} 중 하나")
        if not repo.get(iid):
            raise SupError(f"항목 {iid} 없음")
        if (iid, role) not in seen:
            seen.add((iid, role)); out.append({"id": iid, "role": role})
    return out


def create(repo, cfg, kind: str, b: dict, user: dict) -> dict:
    if kind not in KINDS:
        raise SupError("알 수 없는 종류")
    K = KINDS[kind]
    _need(user, K["create_role"])
    title = (b.get("title") or "").strip()
    if not title:
        raise SupError("제목을 입력하세요")
    fields = _clean_fields(kind, b.get("fields") or {})
    missing = [f["label"] for f in K["fields"] if f["required"] and not fields.get(f["key"]) and f["type"] != "num"]
    if missing:
        raise SupError(f"필수 입력: {', '.join(missing)}")
    if b.get("due") and not re.match(r"^\d{4}-\d{2}-\d{2}$", b["due"]):
        raise SupError("기한 형식은 YYYY-MM-DD")
    rid = _next_id(repo, kind)
    now = time.time()
    start = next(iter(K["statuses"]))
    refs = [r for r in (b.get("refs") or []) if repo.con.execute("SELECT 1 FROM sup_records WHERE id=?", (r,)).fetchone()]
    repo.con.execute("INSERT INTO sup_records VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                     (rid, kind, title, start, (b.get("owner") or "").strip(), b.get("due", ""),
                      json.dumps(fields, ensure_ascii=False), json.dumps(_items(repo, kind, b.get("items")), ensure_ascii=False),
                      json.dumps(refs), now, user["username"], now, user["username"]))
    for r in refs:     # references are kept on both sides
        _add_ref(repo, r, rid)
    _hist(repo, rid, user["username"], "created", "", start, b.get("comment", ""))
    if kind == "QA" and b.get("auto_checklist", True):
        tgt = fields.get("target", "")
        repo.con.execute("UPDATE sup_records SET fields=? WHERE id=?",
                         (json.dumps({**fields, "checklist": checklist_template(cfg, tgt)}, ensure_ascii=False), rid))
    repo.con.commit()
    return detail(repo, cfg, rid)


def _add_ref(repo, rid, other):
    r = repo.con.execute("SELECT refs FROM sup_records WHERE id=?", (rid,)).fetchone()
    if r:
        refs = json.loads(r["refs"])
        if other not in refs:
            refs.append(other)
            repo.con.execute("UPDATE sup_records SET refs=? WHERE id=?", (json.dumps(refs), rid))


def update(repo, cfg, rid: str, b: dict, user: dict) -> dict:
    """Edit fields / owner / due / items / refs (no status change)."""
    cur = get(repo, rid)
    K = KINDS[cur["kind"]]
    if cur["status"] in K["terminal"] and _rank(user) < 2:
        raise PermissionError("종결된 기록은 PM·관리자만 고칠 수 있습니다")
    _need(user, "engineer")
    changes, sets = {}, {}
    if "fields" in b:
        nf = _clean_fields(cur["kind"], b["fields"])
        if cur["kind"] == "CR" and "decision" in nf and nf["decision"] != cur["fields"].get("decision", "") and _rank(user) < 2:
            raise PermissionError("CCB 결정은 PM·관리자만 입력합니다")
        merged = {**cur["fields"], **nf}
        changes.update({k: [cur["fields"].get(k), v] for k, v in nf.items() if cur["fields"].get(k) != v})
        sets["fields"] = json.dumps(merged, ensure_ascii=False)
    for k in ("title", "owner", "due"):
        if k in b and b[k] != cur[k]:
            if k == "due" and b[k] and not re.match(r"^\d{4}-\d{2}-\d{2}$", b[k]):
                raise SupError("기한 형식은 YYYY-MM-DD")
            if k == "title" and not str(b[k]).strip():
                raise SupError("제목을 입력하세요")
            changes[k] = [cur[k], b[k]]; sets[k] = str(b[k]).strip()
    if "items" in b:
        its = _items(repo, cur["kind"], b["items"])
        if its != cur["items"]:
            changes["items"] = [[f"{x['role']}:{x['id']}" for x in cur["items"]], [f"{x['role']}:{x['id']}" for x in its]]
            sets["items"] = json.dumps(its)
    if "refs" in b:
        refs = [r for r in b["refs"] if r != rid and repo.con.execute("SELECT 1 FROM sup_records WHERE id=?", (r,)).fetchone()]
        if refs != cur["refs"]:
            changes["refs"] = [cur["refs"], refs]; sets["refs"] = json.dumps(refs)
            for r in refs:
                _add_ref(repo, r, rid)
    if not sets:
        return detail(repo, cfg, rid)
    sets.update(updated=time.time(), updated_by=user["username"])
    repo.con.execute(f"UPDATE sup_records SET {', '.join(f'{k}=?' for k in sets)} WHERE id=?", (*sets.values(), rid))
    _hist(repo, rid, user["username"], "edited", changes=changes)
    repo.con.commit()
    return detail(repo, cfg, rid)


# ---------------------------------------------------------------- transitions & guards
def _latest(repo, tc) -> dict:
    it = repo.get(tc) or {}
    return (it.get("attrs") or {}).get("result") or {}


def guard_problems(repo, cfg, rec: dict, to: str, user: dict, comment: str) -> list[str]:
    """Return blocking reasons for a transition (empty = allowed)."""
    K = KINDS[rec["kind"]]
    errs = []
    f = rec["fields"]
    labels = {x["key"]: x["label"] for x in K["fields"]}
    for k in K["guards"].get(to, []):
        if f.get(k) in (None, ""):
            errs.append(f"'{labels[k]}' 입력 필요")
    if to in K.get("manager_to", []) and _rank(user) < 2:
        errs.append("이 결정은 PM·관리자만 할 수 있습니다")
    if rec["kind"] == "PR":
        if to == "closed":
            since = _status_time(repo, rec["id"], "resolving") or rec["created"]
            for x in rec["items"]:
                if x["role"] == "verify":
                    r = _latest(repo, x["id"])
                    if r.get("status") != "PASS" or (r.get("at") or 0) < since:
                        errs.append(f"확인 TC {x['id']} 가 조치 이후 PASS 가 아님 (현재 {r.get('status', '미실행')})")
            if not any(x["role"] == "verify" for x in rec["items"]) and not comment:
                errs.append("확인 TC 가 없으면 종결 사유(코멘트)를 적어야 합니다")
        if to == "duplicate" and f.get("duplicate_of") and not repo.con.execute(
                "SELECT 1 FROM sup_records WHERE id=?", (f["duplicate_of"],)).fetchone():
            errs.append(f"중복 대상 {f['duplicate_of']} 없음")
        if to == "closed" and _rank(user) < 2 and f.get("safety", "").startswith("예"):
            errs.append("안전 관련 문제는 PM·관리자가 종결합니다")
    if rec["kind"] == "CR":
        if to == "review" and not f.get("impact"):
            errs.append("영향 분석을 실행·저장해야 CCB 에 올릴 수 있습니다")
        if to == "closed":
            imp = f.get("impact") or {}
            since = _status_time(repo, rec["id"], "approved") or rec["created"]
            for tc in imp.get("retest", []):
                r = _latest(repo, tc)
                if r.get("status") != "PASS" or (r.get("at") or 0) < since:
                    errs.append(f"재시험 TC {tc}: 승인 이후 PASS 아님 ({r.get('status', '미실행')})")
            unchanged = [x["id"] for x in rec["items"] if x["role"] == "scope" and
                         (repo.get(x["id"]) or {}).get("version", 0) <= imp.get("versions", {}).get(x["id"], 0)]
            if unchanged and not comment:
                errs.append(f"변경 대상인데 바뀌지 않은 항목 {', '.join(unchanged[:5])} — 사유를 코멘트로 남기세요")
    if rec["kind"] == "QF" and to == "verified":
        resolver = _status_by(repo, rec["id"], "resolved")
        if resolver == user["username"]:
            errs.append("조치한 사람이 직접 확인할 수 없습니다 (독립성)")
        if _rank(user) < 2 and rec["created_by"] != user["username"]:
            errs.append("확인은 지적한 QA 담당 또는 PM·관리자가 합니다")
    if rec["kind"] == "QA" and to == "done":
        cl = rec["fields"].get("checklist") or []
        if any(not c.get("result") for c in cl):
            errs.append("체크리스트에 판정이 비어 있는 항목이 있습니다")
        _need(user, "manager")
    if rec["kind"] == "REL" and to == "released":
        chk = release_checks(repo, cfg, rec)
        bad = [c for c in chk if c["sev"] == "error" and not c["ok"]]
        if bad and not (f.get("override") or comment):
            errs += [f"{c['label']}: {c['value']}" for c in bad] + ["검사 예외 사유를 적으면 릴리스할 수 있습니다"]
    return errs


def _status_time(repo, rid, status):
    r = repo.con.execute("SELECT MAX(at) a FROM sup_history WHERE rec=? AND to_status=?", (rid, status)).fetchone()
    return r["a"] if r and r["a"] else None


def _status_by(repo, rid, status):
    r = repo.con.execute("SELECT by FROM sup_history WHERE rec=? AND to_status=? ORDER BY at DESC LIMIT 1", (rid, status)).fetchone()
    return r["by"] if r else None


def transition(repo, cfg, rid: str, to: str, comment: str, user: dict) -> dict:
    rec = get(repo, rid)
    K = KINDS[rec["kind"]]
    _need(user, "engineer")
    if to not in K["flow"].get(rec["status"], []):
        raise SupError(f"{K['statuses'][rec['status']]} → {K['statuses'].get(to, to)} 로 바꿀 수 없습니다")
    if rec["status"] in K["terminal"] and _rank(user) < 2:
        raise PermissionError("종결된 기록을 다시 여는 것은 PM·관리자만 할 수 있습니다")
    errs = guard_problems(repo, cfg, rec, to, user, comment)
    if errs:
        raise SupError(" · ".join(errs))
    extra = {}
    if rec["kind"] == "REL" and to == "released":
        name = f"{rid} {rec['fields'].get('sw_version', '')}".strip()
        if not repo.con.execute("SELECT 1 FROM baselines WHERE name=?", (name,)).fetchone():
            repo.create_baseline(name, f"릴리스 {rec['title']}")
        f = dict(rec["fields"], item_baseline=name, checks=release_checks(repo, cfg, rec), released_at=time.time(),
                 released_by=user["username"])
        repo.con.execute("UPDATE sup_records SET fields=? WHERE id=?", (json.dumps(f, ensure_ascii=False), rid))
        extra = {"baseline": name}
    if rec["kind"] == "QF" and to == "verified":
        pass
    repo.con.execute("UPDATE sup_records SET status=?, updated=?, updated_by=? WHERE id=?", (to, time.time(), user["username"], rid))
    _hist(repo, rid, user["username"], "status", rec["status"], to, comment, extra)
    repo.con.commit()
    return detail(repo, cfg, rid)


def escalate(repo, cfg, rid: str, to_whom: str, comment: str, user: dict) -> dict:
    rec = get(repo, rid)
    if not to_whom.strip():
        raise SupError("에스컬레이션 대상을 적으세요")
    f = dict(rec["fields"], escalated_to=to_whom.strip())
    repo.con.execute("UPDATE sup_records SET fields=?, updated=?, updated_by=? WHERE id=?",
                     (json.dumps(f, ensure_ascii=False), time.time(), user["username"], rid))
    _hist(repo, rid, user["username"], "escalated", comment=f"→ {to_whom}: {comment}")
    repo.con.commit()
    return detail(repo, cfg, rid)


def comment(repo, cfg, rid: str, text: str, user: dict) -> dict:
    get(repo, rid)
    if not text.strip():
        raise SupError("내용을 입력하세요")
    _hist(repo, rid, user["username"], "comment", comment=text.strip())
    repo.con.commit()
    return detail(repo, cfg, rid)


# ---------------------------------------------------------------- SUP.9 from a failed test
def problem_from_test(repo, cfg, session: str, tc: str, user: dict, b: dict | None = None) -> dict:
    from . import testruns
    s = testruns.session(repo, cfg, session)
    case = next((c for c in s["cases"] if c["tc"] == tc), None)
    if not case:
        raise KeyError(f"{session} 에 {tc} 없음")
    if case["defect"] and re.match(r"^PR-\d+$", case["defect"]):
        raise SupError(f"이미 문제 {case['defect']} 로 등록되어 있습니다")
    b = b or {}
    items = [{"id": tc, "role": "found_by"}, {"id": tc, "role": "verify"}] + [{"id": v, "role": "affects"} for v in case["verifies"]]
    it = repo.get(tc)
    rec = create(repo, cfg, "PR", {
        "title": b.get("title") or f"{tc} 시험 실패 — {(it or {}).get('title', '')}"[:120],
        "owner": b.get("owner", ""), "due": b.get("due", ""),
        "fields": {"description": b.get("description") or f"[기대] {case['expected'] or case['pass_criteria']}\n[실제] {case['actual']}\n[조건] {case['pre_condition']}",
                   "source": "시험 실패", "severity": b.get("severity", "Major"),
                   "safety": "확인 필요" if any((repo.get(v) or {}).get("attrs", {}).get("asil", "QM") not in ("", "QM", None) for v in case["verifies"]) else "아니오",
                   "found_in": " · ".join(x for x in (s["build"], s["sample"]) if x), "session": session,
                   "reproducibility": b.get("reproducibility", "")},
        "items": items}, user)
    testruns.record(repo, cfg, session, {"tc": tc, "defect": rec["id"]}, user) if s["status"] not in ("completed", "aborted") else \
        repo.con.execute("UPDATE session_cases SET defect=? WHERE session=? AND tc=?", (rec["id"], session, tc))
    repo.con.commit()
    return detail(repo, cfg, rec["id"])


def cr_from_problem(repo, cfg, pid: str, user: dict) -> dict:
    p = get(repo, pid)
    if p["kind"] != "PR":
        raise SupError("문제 기록에서만 만들 수 있습니다")
    scope = [{"id": x["id"], "role": "scope"} for x in p["items"] if x["role"] == "affects"]
    return create(repo, cfg, "CR", {"title": f"[{pid}] {p['title']}"[:120], "refs": [pid], "items": scope,
                                    "fields": {"description": p["fields"].get("corrective") or p["fields"].get("description", ""),
                                               "reason": "문제 조치 (SUP.9)", "priority": "높음" if p["fields"].get("severity") == "Critical" else "보통"}}, user)


# ---------------------------------------------------------------- SUP.10 impact analysis
def impact(repo, cfg, ids: list[str], depth: int = 8) -> dict:
    items = {i["id"]: i for i in repo.items()}
    down, up = defaultdict(list), defaultdict(list)
    for l in repo.links():
        if l["src"] in items and l["dst"] in items:
            down[l["dst"]].append(l)
            up[l["src"]].append(l)
    stage_of = {t: s for s in cfg["stages"] for t in s["types"]}
    seen = {}
    q = deque((i, 0, None, "scope") for i in ids if i in items)
    while q:
        iid, d, via, role = q.popleft()
        if iid in seen or d > depth:
            continue
        seen[iid] = (d, via, role)
        for l in down.get(iid, []):
            if l["role"] != "relates":
                q.append((l["src"], d + 1, iid, l["role"]))
    rows = []
    for iid, (d, via, role) in seen.items():
        it = items[iid]
        st = stage_of.get(it["type"], {})
        tc = it["type"].endswith("_TC")
        action = "modify" if d == 0 else "retest" if tc else "review"
        rows.append({"id": iid, "type": it["type"], "stage": st.get("id", ""), "aspice": st.get("aspice", ""), "title": it["title"][:90],
                     "status": it["status"], "version": it["version"], "asil": it["attrs"].get("asil") or "",
                     "depth": d, "via": via, "role": role, "action": action,
                     "result": (it["attrs"].get("result") or {}).get("status", "") if tc else ""})
    rows.sort(key=lambda r: (r["depth"], r["stage"], r["id"]))
    upstream = sorted({l["dst"] for i in ids for l in up.get(i, []) if l["role"] != "relates"})
    by_stage = Counter(r["stage"] for r in rows)
    effort = {"modify": 0.5, "review": 0.25, "retest": 0.25}
    suggested = round(sum(effort[r["action"]] for r in rows), 1)
    from . import schedule
    try:
        tasks = [{"id": t["id"], "wbs": t["wbs"], "name": t["name"], "stage": t["stage"], "finish": t["finish"], "status": t["status"]}
                 for t in schedule.tasks(repo) if t["stage"] and t["stage"] in by_stage and t["status"] != "done"]
    except sqlite3.OperationalError:
        tasks = []
    return {"scope": [i for i in ids if i in items], "rows": rows, "upstream": upstream, "by_stage": dict(by_stage),
            "asil": sorted({r["asil"] for r in rows if r["asil"] and r["asil"] != "QM"}),
            "safety_items": [r["id"] for r in rows if r["asil"] and r["asil"] != "QM"],
            "retest": [r["id"] for r in rows if r["action"] == "retest"],
            "review": [r["id"] for r in rows if r["action"] == "review"],
            "tasks": tasks, "suggested_effort": suggested, "versions": {r["id"]: r["version"] for r in rows}}


def save_impact(repo, cfg, rid: str, user: dict) -> dict:
    rec = get(repo, rid)
    if rec["kind"] != "CR":
        raise SupError("변경 요청에서만 영향 분석을 저장합니다")
    if rec["status"] not in ("submitted", "analyzing", "review", "deferred"):
        raise SupError("승인 이후에는 영향 분석을 다시 저장할 수 없습니다")
    scope = [x["id"] for x in rec["items"] if x["role"] == "scope"]
    if not scope:
        raise SupError("변경 대상 항목을 먼저 연결하세요")
    imp = impact(repo, cfg, scope)
    imp.update(at=time.time(), by=user["username"])
    f = dict(rec["fields"], impact=imp)
    if f.get("effort_estimate") in (None, ""):
        f["effort_estimate"] = imp["suggested_effort"]
    if not f.get("safety_impact") and imp["safety_items"]:
        f["safety_impact"] = f"ASIL {', '.join(imp['asil'])} 항목 {len(imp['safety_items'])}건 영향 — 안전 분석(FMEA/FMEDA·안전 요구) 재검토 필요"
    elif not f.get("safety_impact"):
        f["safety_impact"] = "영향 범위에 ASIL 항목 없음"
    repo.con.execute("UPDATE sup_records SET fields=?, updated=?, updated_by=? WHERE id=?",
                     (json.dumps(f, ensure_ascii=False), time.time(), user["username"], rid))
    _hist(repo, rid, user["username"], "impact", comment=f"영향 {len(imp['rows'])}건 (검토 {len(imp['review'])}, 재시험 {len(imp['retest'])}, ASIL {len(imp['safety_items'])})")
    if rec["status"] == "submitted":
        repo.con.execute("UPDATE sup_records SET status='analyzing' WHERE id=?", (rid,))
        _hist(repo, rid, user["username"], "status", "submitted", "analyzing", "영향 분석 시작")
    repo.con.commit()
    return detail(repo, cfg, rid)


# ---------------------------------------------------------------- SUP.1 checklists (partly evaluated from ALM data)
def checklist_template(cfg, target: str) -> list[dict]:
    st = {s["id"]: s for s in cfg["stages"]}.get(target)
    if st and st.get("side") == "verify":
        q = [("tc_linked", "모든 TC 가 검증 대상과 추적 연결됨"), ("tc_complete", "TC 필수 항목(사전 조건·기대 결과·합격 기준) 작성됨"),
             ("executed", "모든 TC 가 실행됨 (미실행 없음)"), ("fail_tracked", "FAIL TC 마다 문제 기록(SUP.9)이 있음"),
             ("evidence", "실행 결과에 증빙이 첨부됨 (최근 완료 세션)"), ("stale", "대상 변경 후 재시험 필요 TC 없음"),
             (None, "시험 환경·도구 버전이 기록되어 있음"), (None, "시험 결과가 이해관계자에게 전달됨 (13-52)")]
    elif st and st["id"] in ("SYS3", "HWE2", "SWE2", "SWE3"):
        q = [("wp_complete", "필수 속성 누락 없음"), ("upstream", "모든 요소가 상위 요구사항에 할당·추적됨"),
             ("tests", "통합·검증 TC 와 연결됨"), ("suspect", "미해결 Suspect 없음"), ("approved", "리뷰·승인 완료"),
             (None, "인터페이스·동적 동작이 정의되고 리뷰됨"), (None, "설계 결정의 근거가 기록됨")]
    elif st:
        q = [("lint", "요구사항 문장 품질 (모호어·검증 불가 표현 없음)"), ("wp_complete", "필수 속성(ASIL·검증 기준 등) 누락 없음"),
             ("upstream", "상위 요구사항과 양방향 추적됨"), ("tests", "검증 TC 로 커버됨"), ("suspect", "미해결 Suspect 없음"),
             ("approved", "리뷰·승인 완료"), (None, "고객·이해관계자와 합의 기록 존재 (13-52)")]
    else:
        proc = {"SUP.8": [("baseline", "베이스라인이 존재하고 최근 릴리스에 연결됨"), ("release", "릴리스 기록에 빌드·커밋·검사 결과 있음"),
                          (None, "형상 항목 선정 기준과 저장소가 정의됨"), (None, "백업·복구가 수행·확인됨")],
                "SUP.9": [("pr_overdue", "기한 지난 미해결 문제 없음"), ("pr_rootcause", "종결 문제에 근본 원인·조치·검증 기록"),
                          ("pr_safety", "미해결 안전 관련 문제 없음"), (None, "문제 추세 분석이 정기 보고됨")],
                "SUP.10": [("cr_impact", "CCB 검토 이상 CR 에 영향 분석 있음"), ("cr_decision", "승인·반려 CR 에 결정 근거 있음"),
                           ("cr_verified", "종결 CR 에 반영 검증 기록"), (None, "변경 요청과 변경 산출물 간 추적 유지")],
                "MAN.3": [("plan_baseline", "일정 기준선이 존재함"), ("plan_late", "지연 작업이 조치·보고됨"), (None, "위험이 식별·관리됨")]}
        q = proc.get(target, [(None, "프로세스가 정의된 대로 수행됨"), (None, "산출물이 기준에 맞게 작성됨")])
    return [{"q": text, "auto": key or "", "result": "", "value": "", "comment": ""} for key, text in q]


def evaluate_checklist(repo, cfg, rid: str, user: dict) -> dict:
    from . import testruns, workproducts
    from .audit import run_audit
    rec = get(repo, rid)
    if rec["kind"] != "QA":
        raise SupError("QA 점검에서만 실행합니다")
    _need(user, "manager")
    tgt = rec["fields"].get("target", "")
    stages = {s["id"]: s for s in cfg["stages"]}
    st = stages.get(tgt)
    aud = run_audit(repo, cfg)
    items = {i["id"]: i for i in repo.items()}
    vals = {}
    if st:
        t = workproducts.table(repo, cfg, tgt)
        mb = t["missing_by_column"]
        n = len(t["rows"])
        ids = {r["id"] for r in t["rows"]}
        vals["wp_complete"] = (t["missing"] == 0, f"필수 충족 {t['completeness']}% · 누락 {t['missing']}칸")
        upk = [c["key"] for c in t["columns"] if c["src"].startswith("up:") and c["required"]]
        upm = sum(mb.get(k, 0) for k in upk)
        vals["upstream"] = (upm == 0, f"상위 링크 없음 {upm}/{n}")
        vals["tc_linked"] = vals["upstream"]
        tk = [c["key"] for c in t["columns"] if c["key"] == "tests"]
        tm = sum(mb.get(k, 0) for k in tk)
        vals["tests"] = (tm == 0, f"검증 TC 없음 {tm}/{n}")
        lint = sum(1 for i in ids if any(x.get("sev") == "error" for x in items[i].get("lint", [])))
        vals["lint"] = (lint == 0, f"문장 오류 {lint}건")
        sus = sum(1 for l in repo.links() if l["suspect"] and l["src"] in ids)
        vals["suspect"] = (sus == 0, f"Suspect {sus}건")
        appr = sum(1 for i in ids if items[i]["status"] == "approved")
        vals["approved"] = (n > 0 and appr == n, f"승인 {appr}/{n}")
        tcc = sum(mb.get(k, 0) for k in ("pre_condition", "expected", "pass_criteria"))
        vals["tc_complete"] = (tcc == 0, f"누락 {tcc}칸")
        nr = sum(1 for i in ids if not (items[i]["attrs"].get("result") or {}).get("status"))
        vals["executed"] = (nr == 0, f"미실행 {nr}/{n}")
        fails = [i for i in ids if (items[i]["attrs"].get("result") or {}).get("status") == "FAIL"]
        pr_tcs = {x["id"] for r in list_records(repo, "PR") for x in r["items"]}
        untracked = [i for i in fails if i not in pr_tcs]
        vals["fail_tracked"] = (not untracked, f"FAIL {len(fails)} · 문제 미등록 {len(untracked)}")
        stale = sum(1 for f in aud["findings"] if f["rule"] == "result_stale" and f["item"] in ids)
        vals["stale"] = (stale == 0, f"재시험 필요 {stale}")
        done = [s for s in testruns.sessions(repo, cfg) if s["stage"] == tgt and s["status"] == "completed"]
        if done:
            ses = testruns.session(repo, cfg, done[0]["id"])
            ex = [c for c in ses["cases"] if c["status"] in ("PASS", "FAIL", "BLOCKED")]
            noev = [c for c in ex if not c["evidence"]]
            vals["evidence"] = (not noev, f"{done[0]['id']}: 증빙 없는 실행 {len(noev)}/{len(ex)}")
        else:
            vals["evidence"] = (False, "완료된 세션 없음")
    else:
        prs, crs = list_records(repo, "PR"), list_records(repo, "CR")
        today = dt.date.today().isoformat()
        od = [r["id"] for r in prs if r["due"] and r["due"] < today and r["status"] not in KINDS["PR"]["terminal"]]
        vals["pr_overdue"] = (not od, f"기한 초과 {len(od)}건")
        cl = [r for r in prs if r["status"] == "closed"]
        bad = [r["id"] for r in cl if not (r["fields"].get("root_cause") and r["fields"].get("corrective") and r["fields"].get("verification"))]
        vals["pr_rootcause"] = (not bad, f"종결 {len(cl)} · 기록 불충분 {len(bad)}")
        so = [r["id"] for r in prs if r["fields"].get("safety", "").startswith("예") and r["status"] not in KINDS["PR"]["terminal"]]
        vals["pr_safety"] = (not so, f"미해결 안전 문제 {len(so)}건")
        rv = [r for r in crs if r["status"] in ("review", "approved", "implementing", "verifying", "closed")]
        vals["cr_impact"] = (all(r["fields"].get("impact") for r in rv), f"대상 {len(rv)} · 영향 분석 없음 {sum(1 for r in rv if not r['fields'].get('impact'))}")
        dec = [r for r in crs if r["status"] in ("approved", "implementing", "verifying", "closed", "rejected")]
        vals["cr_decision"] = (all(r["fields"].get("decision") for r in dec), f"대상 {len(dec)} · 근거 없음 {sum(1 for r in dec if not r['fields'].get('decision'))}")
        cc = [r for r in crs if r["status"] == "closed"]
        vals["cr_verified"] = (all(r["fields"].get("verification") for r in cc), f"종결 {len(cc)}")
        bls = repo.baselines()
        rels = [r for r in list_records(repo, "REL") if r["status"] == "released"]
        vals["baseline"] = (bool(bls) and bool(rels), f"베이스라인 {len(bls)} · 릴리스 {len(rels)}")
        vals["release"] = (bool(rels) and all(r["fields"].get("git_commit") and r["fields"].get("checks") for r in rels), f"릴리스 {len(rels)}")
        from . import schedule
        try:
            pl = schedule.computed(repo, cfg)
            vals["plan_baseline"] = (bool(pl["baselines"]), f"기준선 {len(pl['baselines'])}")
            vals["plan_late"] = (pl["metrics"]["late"] == 0, f"지연 작업 {pl['metrics']['late']}")
        except sqlite3.OperationalError:
            vals["plan_baseline"] = (False, "계획 없음"); vals["plan_late"] = (False, "계획 없음")
    cl = rec["fields"].get("checklist") or checklist_template(cfg, tgt)
    for c in cl:
        if c.get("auto") and c["auto"] in vals:
            ok, v = vals[c["auto"]]
            c["result"], c["value"] = ("OK" if ok else "NOK"), v
    f = dict(rec["fields"], checklist=cl, evaluated_at=time.time())
    repo.con.execute("UPDATE sup_records SET fields=?, updated=?, updated_by=? WHERE id=?",
                     (json.dumps(f, ensure_ascii=False), time.time(), user["username"], rid))
    _hist(repo, rid, user["username"], "evaluated", comment=f"자동 판정 {sum(1 for c in cl if c.get('auto'))}항목, NOK {sum(1 for c in cl if c['result'] == 'NOK')}")
    repo.con.commit()
    return detail(repo, cfg, rid)


def set_check(repo, cfg, rid: str, idx: int, result: str, comment_: str, user: dict) -> dict:
    rec = get(repo, rid)
    _need(user, "manager")
    cl = rec["fields"].get("checklist") or []
    if not 0 <= idx < len(cl):
        raise SupError("체크리스트 항목 없음")
    if result not in ("OK", "NOK", "N/A", ""):
        raise SupError("판정은 OK · NOK · N/A")
    cl[idx].update(result=result, comment=comment_ or "")
    repo.con.execute("UPDATE sup_records SET fields=? WHERE id=?", (json.dumps(dict(rec["fields"], checklist=cl), ensure_ascii=False), rid))
    repo.con.commit()
    return detail(repo, cfg, rid)


def findings_from_checklist(repo, cfg, rid: str, user: dict) -> dict:
    rec = get(repo, rid)
    _need(user, "manager")
    made = []
    have = {(r["fields"].get("qa"), r["fields"].get("description")) for r in list_records(repo, "QF")}
    for c in rec["fields"].get("checklist") or []:
        if c["result"] == "NOK":
            desc = f"{c['q']} — {c.get('value') or ''} {c.get('comment') or ''}".strip()
            if (rid, desc) in have:
                continue
            f = create(repo, cfg, "QF", {"title": c["q"][:100], "refs": [rid],
                                         "fields": {"description": desc, "severity": "Major", "process": rec["fields"].get("target", ""), "qa": rid}}, user)
            made.append(f["id"])
    return {"created": made, "record": detail(repo, cfg, rid)}


# ---------------------------------------------------------------- SUP.8 release checks, baseline diff, backup
def release_checks(repo, cfg, rec: dict) -> list[dict]:
    from . import testruns, workproducts
    from .audit import run_audit
    aud = run_audit(repo, cfg)
    sev = Counter(f["sev"] for f in aud["findings"])
    sus = sum(1 for l in repo.links() if l["suspect"])
    items = repo.items()
    req = [i for i in items if cfg["types"].get(i["type"], {}).get("requirement")]
    appr = sum(1 for i in req if i["status"] == "approved")
    sess = [s.strip() for s in (rec["fields"].get("sessions") or "").split(",") if s.strip()]
    sfail, smiss = [], []
    for s in sess:
        try:
            x = testruns.session(repo, cfg, s)
            if x["counts"]["FAIL"]:
                sfail.append(f"{s} FAIL {x['counts']['FAIL']}")
            if x["status"] != "completed":
                smiss.append(f"{s} 미완료")
        except KeyError:
            smiss.append(f"{s} 없음")
    prs = [r for r in list_records(repo, "PR") if r["status"] not in KINDS["PR"]["terminal"]]
    crs = [r for r in list_records(repo, "CR") if r["status"] not in KINDS["CR"]["terminal"] + ["deferred"]
           and rec["fields"].get("sw_version") and rec["fields"]["sw_version"] in (r["fields"].get("target_release") or "")]
    comp = []
    for s in cfg["stages"]:
        try:
            t = workproducts.table(repo, cfg, s["id"])
            if t["rows"]:
                comp.append(t["completeness"])
        except KeyError:
            pass
    out = [
        {"key": "audit_errors", "label": "감사 오류", "ok": sev["error"] == 0, "value": f"{sev['error']}건", "sev": "error"},
        {"key": "suspect", "label": "미해결 Suspect", "ok": sus == 0, "value": f"{sus}건", "sev": "warn"},
        {"key": "approved", "label": "요구사항 승인", "ok": bool(req) and appr == len(req), "value": f"{appr}/{len(req)}", "sev": "warn"},
        {"key": "sessions", "label": "포함 세션 FAIL", "ok": not sfail and bool(sess), "value": ", ".join(sfail) or ("없음" if sess else "세션 미지정"), "sev": "error"},
        {"key": "sessions_done", "label": "포함 세션 완료", "ok": not smiss and bool(sess), "value": ", ".join(smiss) or ("모두 완료" if sess else "세션 미지정"), "sev": "warn"},
        {"key": "safety_pr", "label": "미해결 안전 문제", "ok": not [r for r in prs if r["fields"].get("safety", "").startswith("예")],
         "value": f"{sum(1 for r in prs if r['fields'].get('safety', '').startswith('예'))}건", "sev": "error"},
        {"key": "open_pr", "label": "미해결 문제", "ok": not prs, "value": f"{len(prs)}건", "sev": "warn"},
        {"key": "open_cr", "label": "이 릴리스 대상 미종결 CR", "ok": not crs, "value": ", ".join(r["id"] for r in crs) or "없음", "sev": "error"},
        {"key": "completeness", "label": "산출물 필수 컬럼 충족", "ok": bool(comp) and min(comp) >= 100,
         "value": f"최저 {min(comp) if comp else 0}%", "sev": "warn"},
        {"key": "commit", "label": "소스 커밋 기록", "ok": bool(rec["fields"].get("git_commit")), "value": rec["fields"].get("git_commit") or "없음", "sev": "error"},
    ]
    return out


def baseline_diff(repo, a: str, b: str | None) -> dict:
    A = repo.baseline(a)
    if A is None:
        raise KeyError(f"베이스라인 {a} 없음")
    B = repo.baseline(b) if b else {"items": repo.items(include_deleted=True), "links": repo.links()}
    if B is None:
        raise KeyError(f"베이스라인 {b} 없음")
    ai, bi = {i["id"]: i for i in A["items"]}, {i["id"]: i for i in B["items"]}
    rows = []
    for iid in sorted(set(ai) | set(bi)):
        x, y = ai.get(iid), bi.get(iid)
        if x and not y:
            rows.append({"id": iid, "change": "removed", "type": x["type"], "text": x["text"][:140]})
        elif y and not x:
            rows.append({"id": iid, "change": "added", "type": y["type"], "text": y["text"][:140]})
        elif x["text"] != y["text"] or x["status"] != y["status"] or x.get("version") != y.get("version"):
            rows.append({"id": iid, "change": "modified", "type": y["type"], "text": y["text"][:140], "prev": x["text"][:140],
                         "status": [x["status"], y["status"]], "version": [x.get("version"), y.get("version")]})
    al = {(l["src"], l["role"], l["dst"]) for l in A["links"]}
    bl = {(l["src"], l["role"], l["dst"]) for l in B["links"]}
    return {"a": a, "b": b or "(현재)", "rows": rows, "counts": Counter(r["change"] for r in rows),
            "links_added": len(bl - al), "links_removed": len(al - bl)}


def backup(repo, out: Path) -> Path:
    """Zip of the project data: a consistent copy of alm.db, evidence, assets, alm.json. Never the API key file."""
    with tempfile.TemporaryDirectory() as td:
        snap = Path(td) / "alm.db"
        dst = sqlite3.connect(snap)
        repo.con.commit()
        repo.con.backup(dst)
        dst.close()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(snap, ".alm/alm.db")
            for sub in ("evidence", "assets", "uploads"):
                d = repo.dir / sub
                if d.exists():
                    for f in d.rglob("*"):
                        if f.is_file():
                            z.write(f, f".alm/{f.relative_to(repo.dir).as_posix()}")
            cfgf = repo.root / "alm.json"
            if cfgf.exists():
                z.write(cfgf, "alm.json")
            z.writestr("BACKUP.txt", f"ECU-ALM backup {time.strftime('%Y-%m-%d %H:%M:%S')}\nrestore: ecualm --root <folder> restore <this zip>\n"
                                     "API 키(.alm/secrets.json)는 포함하지 않습니다.\n")
    return out


def restore(root: Path, zpath: Path, force: bool = False) -> dict:
    root = Path(root)
    db = root / ".alm" / "alm.db"
    if db.exists() and not force:
        raise SupError(f"{root} 에 이미 프로젝트 데이터가 있습니다. 덮어쓰려면 --force (현재 데이터는 .alm.bak 으로 옮깁니다)")
    with zipfile.ZipFile(zpath) as z:
        names = z.namelist()
        if ".alm/alm.db" not in names:
            raise SupError("ECU-ALM 백업 파일이 아닙니다")
        for n in names:
            if n.startswith("/") or ".." in Path(n).parts:
                raise SupError(f"잘못된 경로가 들어 있습니다: {n}")
        if (root / ".alm").exists():
            bak = root / ".alm.bak"
            if bak.exists():
                shutil.rmtree(bak)
            (root / ".alm").rename(bak)
        z.extractall(root)
    return {"restored": len(names), "root": str(root)}


# ---------------------------------------------------------------- listing, detail, dashboard
def list_records(repo, kind: str | None = None) -> list[dict]:
    q, a = "SELECT * FROM sup_records", []
    if kind:
        q += " WHERE kind=?"; a.append(kind)
    q += " ORDER BY created DESC"
    return [_row(r) for r in repo.con.execute(q, a)]


def history(repo, rid: str) -> list[dict]:
    return [dict(r, changes=json.loads(r["changes"])) for r in repo.con.execute("SELECT * FROM sup_history WHERE rec=? ORDER BY at", (rid,))]


def detail(repo, cfg, rid: str) -> dict:
    rec = get(repo, rid)
    K = KINDS[rec["kind"]]
    items = {i["id"]: i for i in repo.items(include_deleted=True)}
    rec["item_rows"] = [{**x, "type": (items.get(x["id"]) or {}).get("type", ""), "title": (items.get(x["id"]) or {}).get("title", "(삭제됨)"),
                         "status": (items.get(x["id"]) or {}).get("status", ""), "version": (items.get(x["id"]) or {}).get("version"),
                         "result": ((items.get(x["id"]) or {}).get("attrs", {}).get("result") or {}).get("status", "")} for x in rec["items"]]
    rec["ref_rows"] = []
    for r in rec["refs"]:
        try:
            o = get(repo, r)
            rec["ref_rows"].append({"id": o["id"], "kind": o["kind"], "title": o["title"], "status": KINDS[o["kind"]]["statuses"][o["status"]]})
        except KeyError:
            pass
    rec["history"] = history(repo, rid)
    rec["next"] = [{"to": t, "label": K["statuses"][t]} for t in K["flow"].get(rec["status"], [])]
    rec["status_label"] = K["statuses"][rec["status"]]
    if rec["kind"] == "REL" and rec["status"] == "draft":
        rec["checks"] = release_checks(repo, cfg, rec)
    rec["overdue"] = bool(rec["due"] and rec["due"] < dt.date.today().isoformat() and rec["status"] not in K["terminal"])
    return rec


def for_item(repo, iid: str) -> list[dict]:
    out = []
    for r in list_records(repo):
        roles = [x["role"] for x in r["items"] if x["id"] == iid]
        if roles or iid in json.dumps((r["fields"].get("impact") or {}).get("rows", [])):
            out.append({"id": r["id"], "kind": r["kind"], "title": r["title"], "status": KINDS[r["kind"]]["statuses"][r["status"]],
                        "roles": roles or ["impact"]})
    return out


def summary(repo, cfg) -> dict:
    today = dt.date.today()
    recs = list_records(repo)
    out = {"kinds": {k: {"label": v["label"], "process": v["process"], "statuses": v["statuses"], "terminal": v["terminal"],
                         "fields": v["fields"], "item_roles": v["item_roles"], "create_role": v["create_role"]} for k, v in KINDS.items()},
           "records": [{k: r[k] for k in ("id", "kind", "title", "status", "owner", "due", "created", "created_by", "updated", "updated_by")}
                       | {"severity": r["fields"].get("severity", ""), "safety": r["fields"].get("safety", ""),
                          "source": r["fields"].get("source", r["fields"].get("reason", "")), "process": r["fields"].get("process", r["fields"].get("target", "")),
                          "root_cause_cat": r["fields"].get("root_cause_cat", ""), "priority": r["fields"].get("priority", ""),
                          "sw_version": r["fields"].get("sw_version", ""), "items": len(r["items"]),
                          "overdue": bool(r["due"] and r["due"] < today.isoformat() and r["status"] not in KINDS[r["kind"]]["terminal"]),
                          "status_label": KINDS[r["kind"]]["statuses"][r["status"]]} for r in recs]}
    prs = [r for r in recs if r["kind"] == "PR"]
    closed_at = {}
    for h in repo.con.execute("SELECT rec, MAX(at) at FROM sup_history WHERE to_status IN ('closed','verified','released') GROUP BY rec"):
        closed_at[h["rec"]] = h["at"]
    ttr = [(closed_at[r["id"]] - r["created"]) / 86400 for r in prs if r["status"] == "closed" and r["id"] in closed_at]
    weeks = defaultdict(lambda: {"opened": 0, "closed": 0})
    for r in recs:
        if r["kind"] not in ("PR", "CR"):
            continue
        w = dt.date.fromtimestamp(r["created"]); w -= dt.timedelta(days=w.weekday())
        weeks[(r["kind"], w.isoformat())]["opened"] += 1
        if r["id"] in closed_at:
            c = dt.date.fromtimestamp(closed_at[r["id"]]); c -= dt.timedelta(days=c.weekday())
            weeks[(r["kind"], c.isoformat())]["closed"] += 1
    out["trend"] = sorted([{"kind": k, "week": w, **v} for (k, w), v in weeks.items()], key=lambda x: x["week"])
    openp = [r for r in prs if r["status"] not in KINDS["PR"]["terminal"]]
    out["pr"] = {"open": len(openp), "by_cause": Counter(r["fields"].get("root_cause_cat") or "미분석" for r in prs).most_common(),
                 "by_source": Counter(r["fields"].get("source", "") for r in prs).most_common(),
                 "by_severity_open": Counter(r["fields"].get("severity", "") for r in openp).most_common(),
                 "mttr_days": round(sum(ttr) / len(ttr), 1) if ttr else None,
                 "age_open": sorted([(r["id"], round((time.time() - r["created"]) / 86400, 1)) for r in openp], key=lambda x: -x[1])[:10]}
    out["baselines"] = repo.baselines()
    out["today"] = today.isoformat()
    return out


def audit_findings(repo, cfg) -> list[dict]:
    init(repo)
    out = []
    today = dt.date.today().isoformat()
    for r in list_records(repo):
        K = KINDS[r["kind"]]
        if r["status"] in K["terminal"]:
            continue
        if r["due"] and r["due"] < today:
            out.append({"sev": "warn", "rule": "sup_overdue", "item": r["id"], "msg": f"{K['process']} {K['label']} 기한 초과 ({r['due']}): {r['title']}", "related": []})
        if r["kind"] == "PR" and r["fields"].get("safety", "").startswith("예"):
            out.append({"sev": "warn", "rule": "safety_problem_open", "item": r["id"], "msg": f"안전 관련 문제 미해결: {r['title']}",
                        "related": [x["id"] for x in r["items"]]})
        if r["kind"] == "CR" and r["status"] in ("approved", "implementing", "verifying") and not r["fields"].get("impact"):
            out.append({"sev": "error", "rule": "cr_without_impact", "item": r["id"], "msg": "영향 분석 없이 승인된 변경 요청", "related": []})
    return out


def export_8d(repo, cfg, rid: str, out: Path) -> Path:
    from docx import Document
    from docx.shared import Pt
    r = detail(repo, cfg, rid)
    if r["kind"] != "PR":
        raise SupError("8D 보고서는 문제 기록에서만 만듭니다")
    f = r["fields"]
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Malgun Gothic"; st.font.size = Pt(10)
    d.add_heading(f"8D 보고서 — {r['id']} {r['title']}", 1)
    t = d.add_table(rows=0, cols=2); t.style = "Table Grid"
    for k, v in [("상태", r["status_label"]), ("출처", f.get("source", "")), ("심각도", f.get("severity", "")), ("안전 관련", f.get("safety", "")),
                 ("발견 빌드·샘플", f.get("found_in", "")), ("테스트 세션", f.get("session", "")), ("담당", r["owner"]), ("기한", r["due"]),
                 ("접수", time.strftime("%Y-%m-%d", time.localtime(r["created"])) + f" · {r['created_by']}")]:
        c = t.add_row().cells; c[0].text = k; c[1].text = str(v or "")
    for h, k in [("D1 팀 구성", "team"), ("D2 문제 기술", "description"), ("D3 임시 조치", "containment"), ("D4 근본 원인", "root_cause"),
                 ("수평 전개", "horizontal"), ("D5·D6 영구 조치 및 검증", "corrective"), ("D7 재발 방지", "prevention"), ("D8 종결·검증", "verification")]:
        d.add_heading(h, 2)
        if k == "root_cause" and f.get("root_cause_cat"):
            d.add_paragraph(f"원인 분류: {f['root_cause_cat']}")
        d.add_paragraph(f.get(k) or "—")
    d.add_heading("연결 항목", 2)
    for x in r["item_rows"]:
        d.add_paragraph(f"{KINDS['PR']['item_roles'].get(x['role'], x['role'])}: {x['id']} {x['title']} {('· ' + x['result']) if x['result'] else ''}", style="List Bullet")
    for x in r["ref_rows"]:
        d.add_paragraph(f"관련 기록: {x['id']} {x['title']} ({x['status']})", style="List Bullet")
    d.add_heading("이력", 2)
    for h in r["history"]:
        d.add_paragraph(f"{time.strftime('%Y-%m-%d %H:%M', time.localtime(h['at']))} {h['by']} · {h['action']}"
                        f"{(' ' + h['from_status'] + '→' + h['to_status']) if h['to_status'] else ''} {h['comment']}", style="List Bullet")
    d.save(out)
    return out


def export_xlsx(repo, cfg, kind: str, out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    K = KINDS[kind]
    wb = Workbook(); ws = wb.active; ws.title = kind
    cols = ["ID", "제목", "상태", "담당", "기한", "접수", "접수자"] + [f["label"] for f in K["fields"]] + ["연결 항목", "관련 기록"]
    for j, c in enumerate(cols, 1):
        x = ws.cell(1, j, c); x.font = Font(bold=True, color="FFFFFF"); x.fill = PatternFill("solid", fgColor="1F5FBF")
    for i, r in enumerate(list_records(repo, kind), 2):
        vals = [r["id"], r["title"], K["statuses"][r["status"]], r["owner"], r["due"], time.strftime("%Y-%m-%d", time.localtime(r["created"])),
                r["created_by"]] + [r["fields"].get(f["key"], "") for f in K["fields"]] + \
               [", ".join(f"{x['role']}:{x['id']}" for x in r["items"]), ", ".join(r["refs"])]
        for j, v in enumerate(vals, 1):
            ws.cell(i, j, "" if v is None else v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    wb.save(out)
    return out
