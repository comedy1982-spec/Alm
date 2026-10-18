"""Work-product tables: required columns per ASPICE work product + user-defined columns.

Every stage produces one work product (PAM 4.0 information item). Each has a set of columns that must be
filled for the work product to be complete — e.g. a system requirement needs an ID, the text, a verification
criterion, ASIL, upstream trace and verifying test; a test case needs a pre-condition, expected result and pass
criterion. Projects add their own columns per stage (stored in `wp_columns`; values in item attrs.custom).

Column sources (`src`):
  core:<field>      item field (id, text, title, status, version, domain, doc, section)
  attr:<a.b>        item attribute (dotted path into attrs)
  custom:<key>      user-defined column value (attrs.custom[key])
  up:<STAGES>       upstream linked items (this item -> role -> upstream), filtered by stage ids ('*' = any)
  down:<STAGES>     downstream linked items, filtered by stage ids
  result            latest test result status
"""
from __future__ import annotations

import json
import re
import time
from pathlib import Path

SCHEMA = """
CREATE TABLE IF NOT EXISTS wp_columns (
    stage    TEXT NOT NULL,
    key      TEXT NOT NULL,
    label    TEXT NOT NULL,
    type     TEXT NOT NULL DEFAULT 'text',
    options  TEXT NOT NULL DEFAULT '[]',
    required INTEGER NOT NULL DEFAULT 0,
    help     TEXT NOT NULL DEFAULT '',
    ord      REAL NOT NULL DEFAULT 0,
    created  REAL NOT NULL,
    created_by TEXT NOT NULL DEFAULT '',
    PRIMARY KEY (stage, key)
);
"""

TYPES = {"text": "텍스트", "long": "긴 텍스트", "enum": "선택 목록", "num": "숫자", "date": "날짜", "user": "담당자"}
STATUS_OPTS = ["draft", "review", "approved"]
ASIL_OPTS = ["QM", "A", "B", "C", "D"]
VERIFY_OPTS = ["Test", "Analysis", "Review", "Inspection", "Demonstration", "Simulation"]


def C(key, label, src, type_="text", req=False, edit=True, options=None, help_=""):
    return {"key": key, "label": label, "src": src, "type": type_, "required": req, "edit": edit,
            "options": options or [], "help": help_, "builtin": True}


_ID = C("id", "ID", "core:id", req=True, edit=False)
_STATUS = C("status", "상태", "core:status", "enum", True, options=STATUS_OPTS, help_="승인 상태 (draft·review·approved)")
_VER = C("version", "버전", "core:version", "num", False, False)
_DOM = C("domain", "도메인", "core:domain", edit=False)
_SRC = C("source", "원천 문서", "core:doc", edit=False)
_SEC = C("section", "섹션", "core:section", edit=False)
_NAT = C("native_id", "원문 ID", "attr:native_id", edit=False)
_ASIL = C("asil", "ASIL", "attr:asil", "enum", True, options=ASIL_OPTS, help_="ISO 26262 안전 등급 (QM 포함)")
_VC = C("verification_criteria", "검증 기준", "attr:verification_criteria", "long", True,
        help_="PAM 4.0: 요구사항마다 검증 기준을 정의 (SYS.2 BP / SWE.1 BP)")
_VM = C("verification_method", "검증 방법", "attr:verification_method", "enum", False, options=VERIFY_OPTS)

def _req(text_label="요구사항"):
    return C("text", text_label, "core:text", "long", True, help_="요구사항 본문 (단일·검증 가능한 문장)")


def _tc_cols(target: str, target_label: str):
    return [_ID, C("title", "시험명", "core:title", req=True), C("kind", "시험 종류", "attr:kind", "enum", True,
            options=["nominal", "boundary", "fault_injection", "design_verification", "integration", "regression"]),
            C("pre_condition", "사전 조건", "attr:pre_condition", "long", True),
            C("text", "절차·입력", "core:text", "long", True),
            C("expected", "기대 결과", "attr:expected", "long", True),
            C("pass_criteria", "합격 기준", "attr:pass_criteria", "long", True),
            C("tool", "시험 환경·도구", "attr:tool", req=False),
            C("verifies", f"검증 대상 ({target_label})", f"up:{target}", "ids", True, False,
              help_="PAM 4.0: 검증 수단은 대상 산출물과 양방향 추적 (13-51 일관성 증거)"),
            C("result", "최근 결과", "result", "text", True, False, help_="15-52 검증 결과 (PASS/FAIL/BLOCKED)"),
            _STATUS, _VER, _DOM, _SRC]


# stage -> (work product title, PAM 4.0 information items, columns)
WORK_PRODUCTS = {
    "SYS1": ("고객 요구사항 명세 (CRS)", "17-00 요구사항, 13-52 의사소통 증거", [
        _ID, _NAT, _req("고객 요구사항"), C("compliance", "준수 여부", "attr:compliance", "enum", True,
                                       options=["C (준수)", "PC (부분 준수)", "NC (미준수)", "검토중"],
                                       help_="고객 요구사항별 합의 결과 (Compliance matrix)"),
        C("comment", "공급사 의견", "attr:comment", "long", False),
        C("derived", "SYS.2 전개", "down:SYS2", "ids", True, False, help_="SYS.1 → SYS.2 추적 (13-51)"),
        _STATUS, _ASIL | {"required": False}, _SRC, _SEC, _VER]),
    "SYS2": ("시스템 요구사항 명세 (SyRS)", "17-00 요구사항, 17-54 요구사항 속성, 13-51 일관성 증거", [
        _ID, _NAT, _req(), C("discipline", "분야 (HW/SW)", "attr:discipline", "enum", True, options=["SW", "HW", "HW/SYS", "MECH", "SYS"]),
        _ASIL, _VC, _VM, C("upstream", "상위 (CRS)", "up:SYS1", "ids", True, False, help_="고객 요구사항과 양방향 추적"),
        C("downstream", "하위 (SYS.3·HWE.1·SWE.1)", "down:SYS3,HWE1,SWE1", "ids", False, False),
        C("tests", "검증 TC (SYS.5)", "down:SYS5", "ids", True, False, help_="SYS.5 검증 수단으로 100% 커버"),
        _STATUS, _VER, _DOM, _SRC, _SEC]),
    "SYS3": ("시스템 아키텍처 (SAD)·HSI", "04-06 시스템 아키텍처, 13-51 일관성 증거", [
        _ID, C("title", "요소명", "core:title", req=True), C("text", "설명·역할", "core:text", "long", True),
        C("type", "유형", "core:type", edit=False), C("allocated", "할당 요구사항 (SYS.2)", "up:SYS2", "ids", True, False),
        _ASIL | {"required": False},
        C("tests", "통합 TC (SYS.4)", "down:SYS4", "ids", True, False), _STATUS, _VER, _DOM]),
    "HWE1": ("HW 요구사항 명세 (HWRS)", "17-00 요구사항, 13-51 일관성 증거", [
        _ID, _req("HW 요구사항"), _ASIL, _VC, C("upstream", "상위 (SYS.2)", "up:SYS2", "ids", True, False),
        C("blocks", "HW 블록 (HWE.2)", "down:HWE2", "ids", True, False), C("tests", "적격성 TC (HWE.4)", "down:HWE4", "ids", True, False),
        _STATUS, _VER, _DOM, _SRC]),
    "HWE2": ("HW 아키텍처·설계 (HWAD)", "04-52 HW 아키텍처, 04-53 HW 상세 설계", [
        _ID, C("title", "블록명", "core:title", req=True), C("text", "설명", "core:text", "long", True),
        C("part_number", "부품 번호", "attr:hw.part_number", req=True), _ASIL,
        C("safety_mechanism", "안전 메커니즘", "attr:hw.safety_mechanism", "long", False),
        C("allocated", "할당 요구사항 (HWE.1)", "up:HWE1", "ids", True, False),
        C("tests", "설계 검증 TC (HWE.3)", "down:HWE3", "ids", True, False), _STATUS, _VER, _DOM]),
    "SWE1": ("SW 요구사항 명세 (SWRS)", "17-00 요구사항, 17-54 요구사항 속성, 13-51 일관성 증거", [
        _ID, _NAT, _req("SW 요구사항"), _ASIL,
        C("pre", "Pre-condition", "attr:spec.pre_condition", "long", True), C("trigger", "Trigger", "attr:spec.trigger", "long", False),
        C("timing", "Timing", "attr:spec.timing", req=True), C("output", "Output", "attr:spec.output", "long", True),
        _VC, C("upstream", "상위 (SYS.2)", "up:SYS2,SYS3", "ids", True, False),
        C("arch", "SW 컴포넌트 (SWE.2)", "down:SWE2", "ids", True, False), C("tests", "검증 TC (SWE.6)", "down:SWE6", "ids", True, False),
        _STATUS, _VER, _DOM, _SRC]),
    "SWE2": ("SW 아키텍처 (SWAD)", "04-04 SW 아키텍처, 13-51 일관성 증거", [
        _ID, C("title", "컴포넌트", "core:title", req=True), C("text", "책임·동작", "core:text", "long", True),
        C("layer", "AUTOSAR 계층", "attr:autosar_layer", "enum", True, options=["ASW", "RTE", "BSW", "CDD", "MCAL"]),
        C("cycle", "실행 주기", "attr:runnable_cycle", req=True), _ASIL | {"required": False},
        C("allocated", "할당 요구사항 (SWE.1)", "up:SWE1", "ids", True, False),
        C("units", "유닛 (SWE.3)", "down:SWE3", "ids", False, False), C("tests", "통합 TC (SWE.5)", "down:SWE5", "ids", True, False),
        _STATUS, _VER, _DOM]),
    "SWE3": ("SW 상세 설계 (SWDD)·유닛", "04-05 SW 상세 설계, 11-05 SW 유닛", [
        _ID, C("title", "유닛", "core:title", req=True), C("text", "동작 설명", "core:text", "long", True),
        C("file", "소스 파일", "attr:c_file", req=True), C("signature", "함수 시그니처", "attr:function_signature", req=True),
        C("complexity", "복잡도 목표", "attr:complexity_target", "num", False),
        C("component", "상위 컴포넌트 (SWE.2)", "up:SWE2", "ids", True, False), C("tests", "유닛 TC (SWE.4)", "down:SWE4", "ids", True, False),
        _STATUS, _VER, _DOM]),
    "SWE4": ("SW 유닛 검증 명세·결과", "08-60 검증 수단, 15-52 검증 결과", _tc_cols("SWE3", "SWE.3 유닛")),
    "SWE5": ("SW 통합 검증 명세·결과", "08-60 검증 수단, 15-52 검증 결과", _tc_cols("SWE2", "SWE.2 컴포넌트")),
    "SWE6": ("SW 검증 명세·결과", "08-60 검증 수단, 15-52 검증 결과", _tc_cols("SWE1", "SWE.1 요구사항")),
    "HWE3": ("HW 설계 검증 명세·결과", "08-60 검증 수단, 15-52 검증 결과", _tc_cols("HWE2", "HWE.2 블록")),
    "HWE4": ("HW 적격성 시험 명세·결과", "08-60 검증 수단, 15-52 검증 결과", _tc_cols("HWE1", "HWE.1 요구사항")),
    "SYS4": ("시스템 통합 검증 명세·결과", "08-60 검증 수단, 15-52 검증 결과", _tc_cols("SYS3", "SYS.3 요소")),
    "SYS5": ("시스템 검증 명세·결과", "08-60 검증 수단, 15-52 검증 결과", _tc_cols("SYS2", "SYS.2 요구사항")),
}
# attributes that are bookkeeping, not requirement content: changing them does not make downstream items Suspect
NO_SUSPECT = {"compliance", "comment", "verification_method", "tool", "discipline", "complexity_target"}
FIELD_ALIASES = {"asil": ["asil"], "verification_method": ["검증방법", "verification", "verificationmethod", "검증"],
                 "verification_criteria": ["검증기준", "acceptancecriteria"], "comment": ["비고", "comment"]}


def init(repo):
    repo.con.executescript(SCHEMA)
    repo.con.commit()


def custom_columns(repo, stage: str) -> list[dict]:
    init(repo)
    out = []
    for r in repo.con.execute("SELECT * FROM wp_columns WHERE stage=? ORDER BY ord, created", (stage,)):
        out.append({"key": r["key"], "label": r["label"], "src": f"custom:{r['key']}", "type": r["type"],
                    "required": bool(r["required"]), "edit": True, "options": json.loads(r["options"]),
                    "help": r["help"], "builtin": False, "created_by": r["created_by"]})
    return out


def columns(repo, cfg: dict, stage: str) -> tuple[str, str, list[dict]]:
    if stage not in {s["id"] for s in cfg["stages"]}:
        raise KeyError(stage)
    title, info, cols = WORK_PRODUCTS.get(stage, (stage, "", [_ID, _req(), _STATUS]))
    return title, info, [dict(c) for c in cols] + custom_columns(repo, stage)


def _slug(label: str, taken: set[str]) -> str:
    base = re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_") or "col"
    k, n = f"c_{base}"[:40], 2
    while k in taken:
        k, n = f"c_{base}_{n}"[:40], n + 1
    return k


def add_column(repo, cfg, stage: str, label: str, type_: str = "text", options=None, required=False, help_="", user="") -> dict:
    label = (label or "").strip()
    if not label or len(label) > 40:
        raise ValueError("컬럼 이름은 1~40자")
    if type_ not in TYPES:
        raise ValueError(f"유형은 {', '.join(TYPES)} 중 하나")
    _, _, cols = columns(repo, cfg, stage)
    if any(c["label"] == label for c in cols):
        raise ValueError("같은 이름의 컬럼이 있습니다")
    opts = [str(o).strip() for o in (options or []) if str(o).strip()]
    if type_ == "enum" and not opts:
        raise ValueError("선택 목록 컬럼에는 항목을 하나 이상 넣으세요")
    key = _slug(label, {c["key"] for c in cols})
    n = repo.con.execute("SELECT COUNT(*) FROM wp_columns WHERE stage=?", (stage,)).fetchone()[0]
    repo.con.execute("INSERT INTO wp_columns VALUES(?,?,?,?,?,?,?,?,?,?)",
                     (stage, key, label, type_, json.dumps(opts, ensure_ascii=False), int(bool(required)), help_ or "", n, time.time(), user))
    repo.con.commit()
    return next(c for c in custom_columns(repo, stage) if c["key"] == key)


def update_column(repo, stage: str, key: str, label=None, required=None, options=None, help_=None) -> dict:
    r = repo.con.execute("SELECT * FROM wp_columns WHERE stage=? AND key=?", (stage, key)).fetchone()
    if not r:
        raise KeyError(f"{stage}.{key} (기본 필수 컬럼은 바꿀 수 없습니다)")
    if label is not None:
        repo.con.execute("UPDATE wp_columns SET label=? WHERE stage=? AND key=?", (label.strip()[:40] or r["label"], stage, key))
    if required is not None:
        repo.con.execute("UPDATE wp_columns SET required=? WHERE stage=? AND key=?", (int(bool(required)), stage, key))
    if options is not None:
        repo.con.execute("UPDATE wp_columns SET options=? WHERE stage=? AND key=?",
                         (json.dumps([str(o).strip() for o in options if str(o).strip()], ensure_ascii=False), stage, key))
    if help_ is not None:
        repo.con.execute("UPDATE wp_columns SET help=? WHERE stage=? AND key=?", (help_, stage, key))
    repo.con.commit()
    return next(c for c in custom_columns(repo, stage) if c["key"] == key)


def delete_column(repo, stage: str, key: str) -> dict:
    n = repo.con.execute("DELETE FROM wp_columns WHERE stage=? AND key=?", (stage, key)).rowcount
    repo.con.commit()
    if not n:
        raise KeyError(f"{stage}.{key} (기본 필수 컬럼은 지울 수 없습니다)")
    return {"deleted": key, "note": "입력된 값은 아이템 속성에 남아 있어 같은 이름으로 다시 만들면 복원됩니다"}


# ---------------------------------------------------------------- values
def _dig(d: dict, path: str):
    for p in path.split("."):
        if not isinstance(d, dict):
            return None
        d = d.get(p)
    return d


def _stage_types(cfg, stage_ids: str) -> set[str] | None:
    if stage_ids in ("*", ""):
        return None
    want = set(stage_ids.split(","))
    return {t for s in cfg["stages"] if s["id"] in want for t in s["types"]}


def _ctx(repo):
    items = {i["id"]: i for i in repo.items()}
    up, down = {}, {}
    for l in repo.links():
        if l["src"] in items and l["dst"] in items:
            up.setdefault(l["src"], []).append(l)
            down.setdefault(l["dst"], []).append(l)
    return items, up, down


def value(it: dict, col: dict, cfg: dict, ctx) -> object:
    items, up, down = ctx
    kind, _, arg = col["src"].partition(":")
    if kind == "core":
        return it.get(arg)
    if kind == "attr":
        v = _dig(it["attrs"], arg)
        if v in (None, "") and arg in FIELD_ALIASES:       # imported table rows keep columns in attrs.fields
            f = {str(k).replace(" ", "").lower(): x for k, x in (it["attrs"].get("fields") or {}).items()}
            v = next((f[a] for a in FIELD_ALIASES[arg] if f.get(a)), v)
        return v
    if kind == "custom":
        return (it["attrs"].get("custom") or {}).get(arg)
    if kind in ("up", "down"):
        types = _stage_types(cfg, arg)
        ls = up.get(it["id"], []) if kind == "up" else down.get(it["id"], [])
        other = "dst" if kind == "up" else "src"
        out = []
        for l in ls:
            o = items[l[other]]
            if types is None or o["type"] in types:
                out.append({"id": o["id"], "suspect": bool(l["suspect"]), "proposed": l["state"] == "proposed"})
        return out
    if kind == "result":
        r = it["attrs"].get("result") or {}
        return r.get("status") or ""
    return None


def empty(v) -> bool:
    if v is None:
        return True
    if isinstance(v, (list, dict)):
        return not v or (isinstance(v, list) and all(x.get("proposed") for x in v if isinstance(x, dict)))
    s = str(v).strip()
    return not s or s.startswith("[") and s.endswith("]")   # template placeholders like "[전제 조건]"


def table(repo, cfg: dict, stage: str) -> dict:
    title, info, cols = columns(repo, cfg, stage)
    st = next(s for s in cfg["stages"] if s["id"] == stage)
    ctx = _ctx(repo)
    rows, missing_total, per_col = [], 0, {c["key"]: 0 for c in cols}
    for it in ctx[0].values():
        if it["type"] not in st["types"]:
            continue
        vals, miss = {}, []
        for c in cols:
            v = value(it, c, cfg, ctx)
            vals[c["key"]] = v
            if c["required"] and empty(v):
                miss.append(c["key"]); per_col[c["key"]] += 1
        missing_total += len(miss)
        rows.append({"id": it["id"], "type": it["type"], "values": vals, "missing": miss, "ord": (it["doc"], it["ord"])})
    rows.sort(key=lambda r: r["ord"])
    for r in rows:
        r.pop("ord")
    req_cols = [c for c in cols if c["required"]]
    filled = sum(len(req_cols) - len(r["missing"]) for r in rows)
    return {"stage": stage, "work_product": title, "info_items": info, "columns": cols, "rows": rows,
            "types": TYPES, "missing": missing_total, "missing_by_column": per_col,
            "complete_rows": sum(1 for r in rows if not r["missing"]),
            "completeness": round(filled * 100 / (len(req_cols) * len(rows)), 1) if rows and req_cols else 100.0}


def set_cell(repo, cfg, item_id: str, key: str, val, user: str) -> dict:
    """Write one editable cell through ops.edit_item (keeps history, lint and suspect propagation)."""
    from . import ops
    it = repo.get(item_id)
    if not it:
        raise KeyError(item_id)
    st = next((s for s in cfg["stages"] if it["type"] in s["types"]), None)
    _, _, cols = columns(repo, cfg, st["id"])
    col = next((c for c in cols if c["key"] == key), None)
    if not col or not col["edit"]:
        raise ValueError(f"{key} 는 편집할 수 없는 컬럼입니다")
    if col["type"] == "enum" and val not in ("", None) and col["options"] and val not in col["options"]:
        raise ValueError(f"{col['label']}: {', '.join(col['options'])} 중에서 고르세요")
    if col["type"] == "num" and val not in ("", None):
        val = float(val)
        val = int(val) if val.is_integer() else val
    if col["type"] == "date" and val and not re.match(r"^\d{4}-\d{2}-\d{2}$", str(val)):
        raise ValueError("날짜 형식은 YYYY-MM-DD")
    kind, _, arg = col["src"].partition(":")
    if kind == "core":
        if arg == "text":
            return ops.edit_item(repo, cfg, item_id, text=str(val or ""), user=user)
        if arg == "status":
            return ops.edit_item(repo, cfg, item_id, status=str(val), user=user)
        if arg == "title":
            r = repo.update(item_id, {"title": str(val or "")}, change="edited", by=user, mark_suspect=False)
            repo.commit()
            return r
    if kind == "custom":       # project bookkeeping: downstream items do not need re-review
        return ops.edit_item(repo, cfg, item_id, custom={arg: val}, user=user, suspect=False)
    if kind == "attr":
        parts = arg.split(".")
        sus = None if arg not in NO_SUSPECT else False
        if len(parts) == 1:
            return ops.edit_item(repo, cfg, item_id, attrs={arg: val}, user=user, suspect=sus)
        top = dict(it["attrs"].get(parts[0]) or {})
        cur = top
        for p in parts[1:-1]:
            cur[p] = dict(cur.get(p) or {}); cur = cur[p]
        cur[parts[-1]] = val
        return ops.edit_item(repo, cfg, item_id, attrs={parts[0]: top}, user=user, suspect=sus)
    raise ValueError("편집할 수 없는 컬럼")


def audit_findings(repo, cfg) -> list[dict]:
    out = []
    for s in cfg["stages"]:
        try:
            t = table(repo, cfg, s["id"])
        except KeyError:
            continue
        labels = {c["key"]: c["label"] for c in t["columns"]}
        for r in t["rows"]:
            if r["missing"]:
                out.append({"sev": "warn", "rule": "wp_required", "item": r["id"],
                            "msg": f"{t['work_product']} 필수 컬럼 누락: {', '.join(labels[k] for k in r['missing'])}", "related": []})
    return out


def export_xlsx(repo, cfg, stage: str, out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    t = table(repo, cfg, stage)
    wb = Workbook(); ws = wb.active; ws.title = stage
    ws["A1"] = f"{t['work_product']}  —  {t['info_items']}"; ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"{len(t['rows'])}건 · 필수 컬럼 충족률 {t['completeness']}% · 누락 {t['missing']}칸 (빨간 칸)  ·  * = 필수"
    thin = Side(style="thin", color="C9D1DB"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    for j, c in enumerate(t["columns"], 1):
        h = ws.cell(4, j, c["label"] + (" *" if c["required"] else ""))
        h.font = Font(bold=True, color="FFFFFF"); h.fill = PatternFill("solid", fgColor="1F5FBF" if c["required"] else "6B7785")
        h.alignment = Alignment(wrap_text=True, vertical="center"); h.border = bd
        ws.column_dimensions[get_column_letter(j)].width = 48 if c["type"] == "long" else 22 if c["type"] == "ids" else 14
    miss = PatternFill("solid", fgColor="FDF3F2")
    for i, r in enumerate(t["rows"], 5):
        for j, c in enumerate(t["columns"], 1):
            v = r["values"][c["key"]]
            if isinstance(v, list):
                v = ", ".join(x["id"] + (" (Suspect)" if x["suspect"] else "") for x in v if not x.get("proposed"))
            cell = ws.cell(i, j, "" if v is None else v)
            cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.border = bd
            if c["key"] in r["missing"]:
                cell.fill = miss
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(t['columns']))}{4 + len(t['rows'])}"
    wb.save(out)
    return out
