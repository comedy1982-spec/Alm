"""Project plan in the ASPICE MAN.3 form: WBS, schedule, effort estimates and actuals, baselines, EVM.

A plan row ("task") carries the WBS number, the activity, the ASPICE process it belongs to, the work products
(PAM 4.0 information items) it produces, the owner, planned dates and effort (man-days), actual effort,
progress, predecessors (finish-to-start) and the baseline it was last approved at. Rows whose WBS number
has children ("2" over "2.1", "2.2") are summary rows: their dates, effort and progress are rolled up.

Effort is spread evenly over the working days (Mon–Fri minus holidays) of a task to get the weekly load
per owner. Earned value at "today": PV from the baseline (or plan) dates, EV = effort × progress,
AC = actual effort; SPI = EV / PV, CPI = EV / AC.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import time
from collections import defaultdict
from pathlib import Path

SCHEMA = """
CREATE TABLE IF NOT EXISTS plan_tasks (
    id            TEXT PRIMARY KEY,
    wbs           TEXT NOT NULL DEFAULT '',
    name          TEXT NOT NULL DEFAULT '',
    process       TEXT NOT NULL DEFAULT '',
    stage         TEXT NOT NULL DEFAULT '',
    work_product  TEXT NOT NULL DEFAULT '',
    owner         TEXT NOT NULL DEFAULT '',
    start         TEXT NOT NULL DEFAULT '',
    finish        TEXT NOT NULL DEFAULT '',
    effort_plan   REAL NOT NULL DEFAULT 0,
    effort_actual REAL NOT NULL DEFAULT 0,
    progress      REAL NOT NULL DEFAULT 0,
    status        TEXT NOT NULL DEFAULT 'not_started',
    milestone     INTEGER NOT NULL DEFAULT 0,
    deps          TEXT NOT NULL DEFAULT '[]',
    actual_start  TEXT NOT NULL DEFAULT '',
    actual_finish TEXT NOT NULL DEFAULT '',
    base_start    TEXT NOT NULL DEFAULT '',
    base_finish   TEXT NOT NULL DEFAULT '',
    base_effort   REAL,
    note          TEXT NOT NULL DEFAULT '',
    updated       REAL NOT NULL DEFAULT 0,
    updated_by    TEXT NOT NULL DEFAULT ''
);
CREATE TABLE IF NOT EXISTS plan_history (
    task TEXT NOT NULL, at REAL NOT NULL, by TEXT NOT NULL, change TEXT NOT NULL, fields TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS plan_baselines (
    name TEXT PRIMARY KEY, at REAL NOT NULL, by TEXT NOT NULL, note TEXT NOT NULL, snapshot TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS plan_meta (key TEXT PRIMARY KEY, value TEXT NOT NULL);
"""

STATUSES = {"not_started": "미착수", "in_progress": "진행", "done": "완료", "on_hold": "보류"}
PROCESSES = ["MAN.3", "MAN.5", "MAN.6", "SYS.1", "SYS.2", "SYS.3", "SYS.4", "SYS.5", "HWE.1", "HWE.2", "HWE.3", "HWE.4",
             "SWE.1", "SWE.2", "SWE.3", "SWE.4", "SWE.5", "SWE.6", "SUP.1", "SUP.8", "SUP.9", "SUP.10", "VAL.1", "기타"]
EDIT_FIELDS = ("wbs", "name", "process", "stage", "work_product", "owner", "start", "finish", "effort_plan",
               "effort_actual", "progress", "status", "milestone", "deps", "actual_start", "actual_finish", "note")
SELF_FIELDS = ("effort_actual", "progress", "status", "actual_start", "actual_finish", "note")   # engineer on own task
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


class PlanError(ValueError):
    pass


def init(repo):
    repo.con.executescript(SCHEMA)
    repo.con.commit()


# ---------------------------------------------------------------- calendar
def _d(s: str) -> dt.date:
    return dt.date.fromisoformat(s)


def holidays(repo) -> set[str]:
    return set(meta(repo).get("holidays", []))


def workdays(a: str, b: str, hol: set[str]) -> list[dt.date]:
    if not a or not b:
        return []
    d, e, out = _d(a), _d(b), []
    while d <= e:
        if d.weekday() < 5 and d.isoformat() not in hol:
            out.append(d)
        d += dt.timedelta(days=1)
    return out


def week_of(d: dt.date) -> str:
    return (d - dt.timedelta(days=d.weekday())).isoformat()


# ---------------------------------------------------------------- meta
def meta(repo) -> dict:
    m = {r["key"]: json.loads(r["value"]) for r in repo.con.execute("SELECT key,value FROM plan_meta")}
    m.setdefault("holidays", [])
    m.setdefault("capacity", 5.0)        # MD per person per week
    m.setdefault("title", "")
    m.setdefault("doc_no", "")
    return m


def set_meta(repo, data: dict, user: str) -> dict:
    allowed = {"holidays", "capacity", "title", "doc_no", "today"}
    for k, v in data.items():
        if k not in allowed:
            continue
        if k == "holidays":
            v = sorted({str(x).strip() for x in v if _DATE.match(str(x).strip())})
        if k == "capacity":
            v = float(v)
            if not 0 < v <= 7:
                raise PlanError("주간 가용 공수는 0 초과 7 이하 (MD/주)")
        repo.con.execute("INSERT INTO plan_meta VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                         (k, json.dumps(v, ensure_ascii=False)))
    repo.con.commit()
    return meta(repo)


# ---------------------------------------------------------------- tasks
def _row(r) -> dict:
    d = dict(r)
    d["deps"] = json.loads(d["deps"] or "[]")
    d["milestone"] = bool(d["milestone"])
    return d


def tasks(repo) -> list[dict]:
    rows = [_row(r) for r in repo.con.execute("SELECT * FROM plan_tasks")]
    rows.sort(key=lambda t: _wbs_key(t["wbs"]))
    return rows


def _wbs_key(w: str):
    return [int(x) if x.isdigit() else x for x in re.split(r"[.\-]", w or "999")]


def _summary_ids(rows: list[dict]) -> set[str]:
    wbs = [t["wbs"] for t in rows if t["wbs"]]
    return {t["id"] for t in rows if t["wbs"] and any(w.startswith(t["wbs"] + ".") for w in wbs)}


def _next_id(repo) -> str:
    n = 0
    for (i,) in repo.con.execute("SELECT id FROM plan_tasks"):
        m = re.match(r"T(\d+)$", i)
        if m:
            n = max(n, int(m.group(1)))
    return f"T{n + 1:03d}"


def _hist(repo, task, by, change, fields):
    repo.con.execute("INSERT INTO plan_history VALUES(?,?,?,?,?)",
                     (task, time.time(), by, change, json.dumps(fields, ensure_ascii=False)))


def _clean(data: dict, cur: dict | None, rows: list[dict]) -> dict:
    out = {}
    for k in EDIT_FIELDS:
        if k not in data:
            continue
        v = data[k]
        if k in ("effort_plan", "effort_actual"):
            v = round(float(v or 0), 2)
            if v < 0 or v > 100000:
                raise PlanError("공수는 0 이상 (MD)")
        elif k == "progress":
            v = float(v or 0)
            if not 0 <= v <= 100:
                raise PlanError("진척률은 0 ~ 100 %")
        elif k == "status":
            if v not in STATUSES:
                raise PlanError(f"상태는 {', '.join(STATUSES)} 중 하나")
        elif k == "milestone":
            v = bool(v)
        elif k == "deps":
            v = [str(x).strip() for x in (v if isinstance(v, list) else str(v).split(",")) if str(x).strip()]
        elif k in ("start", "finish", "actual_start", "actual_finish"):
            v = str(v or "").strip()
            if v and not _DATE.match(v):
                raise PlanError(f"{k}: 날짜 형식은 YYYY-MM-DD")
            if v:
                _d(v)
        else:
            v = str(v or "").strip()
        out[k] = v
    m = {**(cur or {}), **out}
    if m.get("milestone"):
        if m.get("start"):
            out["finish"] = m["start"]
        if not cur:
            out.setdefault("effort_plan", 0.0)
    if m.get("start") and m.get("finish") and m["finish"] < m["start"]:
        raise PlanError("종료일이 시작일보다 빠릅니다")
    if "progress" in out and "status" not in data:
        if out["progress"] >= 100:
            out["status"] = "done"
        elif out["progress"] > 0 and (cur or {}).get("status", "not_started") == "not_started":
            out["status"] = "in_progress"
    if out.get("status") == "done" and "progress" not in data:
        out["progress"] = 100.0
    if "deps" in out:
        ids = {t["id"] for t in rows}
        by_wbs = {t["wbs"]: t["id"] for t in rows if t["wbs"]}
        deps = []
        for x in out["deps"]:
            x = x if x in ids else by_wbs.get(x, x)
            if x not in ids:
                raise PlanError(f"선행 작업 {x} 를 찾을 수 없습니다")
            if cur and x == cur["id"]:
                raise PlanError("자기 자신을 선행 작업으로 둘 수 없습니다")
            deps.append(x)
        out["deps"] = deps
    return out


def _cycle(rows: list[dict], tid: str, deps: list[str]) -> bool:
    g = {t["id"]: t["deps"] for t in rows}
    g[tid] = deps
    seen, stack = set(), list(deps)
    while stack:
        x = stack.pop()
        if x == tid:
            return True
        if x in seen:
            continue
        seen.add(x)
        stack.extend(g.get(x, []))
    return False


def can_edit(user: dict, task: dict | None, fields) -> bool:
    from .auth import rank
    if rank(user["role"]) >= rank("manager"):
        return True
    if user["role"] == "engineer" and task is not None:
        mine = task["owner"] in (user["username"], user.get("name"))
        return mine and all(f in SELF_FIELDS for f in fields)
    return False


def save_task(repo, data: dict, user: dict) -> dict:
    rows = tasks(repo)
    tid = data.get("id")
    cur = next((t for t in rows if t["id"] == tid), None) if tid else None
    if tid and not cur:
        raise PlanError(f"작업 {tid} 없음")
    fields = [k for k in data if k in EDIT_FIELDS]
    if not can_edit(user, cur, fields):
        raise PermissionError("이 작업을 고칠 권한이 없습니다 (엔지니어는 본인 작업의 진척·실적만 입력)")
    vals = _clean(data, cur, rows)
    if cur:
        if "deps" in vals and _cycle(rows, cur["id"], vals["deps"]):
            raise PlanError("선행 관계가 순환합니다")
        changed = {k: [cur.get(k), v] for k, v in vals.items() if cur.get(k) != v}
        if not changed:
            return cur
        sets = {k: (json.dumps(v) if k == "deps" else int(v) if k == "milestone" else v) for k, v in vals.items()}
        sets.update(updated=time.time(), updated_by=user["username"])
        repo.con.execute(f"UPDATE plan_tasks SET {', '.join(f'{k}=?' for k in sets)} WHERE id=?", (*sets.values(), cur["id"]))
        _hist(repo, cur["id"], user["username"], "modified", changed)
        tid = cur["id"]
    else:
        if not vals.get("name"):
            raise PlanError("작업명을 입력하세요")
        tid = _next_id(repo)
        if not vals.get("wbs"):
            top = [int(t["wbs"].split(".")[0]) for t in rows if t["wbs"].split(".")[0].isdigit()]
            vals["wbs"] = str(max(top or [0]) + 1)
        row = {"id": tid, **{k: v for k, v in vals.items()}}
        row["deps"] = json.dumps(row.get("deps", []))
        row["milestone"] = int(row.get("milestone", False))
        row.update(updated=time.time(), updated_by=user["username"])
        repo.con.execute(f"INSERT INTO plan_tasks({', '.join(row)}) VALUES({', '.join('?' * len(row))})", tuple(row.values()))
        _hist(repo, tid, user["username"], "created", {k: [None, v] for k, v in vals.items()})
    repo.con.commit()
    return _row(repo.con.execute("SELECT * FROM plan_tasks WHERE id=?", (tid,)).fetchone())


def delete_task(repo, tid: str, user: dict) -> dict:
    if not can_edit(user, None, ["wbs"]):
        raise PermissionError("작업 삭제는 PM·관리자만 할 수 있습니다")
    cur = repo.con.execute("SELECT * FROM plan_tasks WHERE id=?", (tid,)).fetchone()
    if not cur:
        raise PlanError(f"작업 {tid} 없음")
    repo.con.execute("DELETE FROM plan_tasks WHERE id=?", (tid,))
    for t in tasks(repo):
        if tid in t["deps"]:
            repo.con.execute("UPDATE plan_tasks SET deps=? WHERE id=?", (json.dumps([x for x in t["deps"] if x != tid]), t["id"]))
    _hist(repo, tid, user["username"], "deleted", {"name": [cur["name"], None], "wbs": [cur["wbs"], None]})
    repo.con.commit()
    return {"deleted": tid}


def history(repo, tid: str | None = None, limit: int = 200) -> list[dict]:
    q, a = "SELECT * FROM plan_history", []
    if tid:
        q += " WHERE task=?"; a.append(tid)
    q += " ORDER BY at DESC LIMIT ?"; a.append(limit)
    return [dict(r, fields=json.loads(r["fields"])) for r in repo.con.execute(q, a)]


# ---------------------------------------------------------------- baselines
def set_baseline(repo, name: str, note: str, user: dict) -> dict:
    if not can_edit(user, None, ["wbs"]):
        raise PermissionError("기준선 저장은 PM·관리자만 할 수 있습니다")
    name = (name or "").strip()
    if not name:
        raise PlanError("기준선 이름을 입력하세요")
    if repo.con.execute("SELECT 1 FROM plan_baselines WHERE name=?", (name,)).fetchone():
        raise PlanError("같은 이름의 기준선이 있습니다")
    rows = tasks(repo)
    repo.con.execute("INSERT INTO plan_baselines VALUES(?,?,?,?,?)",
                     (name, time.time(), user["username"], note or "", json.dumps(rows, ensure_ascii=False)))
    repo.con.execute("UPDATE plan_tasks SET base_start=start, base_finish=finish, base_effort=effort_plan")
    _hist(repo, "*", user["username"], "baseline", {"name": [None, name], "tasks": [None, len(rows)]})
    repo.con.commit()
    return {"name": name, "tasks": len(rows)}


def baselines(repo) -> list[dict]:
    return [dict(name=r["name"], at=r["at"], by=r["by"], note=r["note"])
            for r in repo.con.execute("SELECT name,at,by,note FROM plan_baselines ORDER BY at")]


# ---------------------------------------------------------------- roll-up, load, EVM
def _today(repo) -> dt.date:
    t = meta(repo).get("today")
    return _d(t) if t else dt.date.today()


def computed(repo, cfg: dict | None = None, stage_counts: dict | None = None) -> dict:
    rows = tasks(repo)
    hol = holidays(repo)
    today = _today(repo)
    summ = _summary_ids(rows)
    leaves = [t for t in rows if t["id"] not in summ]
    # roll up summary rows
    for t in rows:
        t["summary"] = t["id"] in summ
        if t["summary"]:
            kids = [k for k in leaves if k["wbs"].startswith(t["wbs"] + ".")]
            st = [k["start"] for k in kids if k["start"]]
            fi = [k["finish"] for k in kids if k["finish"]]
            t["start"], t["finish"] = (min(st) if st else ""), (max(fi) if fi else "")
            bs = [k["base_start"] for k in kids if k["base_start"]]
            bf = [k["base_finish"] for k in kids if k["base_finish"]]
            t["base_start"], t["base_finish"] = (min(bs) if bs else ""), (max(bf) if bf else "")
            t["effort_plan"] = round(sum(k["effort_plan"] for k in kids), 2)
            t["effort_actual"] = round(sum(k["effort_actual"] for k in kids), 2)
            w = sum(k["effort_plan"] for k in kids)
            t["progress"] = round(sum(k["effort_plan"] * k["progress"] for k in kids) / w, 1) if w else 0
            t["status"] = "done" if kids and all(k["status"] == "done" for k in kids) else (
                "in_progress" if any(k["status"] != "not_started" for k in kids) else "not_started")
    # flags
    for t in rows:
        flags = []
        if not t["summary"] and t["status"] != "done":
            if t["finish"] and _d(t["finish"]) < today and t["progress"] < 100:
                flags.append("late")
            elif t["start"] and _d(t["start"]) < today and t["progress"] == 0 and not t["milestone"]:
                flags.append("late_start")
        if t["base_finish"] and t["finish"] and t["finish"] > t["base_finish"]:
            flags.append("slip")
        if not t["summary"] and t["effort_plan"] and t["effort_actual"] > t["effort_plan"] * 1.1:
            flags.append("overrun")
        byid = {x["id"]: x for x in rows}
        for d in t["deps"]:
            p = byid.get(d)
            if not t["summary"] and p and p["finish"] and t["start"] and (
                    p["finish"] > t["start"] if p["milestone"] else p["finish"] >= t["start"]):
                flags.append("dep_conflict")
                break
        t["flags"] = flags
        t["slip_days"] = (_d(t["finish"]) - _d(t["base_finish"])).days if t["base_finish"] and t["finish"] else 0
        if t.get("stage") and stage_counts and stage_counts.get(t["stage"], {}).get("total"):
            c = stage_counts[t["stage"]]
            tot = c.get("total", 0)
            appr = c.get("status", {}).get("approved", 0)
            t["alm"] = {"total": tot, "approved": appr, "pct": round(appr * 100 / tot) if tot else 0}
    # weekly load per owner (planned effort spread evenly over working days)
    load: dict = defaultdict(lambda: defaultdict(float))
    weeks = set()
    for t in leaves:
        days = workdays(t["start"], t["finish"], hol)
        if not days or not t["effort_plan"]:
            continue
        per = t["effort_plan"] / len(days)
        for d in days:
            w = week_of(d)
            load[t["owner"] or "(미지정)"][w] += per
            weeks.add(w)
    load_out = {o: {w: round(v, 2) for w, v in sorted(ws.items())} for o, ws in sorted(load.items())}
    # earned value
    pv = ev = ac = bac = 0.0
    for t in leaves:
        e = t["base_effort"] if t["base_effort"] is not None else t["effort_plan"]
        s, f = (t["base_start"] or t["start"]), (t["base_finish"] or t["finish"])
        bac += e
        days = workdays(s, f, hol)
        if days:
            done = sum(1 for d in days if d <= today)
            pv += e * done / len(days)
        elif f and _d(f) <= today:
            pv += e
        ev += e * t["progress"] / 100
        ac += t["effort_actual"]
    nxt = sorted((t for t in rows if t["milestone"] and t["start"] and _d(t["start"]) >= today), key=lambda t: t["start"])
    st = [t["start"] for t in rows if t["start"]]
    fi = [t["finish"] for t in rows if t["finish"]]
    return {
        "tasks": rows, "meta": meta(repo), "baselines": baselines(repo), "today": today.isoformat(),
        "range": [min(st) if st else today.isoformat(), max(fi) if fi else today.isoformat()],
        "load": load_out, "weeks": sorted(weeks), "statuses": STATUSES, "processes": PROCESSES,
        "metrics": {"bac": round(bac, 1), "pv": round(pv, 1), "ev": round(ev, 1), "ac": round(ac, 1),
                    "spi": round(ev / pv, 2) if pv else None, "cpi": round(ev / ac, 2) if ac else None,
                    "progress": round(ev * 100 / bac, 1) if bac else 0,
                    "late": sum(1 for t in rows if "late" in t["flags"] or "late_start" in t["flags"]),
                    "overrun": sum(1 for t in rows if "overrun" in t["flags"]),
                    "next_milestone": {"name": nxt[0]["name"], "date": nxt[0]["start"]} if nxt else None},
    }


# ---------------------------------------------------------------- standard ASPICE WBS
# (wbs, name, process, ALM stage, work products, owner, start week, weeks, effort MD, predecessors, milestone)
TEMPLATE = [
    ("1", "프로젝트 관리", "MAN.3", "", "", "", 0, 0, 0, [], False),
    ("1.1", "프로젝트 착수·작업 범위·타당성 검토", "MAN.3", "", "08-53 작업 범위, 08-54 타당성 분석", "PM", 0, 2, 8, [], False),
    ("1.2", "WBS·일정·공수 계획 수립", "MAN.3", "", "14-10 작업 패키지, 08-56 일정", "PM", 2, 2, 6, ["1.1"], False),
    ("1.3", "프로젝트 모니터링·주간 보고", "MAN.3", "", "15-06 프로젝트 상태, 13-52 의사소통 증거", "PM", 4, 59, 30, ["1.2"], False),
    ("1.4", "위험 식별·대응 관리", "MAN.5", "", "위험 목록, 대응 계획", "PM", 2, 61, 15, [], False),
    ("2", "시스템 엔지니어링", "SYS.2", "", "", "", 0, 0, 0, [], False),
    ("2.1", "SYS.1 고객 요구사항 접수·검토", "SYS.1", "SYS1", "17-00 요구사항, 13-52 의사소통 증거", "SYS", 0, 4, 20, [], False),
    ("2.2", "SYS.2 시스템 요구사항 분석", "SYS.2", "SYS2", "17-00 요구사항, 17-54 요구사항 속성, 13-51 일관성 증거", "SYS", 4, 6, 45, ["2.1"], False),
    ("2.3", "SYS.3 시스템 아키텍처·HSI 설계", "SYS.3", "SYS3", "04-06 시스템 아키텍처, 13-51 일관성 증거", "SYS", 10, 6, 40, ["2.2"], False),
    ("2.4", "SYS.4 시스템 통합·통합 검증", "SYS.4", "SYS4", "08-60 검증 수단, 15-52 검증 결과", "검증", 50, 6, 50, ["3.4", "4.6"], False),
    ("2.5", "SYS.5 시스템 검증 (차량 포함)", "SYS.5", "SYS5", "08-60 검증 수단, 15-52 검증 결과", "검증", 56, 6, 60, ["2.4"], False),
    ("3", "HW 엔지니어링", "HWE.2", "", "", "", 0, 0, 0, [], False),
    ("3.1", "HWE.1 HW 요구사항 분석", "HWE.1", "HWE1", "17-00 요구사항, 13-51 일관성 증거", "HW", 10, 4, 25, ["2.2"], False),
    ("3.2", "HWE.2 HW 설계 (회로도·블록도)", "HWE.2", "HWE2", "04-52 HW 아키텍처, 04-53 HW 상세 설계", "HW", 14, 8, 60, ["3.1"], False),
    ("3.3", "HWE.3 HW 설계 검증", "HWE.3", "HWE3", "08-60 검증 수단, 15-52 검증 결과", "HW", 22, 6, 30, ["3.2"], False),
    ("3.4", "HWE.4 HW 적격성 시험 (환경·EMC)", "HWE.4", "HWE4", "08-60 검증 수단, 15-52 검증 결과", "HW", 30, 12, 45, ["3.3"], False),
    ("4", "SW 엔지니어링", "SWE.2", "", "", "", 0, 0, 0, [], False),
    ("4.1", "SWE.1 SW 요구사항 분석", "SWE.1", "SWE1", "17-00 요구사항, 17-54 요구사항 속성", "SW 아키텍트", 10, 5, 40, ["2.2"], False),
    ("4.2", "SWE.2 SW 아키텍처 설계 (AUTOSAR)", "SWE.2", "SWE2", "04-04 SW 아키텍처, ARXML", "SW 아키텍트", 16, 6, 45, ["4.1", "2.3"], False),
    ("4.3", "SWE.3 SW 상세 설계·구현", "SWE.3", "SWE3", "04-05 SW 상세 설계, 11-05 SW 유닛", "SW 개발", 22, 14, 160, ["4.2"], False),
    ("4.4", "SWE.4 SW 유닛 검증 (MC/DC)", "SWE.4", "SWE4", "08-60 검증 수단, 15-52 검증 결과", "SW 개발", 26, 14, 70, [], False),
    ("4.5", "SWE.5 SW 통합·통합 검증", "SWE.5", "SWE5", "01-50 통합 SW, 08-60 검증 수단, 15-52 검증 결과", "통합", 36, 8, 50, ["4.3"], False),
    ("4.6", "SWE.6 SW 검증 (HIL)", "SWE.6", "SWE6", "08-60 검증 수단, 15-52 검증 결과", "검증", 44, 6, 55, ["4.5"], False),
    ("5", "지원 프로세스", "SUP.1", "", "", "", 0, 0, 0, [], False),
    ("5.1", "SUP.1 품질 보증 (공정·산출물 감사)", "SUP.1", "", "13-18 품질 기록, 14-02 시정 조치", "QA", 2, 60, 30, [], False),
    ("5.2", "SUP.8 형상 관리 (베이스라인·릴리스)", "SUP.8", "", "06-52 백업·복구 기록, 18-53 형상 항목 선택 기준", "CM", 2, 60, 20, [], False),
    ("5.3", "SUP.9 문제 해결 관리", "SUP.9", "", "13-07 문제 기록, 15-55 문제 분석", "검증", 22, 40, 15, [], False),
    ("5.4", "SUP.10 변경 요청 관리", "SUP.10", "", "13-16 변경 요청, 18-57 변경 분석 기준", "PM", 4, 58, 15, [], False),
    ("6", "마일스톤", "MAN.3", "", "", "", 0, 0, 0, [], False),
    ("6.1", "프로젝트 킥오프", "MAN.3", "", "킥오프 회의록", "PM", 0, 0, 0, [], True),
    ("6.2", "A 샘플 출하", "MAN.3", "", "A 샘플, 릴리스 노트", "PM", 18, 0, 0, ["2.3"], True),
    ("6.3", "B 샘플 출하", "MAN.3", "", "B 샘플, 추적 매트릭스", "PM", 36, 0, 0, ["4.3"], True),
    ("6.4", "C 샘플 출하", "MAN.3", "", "C 샘플, FMEDA", "PM", 50, 0, 0, ["4.6"], True),
    ("6.5", "D 샘플 (PPAP)", "MAN.3", "", "PPAP, 안전 사례", "PM", 62, 0, 0, ["2.5"], True),
]


def template(repo, start: str, user: dict, replace: bool = False) -> dict:
    if not can_edit(user, None, ["wbs"]):
        raise PermissionError("표준 WBS 생성은 PM·관리자만 할 수 있습니다")
    if not _DATE.match(start or ""):
        raise PlanError("시작일 형식은 YYYY-MM-DD")
    if tasks(repo) and not replace:
        raise PlanError("이미 작업이 있습니다. 덮어쓰려면 replace 를 선택하세요")
    if replace:
        repo.con.execute("DELETE FROM plan_tasks")
        _hist(repo, "*", user["username"], "replaced", {})
    d0 = _d(start)
    d0 -= dt.timedelta(days=d0.weekday())          # Monday of the start week
    ids = {}
    for i, (wbs, *_r) in enumerate(TEMPLATE, 1):
        ids[wbs] = f"T{i:03d}"
    now = time.time()
    for i, (wbs, name, proc, stage, wp, owner, w0, wn, eff, deps, ms) in enumerate(TEMPLATE, 1):
        s = d0 + dt.timedelta(weeks=w0)
        f = s if ms else s + dt.timedelta(weeks=wn) - dt.timedelta(days=3)
        summary = wn == 0 and not ms
        repo.con.execute(
            "INSERT INTO plan_tasks(id,wbs,name,process,stage,work_product,owner,start,finish,effort_plan,milestone,deps,updated,updated_by)"
            " VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (ids[wbs], wbs, name, proc, stage, wp, owner, "" if summary else s.isoformat(), "" if summary else f.isoformat(),
             float(eff), int(ms), json.dumps([ids[x] for x in deps]), now, user["username"]))
    _hist(repo, "*", user["username"], "template", {"start": [None, d0.isoformat()], "tasks": [None, len(TEMPLATE)]})
    repo.con.commit()
    return {"created": len(TEMPLATE), "start": d0.isoformat()}


# ---------------------------------------------------------------- Excel (ASPICE MAN.3 form)
def export_xlsx(repo, cfg: dict, out: Path, stage_counts: dict | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    data = computed(repo, cfg, stage_counts)
    rows, m = data["tasks"], data["meta"]
    wb = Workbook()
    ws = wb.active
    ws.title = "일정·공수 계획"
    thin = Side(style="thin", color="C9D1DB")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr = PatternFill("solid", fgColor="1F5FBF")
    sumf = PatternFill("solid", fgColor="EEF1F5")
    bar = {"SYS": "0F6E7C", "HWE": "A8501E", "SWE": "1F5FBF", "SUP": "6B7785", "MAN": "6B4BB5", "VAL": "1E7A50"}
    proj = m.get("title") or str(cfg.get("project", ""))
    ws["A1"] = "프로젝트 일정·공수 계획서 (ASPICE MAN.3)"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"프로젝트: {proj}    문서번호: {m.get('doc_no') or '-'}    기준일: {data['today']}    " \
               f"기준선: {data['baselines'][-1]['name'] if data['baselines'] else '없음'}"
    mt = data["metrics"]
    ws["A3"] = (f"총 계획 공수(BAC) {mt['bac']} MD · 계획 가치(PV) {mt['pv']} · 획득 가치(EV) {mt['ev']} · 실적(AC) {mt['ac']} · "
                f"SPI {mt['spi'] if mt['spi'] is not None else '-'} · CPI {mt['cpi'] if mt['cpi'] is not None else '-'} · 진척 {mt['progress']}%")
    cols = ["WBS", "작업 (Activity)", "ASPICE 프로세스", "산출물 (Information Item)", "담당", "계획 시작", "계획 종료", "기간(일)",
            "계획 공수(MD)", "실적 공수(MD)", "진척률(%)", "상태", "선행 작업", "기준선 종료", "지연(일)", "ALM 승인률", "비고"]
    widths = [7, 34, 10, 34, 11, 11, 11, 7, 9, 9, 8, 8, 10, 11, 7, 9, 24]
    R0 = 5
    # week columns
    weeks = []
    if data["range"][0]:
        w = _d(data["range"][0]); w -= dt.timedelta(days=w.weekday())
        end = _d(data["range"][1])
        while w <= end:
            weeks.append(w); w += dt.timedelta(weeks=1)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(R0, j, c)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = hdr; cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = bd
        ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    W0 = len(cols) + 1
    for k, w in enumerate(weeks):
        c = ws.cell(R0, W0 + k, f"{w.month}/{w.day}")
        c.font = Font(bold=True, color="FFFFFF", size=8); c.fill = hdr; c.alignment = Alignment(horizontal="center", text_rotation=90); c.border = bd
        ws.column_dimensions[get_column_letter(W0 + k)].width = 3.2
        if w.day <= 7:
            ws.cell(R0 - 1, W0 + k, f"{w.year % 100}.{w.month:02d}").font = Font(size=8, bold=True)
    idwbs = {t["id"]: t["wbs"] for t in rows}
    for i, t in enumerate(rows, R0 + 1):
        dur = (_d(t["finish"]) - _d(t["start"])).days + 1 if t["start"] and t["finish"] and not t["milestone"] else ""
        vals = [t["wbs"], ("◆ " if t["milestone"] else "") + t["name"], t["process"], t["work_product"], t["owner"], t["start"], t["finish"], dur,
                t["effort_plan"] or "", t["effort_actual"] or "", t["progress"], STATUSES.get(t["status"], t["status"]),
                ", ".join(idwbs.get(d, d) for d in t["deps"]), t["base_finish"], t["slip_days"] or "",
                f"{t['alm']['pct']}% ({t['alm']['approved']}/{t['alm']['total']})" if t.get("alm") else "", t["note"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v); c.border = bd; c.alignment = Alignment(vertical="center", wrap_text=j in (2, 4, 17))
            if t["summary"]:
                c.fill = sumf; c.font = Font(bold=True)
        if "late" in t["flags"]:
            ws.cell(i, 12).font = Font(bold=True, color="B3261E")
            ws.cell(i, 12).value = "지연"
        color = bar.get(t["process"][:3], "6B7785")
        for k, w in enumerate(weeks):
            c = ws.cell(i, W0 + k); c.border = Border(left=Side(style="hair", color="E3E8EE"), bottom=Side(style="hair", color="E3E8EE"))
            if not t["start"]:
                continue
            we = w + dt.timedelta(days=6)
            if t["milestone"]:
                if w <= _d(t["start"]) <= we:
                    c.value = "◆"; c.font = Font(color="B3261E", bold=True); c.alignment = Alignment(horizontal="center")
            elif _d(t["start"]) <= we and _d(t["finish"]) >= w:
                c.fill = PatternFill("solid", fgColor="9AA6B4" if t["summary"] else color)
    ws.freeze_panes = ws.cell(R0 + 1, 3)
    ws.sheet_view.zoomScale = 85
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{R0}:{R0}"
    ws.row_dimensions[R0].height = 34
    # load sheet
    ls = wb.create_sheet("공수 부하")
    ls["A1"] = f"담당별 주간 계획 공수 (MD) — 가용 {m['capacity']} MD/주 초과 시 빨강"; ls["A1"].font = Font(bold=True)
    ls.cell(3, 1, "담당").font = Font(bold=True)
    for k, w in enumerate(data["weeks"]):
        ls.cell(3, 2 + k, w[5:]).font = Font(bold=True, size=8)
        ls.column_dimensions[get_column_letter(2 + k)].width = 5.5
    ls.column_dimensions["A"].width = 14
    for r, (o, ws_) in enumerate(data["load"].items(), 4):
        ls.cell(r, 1, o)
        for k, w in enumerate(data["weeks"]):
            v = ws_.get(w)
            if v:
                c = ls.cell(r, 2 + k, round(v, 1))
                c.font = Font(size=8, color="B3261E" if v > m["capacity"] else "16202C", bold=v > m["capacity"])
                c.fill = PatternFill("solid", fgColor="FDF3F2" if v > m["capacity"] else "E8EFFA")
    # change history (MAN.3 plan adjustments)
    hs = wb.create_sheet("계획 변경 이력")
    for j, c in enumerate(["일시", "작업", "변경자", "구분", "변경 내용"], 1):
        hs.cell(1, j, c).font = Font(bold=True)
    for r, h in enumerate(history(repo, limit=2000), 2):
        hs.cell(r, 1, time.strftime("%Y-%m-%d %H:%M", time.localtime(h["at"])))
        hs.cell(r, 2, idwbs.get(h["task"], h["task"])); hs.cell(r, 3, h["by"]); hs.cell(r, 4, h["change"])
        hs.cell(r, 5, "; ".join(f"{k}: {v[0]} → {v[1]}" for k, v in h["fields"].items()))
    for col, wd in zip("ABCDE", (16, 8, 12, 10, 90)):
        hs.column_dimensions[col].width = wd
    wb.save(out)
    return out
