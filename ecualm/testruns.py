"""Test sessions (test runs), per-case results, evidence files and the test dashboard.

A session is one execution campaign: a verification level (SWE.4 … SYS.5), the build / HW sample / bench it ran
on, a set of test cases and their outcome. Recording a case result also updates the test case itself
(attrs.result, through arch.ingest_results), so the audit rules test_fail / result_stale keep working and the
item drawer shows the latest outcome. Evidence (logs, screenshots, waveforms, reports) is stored content-addressed
under .alm/evidence with its SHA-256, so the file named in a report can be shown to be the one that was uploaded.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import mimetypes
import re
import time
from collections import Counter, defaultdict
from pathlib import Path

from . import arch

SCHEMA = """
CREATE TABLE IF NOT EXISTS test_sessions (
    id          TEXT PRIMARY KEY,
    name        TEXT NOT NULL,
    stage       TEXT NOT NULL,
    build       TEXT NOT NULL DEFAULT '',
    sample      TEXT NOT NULL DEFAULT '',
    environment TEXT NOT NULL DEFAULT '',
    description TEXT NOT NULL DEFAULT '',
    planned_start TEXT NOT NULL DEFAULT '',
    planned_end   TEXT NOT NULL DEFAULT '',
    status      TEXT NOT NULL DEFAULT 'planned',
    owner       TEXT NOT NULL DEFAULT '',
    created     REAL NOT NULL,
    created_by  TEXT NOT NULL DEFAULT '',
    started     REAL,
    closed      REAL,
    closed_by   TEXT NOT NULL DEFAULT '',
    verdict     TEXT NOT NULL DEFAULT ''
);
CREATE TABLE IF NOT EXISTS session_cases (
    session     TEXT NOT NULL,
    tc          TEXT NOT NULL,
    ord         INTEGER NOT NULL DEFAULT 0,
    assignee    TEXT NOT NULL DEFAULT '',
    status      TEXT NOT NULL DEFAULT 'NOT_RUN',
    actual      TEXT NOT NULL DEFAULT '',
    defect      TEXT NOT NULL DEFAULT '',
    duration_s  REAL,
    executed_at REAL,
    executed_by TEXT NOT NULL DEFAULT '',
    tc_version  INTEGER,
    history     TEXT NOT NULL DEFAULT '[]',
    PRIMARY KEY (session, tc)
);
CREATE TABLE IF NOT EXISTS evidence (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    session     TEXT NOT NULL,
    tc          TEXT NOT NULL DEFAULT '',
    filename    TEXT NOT NULL,
    path        TEXT NOT NULL,
    mime        TEXT NOT NULL,
    size        INTEGER NOT NULL,
    sha256      TEXT NOT NULL,
    note        TEXT NOT NULL DEFAULT '',
    at          REAL NOT NULL,
    by          TEXT NOT NULL DEFAULT ''
);
CREATE INDEX IF NOT EXISTS ix_ev_session ON evidence(session, tc);
"""

CASE_STATES = ("NOT_RUN", "PASS", "FAIL", "BLOCKED", "N_A")
SESSION_STATES = {"planned": "계획", "in_progress": "진행", "completed": "완료", "aborted": "중단"}
MAX_EVIDENCE = 50 * 1024 * 1024
INLINE_MIME = {"image/png", "image/jpeg", "image/gif", "image/webp", "application/pdf", "text/plain", "text/csv"}


class TestError(ValueError):
    pass


def init(repo):
    repo.con.executescript(SCHEMA)
    repo.con.commit()


def _verify_stages(cfg) -> dict:
    return {s["id"]: s for s in cfg["stages"] if s.get("side") == "verify"}


def _sid(repo) -> str:
    n = 0
    for (i,) in repo.con.execute("SELECT id FROM test_sessions"):
        m = re.match(r"TS-(\d+)$", i)
        if m:
            n = max(n, int(m.group(1)))
    return f"TS-{n + 1:03d}"


def _session(repo, sid: str) -> dict:
    r = repo.con.execute("SELECT * FROM test_sessions WHERE id=?", (sid,)).fetchone()
    if not r:
        raise KeyError(f"세션 {sid} 없음")
    return dict(r)


def _role(user) -> int:
    from .auth import rank
    return rank(user["role"])


# ---------------------------------------------------------------- candidate selection
def candidates(repo, cfg, stage: str, domain: str = "", only: str = "all", q: str = "") -> list[dict]:
    """Test cases of a verification stage, filtered. only: all | not_passed | failed | never_run | stale"""
    vs = _verify_stages(cfg)
    if stage not in vs:
        raise TestError(f"{stage} 는 검증 단계가 아닙니다 ({', '.join(vs)})")
    types = set(vs[stage]["types"])
    items = {i["id"]: i for i in repo.items()}
    out = []
    for it in items.values():
        if it["type"] not in types or (domain and it["domain"] != domain):
            continue
        res = it["attrs"].get("result") or {}
        st = res.get("status", "NOT_RUN")
        stale = any(items.get(d) and v is not None and items[d]["version"] > v
                    for d, v in (res.get("verified_versions") or {}).items())
        if only == "not_passed" and st == "PASS" and not stale:
            continue
        if only == "failed" and st not in ("FAIL", "BLOCKED"):
            continue
        if only == "never_run" and res:
            continue
        if only == "stale" and not stale:
            continue
        if q and q.lower() not in (it["id"] + " " + it["title"] + " " + it["text"]).lower():
            continue
        out.append({"id": it["id"], "title": it["title"], "domain": it["domain"], "last": st, "stale": stale})
    out.sort(key=lambda x: x["id"])
    return out


# ---------------------------------------------------------------- sessions
def create_session(repo, cfg, b: dict, user: dict) -> dict:
    stage = b.get("stage", "")
    name = (b.get("name") or "").strip()
    if not name:
        raise TestError("세션 이름을 입력하세요")
    ids = b.get("tcs")
    if ids is None:
        ids = [c["id"] for c in candidates(repo, cfg, stage, b.get("domain", ""), b.get("only", "all"), b.get("q", ""))]
    else:
        vs = _verify_stages(cfg)
        if stage not in vs:
            raise TestError("검증 단계를 고르세요")
        ok = set(vs[stage]["types"])
        bad = [i for i in ids if not (repo.get(i) or {}).get("type") in ok]
        if bad:
            raise TestError(f"이 단계의 TC 가 아닙니다: {', '.join(bad[:5])}")
    if not ids:
        raise TestError("조건에 맞는 TC 가 없습니다")
    for k in ("planned_start", "planned_end"):
        if b.get(k) and not re.match(r"^\d{4}-\d{2}-\d{2}$", b[k]):
            raise TestError("날짜 형식은 YYYY-MM-DD")
    sid = _sid(repo)
    now = time.time()
    repo.con.execute("INSERT INTO test_sessions(id,name,stage,build,sample,environment,description,planned_start,planned_end,"
                     "status,owner,created,created_by) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                     (sid, name, stage, b.get("build", "").strip(), b.get("sample", "").strip(), b.get("environment", "").strip(),
                      b.get("description", "").strip(), b.get("planned_start", ""), b.get("planned_end", ""), "planned",
                      (b.get("owner") or user["username"]).strip(), now, user["username"]))
    for k, tc in enumerate(ids):
        repo.con.execute("INSERT INTO session_cases(session,tc,ord,assignee) VALUES(?,?,?,?)", (sid, tc, k, b.get("assignee", "")))
    repo.con.commit()
    return session(repo, cfg, sid)


def update_session(repo, cfg, sid: str, b: dict, user: dict) -> dict:
    s = _session(repo, sid)
    action = b.get("action")
    now = time.time()
    if action == "start":
        if s["status"] != "planned":
            raise TestError("계획 상태의 세션만 시작할 수 있습니다")
        repo.con.execute("UPDATE test_sessions SET status='in_progress', started=? WHERE id=?", (now, sid))
    elif action in ("complete", "abort"):
        if s["status"] in ("completed", "aborted"):
            raise TestError("이미 닫힌 세션입니다")
        cnt = Counter(r["status"] for r in repo.con.execute("SELECT status FROM session_cases WHERE session=?", (sid,)))
        verdict = "FAIL" if cnt["FAIL"] else "BLOCKED" if cnt["BLOCKED"] or cnt["NOT_RUN"] else "PASS"
        repo.con.execute("UPDATE test_sessions SET status=?, closed=?, closed_by=?, verdict=?, started=COALESCE(started, ?) WHERE id=?",
                         ("completed" if action == "complete" else "aborted", now, user["username"],
                          verdict if action == "complete" else "ABORTED", now, sid))
    elif action == "reopen":
        if _role(user) < 2:
            raise PermissionError("닫힌 세션은 PM·관리자만 다시 열 수 있습니다")
        repo.con.execute("UPDATE test_sessions SET status='in_progress', closed=NULL, closed_by='', verdict='' WHERE id=?", (sid,))
    elif action == "edit":
        fields = {k: str(b[k]).strip() for k in ("name", "build", "sample", "environment", "description", "planned_start",
                                                  "planned_end", "owner") if k in b}
        if fields:
            repo.con.execute(f"UPDATE test_sessions SET {', '.join(f'{k}=?' for k in fields)} WHERE id=?", (*fields.values(), sid))
    elif action == "add_cases":
        if s["status"] in ("completed", "aborted"):
            raise TestError("닫힌 세션에는 TC 를 추가할 수 없습니다")
        have = {r["tc"] for r in repo.con.execute("SELECT tc FROM session_cases WHERE session=?", (sid,))}
        n = len(have)
        ok = set(_verify_stages(cfg)[s["stage"]]["types"])
        for tc in b.get("tcs", []):
            if tc not in have and (repo.get(tc) or {}).get("type") in ok:
                repo.con.execute("INSERT INTO session_cases(session,tc,ord) VALUES(?,?,?)", (sid, tc, n)); n += 1
    elif action == "remove_case":
        if s["status"] in ("completed", "aborted"):
            raise TestError("닫힌 세션입니다")
        repo.con.execute("DELETE FROM session_cases WHERE session=? AND tc=? AND status='NOT_RUN'", (sid, b.get("tc")))
    elif action == "delete":
        if _role(user) < 2:
            raise PermissionError("세션 삭제는 PM·관리자만 할 수 있습니다")
        if repo.con.execute("SELECT 1 FROM evidence WHERE session=?", (sid,)).fetchone() or s["status"] == "completed":
            raise TestError("완료됐거나 증빙이 있는 세션은 지울 수 없습니다 (중단 처리하세요)")
        repo.con.execute("DELETE FROM session_cases WHERE session=?", (sid,))
        repo.con.execute("DELETE FROM test_sessions WHERE id=?", (sid,))
        repo.con.commit()
        return {"deleted": sid}
    else:
        raise TestError("알 수 없는 동작")
    repo.con.commit()
    return session(repo, cfg, sid)


def record(repo, cfg, sid: str, b: dict, user: dict) -> dict:
    """Set one case result. Also pushes PASS/FAIL/BLOCKED to the test case (arch.ingest_results)."""
    s = _session(repo, sid)
    if s["status"] in ("completed", "aborted"):
        raise TestError("닫힌 세션의 결과는 바꿀 수 없습니다 (PM 이 다시 열 수 있습니다)")
    tc = b.get("tc")
    row = repo.con.execute("SELECT * FROM session_cases WHERE session=? AND tc=?", (sid, tc)).fetchone()
    if not row:
        raise KeyError(f"{sid} 에 {tc} 없음")
    row = dict(row)
    sets = {}
    if "status" in b:
        st = str(b["status"]).upper()
        if st not in CASE_STATES:
            raise TestError(f"결과는 {', '.join(CASE_STATES)} 중 하나")
        sets["status"] = st
    for k in ("actual", "defect", "assignee"):
        if k in b:
            sets[k] = str(b[k] or "").strip()
    if "duration_s" in b:
        sets["duration_s"] = float(b["duration_s"]) if b["duration_s"] not in ("", None) else None
    if not sets:
        return case_row(repo, sid, tc)
    if sets.get("status") == "FAIL" and not (sets.get("actual") or row["actual"]):
        raise TestError("FAIL 은 실제 결과(관찰 내용)를 함께 적어야 합니다")
    now = time.time()
    it = repo.get(tc)
    if "status" in sets:
        sets.update(executed_at=now if sets["status"] != "NOT_RUN" else None, executed_by=user["username"],
                    tc_version=it["version"] if it else None)
        hist = json.loads(row["history"]) + [{"status": sets["status"], "at": now, "by": user["username"]}]
        sets["history"] = json.dumps(hist[-30:])
    repo.con.execute(f"UPDATE session_cases SET {', '.join(f'{k}=?' for k in sets)} WHERE session=? AND tc=?",
                     (*sets.values(), sid, tc))
    if s["status"] == "planned" and sets.get("status", "NOT_RUN") != "NOT_RUN":
        repo.con.execute("UPDATE test_sessions SET status='in_progress', started=? WHERE id=?", (now, sid))
    repo.con.commit()
    final = {**row, **sets}
    if sets.get("status") in ("PASS", "FAIL", "BLOCKED") or ("status" not in sets and final["status"] in ("PASS", "FAIL", "BLOCKED")):
        arch.ingest_results(repo, cfg, {"results": [{
            "tc": tc, "status": final["status"], "run_id": sid, "tool": s["environment"],
            "note": final["actual"] + (f" · 결함 {final['defect']}" if final["defect"] else ""),
            "metrics": {"build": s["build"], "sample": s["sample"], **({"duration_s": final["duration_s"]} if final["duration_s"] else {})}}]},
            user=user["username"])
    return case_row(repo, sid, tc)


def case_row(repo, sid, tc) -> dict:
    r = dict(repo.con.execute("SELECT * FROM session_cases WHERE session=? AND tc=?", (sid, tc)).fetchone())
    r["history"] = json.loads(r["history"])
    return r


def _counts(rows) -> dict:
    c = Counter(r["status"] for r in rows)
    n = len(rows)
    run = n - c["NOT_RUN"]
    eff = n - c["N_A"]
    return {"total": n, **{k: c[k] for k in CASE_STATES}, "executed": run - c["N_A"],
            "progress": round(run * 100 / n, 1) if n else 0,
            "pass_rate": round(c["PASS"] * 100 / (run - c["N_A"]), 1) if run - c["N_A"] > 0 else None,
            "effective": eff}


def sessions(repo, cfg) -> list[dict]:
    out = []
    for s in repo.con.execute("SELECT * FROM test_sessions ORDER BY created DESC"):
        s = dict(s)
        rows = [dict(r) for r in repo.con.execute("SELECT status FROM session_cases WHERE session=?", (s["id"],))]
        s["counts"] = _counts(rows)
        s["evidence"] = repo.con.execute("SELECT COUNT(*) FROM evidence WHERE session=?", (s["id"],)).fetchone()[0]
        out.append(s)
    return out


def session(repo, cfg, sid: str) -> dict:
    s = _session(repo, sid)
    items = {i["id"]: i for i in repo.items(include_deleted=True)}
    ev = defaultdict(list)
    for e in repo.con.execute("SELECT * FROM evidence WHERE session=? ORDER BY at", (sid,)):
        ev[e["tc"]].append({k: e[k] for k in ("id", "tc", "filename", "mime", "size", "sha256", "note", "at", "by")})
    verifies = defaultdict(list)
    for l in repo.links():
        if l["role"] == "verifies":
            verifies[l["src"]].append(l["dst"])
    cases = []
    for r in repo.con.execute("SELECT * FROM session_cases WHERE session=? ORDER BY ord", (sid,)):
        r = dict(r)
        r["history"] = json.loads(r["history"])
        it = items.get(r["tc"]) or {}
        a = it.get("attrs", {})
        r.update(title=it.get("title", "(삭제됨)"), text=it.get("text", ""), domain=it.get("domain", ""),
                 kind=a.get("kind", ""), pre_condition=a.get("pre_condition", ""), expected=a.get("expected", ""),
                 pass_criteria=a.get("pass_criteria", ""), verifies=verifies.get(r["tc"], []),
                 changed_since=bool(r["tc_version"] and it.get("version", 0) > r["tc_version"]),
                 evidence=ev.get(r["tc"], []))
        cases.append(r)
    s["cases"] = cases
    s["counts"] = _counts(cases)
    s["session_evidence"] = ev.get("", [])
    s["stage_name"] = next((f"{x['aspice']} {x['name']}" for x in cfg["stages"] if x["id"] == s["stage"]), s["stage"])
    s["states"] = SESSION_STATES
    return s


# ---------------------------------------------------------------- evidence
def add_evidence(repo, sid: str, tc: str, filename: str, data: bytes, note: str, user: dict) -> dict:
    s = _session(repo, sid)
    if s["status"] == "aborted":
        raise TestError("중단된 세션에는 증빙을 올릴 수 없습니다")
    if not data:
        raise TestError("빈 파일입니다")
    if len(data) > MAX_EVIDENCE:
        raise TestError(f"파일은 {MAX_EVIDENCE // (1024 * 1024)}MB 이하만 올릴 수 있습니다")
    if tc and not repo.con.execute("SELECT 1 FROM session_cases WHERE session=? AND tc=?", (sid, tc)).fetchone():
        raise KeyError(f"{sid} 에 {tc} 없음")
    name = Path(filename or "evidence.bin").name[:180] or "evidence.bin"
    sha = hashlib.sha256(data).hexdigest()
    ext = Path(name).suffix.lower()[:10]
    store = repo.dir / "evidence" / sha[:2]
    store.mkdir(parents=True, exist_ok=True)
    if not re.match(r"^\.[a-z0-9]+$", ext):
        ext = ""
    f = store / f"{sha}{ext}"
    if not f.exists():
        f.write_bytes(data)
    mime = mimetypes.guess_type(name)[0] or "application/octet-stream"
    cur = repo.con.execute("INSERT INTO evidence(session,tc,filename,path,mime,size,sha256,note,at,by) VALUES(?,?,?,?,?,?,?,?,?,?)",
                           (sid, tc or "", name, str(f.relative_to(repo.dir)), mime, len(data), sha, note or "", time.time(), user["username"]))
    repo.con.commit()
    return dict(repo.con.execute("SELECT id,session,tc,filename,mime,size,sha256,note,at,by FROM evidence WHERE id=?",
                                 (cur.lastrowid,)).fetchone())


def get_evidence(repo, eid: int) -> tuple[dict, bytes]:
    r = repo.con.execute("SELECT * FROM evidence WHERE id=?", (eid,)).fetchone()
    if not r:
        raise KeyError("증빙 없음")
    p = repo.dir / r["path"]
    return dict(r), p.read_bytes()


def delete_evidence(repo, eid: int, user: dict) -> dict:
    r = repo.con.execute("SELECT * FROM evidence WHERE id=?", (eid,)).fetchone()
    if not r:
        raise KeyError("증빙 없음")
    s = _session(repo, r["session"])
    if s["status"] == "completed":
        raise TestError("완료된 세션의 증빙은 지울 수 없습니다")
    if r["by"] != user["username"] and _role(user) < 2:
        raise PermissionError("올린 사람 또는 PM·관리자만 지울 수 있습니다")
    repo.con.execute("DELETE FROM evidence WHERE id=?", (eid,))
    if not repo.con.execute("SELECT 1 FROM evidence WHERE path=?", (r["path"],)).fetchone():
        try:
            (repo.dir / r["path"]).unlink()
        except OSError:
            pass
    repo.con.commit()
    return {"deleted": eid}


# ---------------------------------------------------------------- dashboard
def dashboard(repo, cfg) -> dict:
    items = {i["id"]: i for i in repo.items()}
    vs = _verify_stages(cfg)
    links = repo.links()
    verifies = defaultdict(list)       # target -> [tc]
    for l in links:
        if l["role"] == "verifies" and l["src"] in items and l["dst"] in items:
            verifies[l["dst"]].append(l["src"])

    def latest(tc):
        return (items[tc]["attrs"].get("result") or {}).get("status", "NOT_RUN")

    def stale(tc):
        res = items[tc]["attrs"].get("result") or {}
        return any(items.get(d) and v is not None and items[d]["version"] > v for d, v in (res.get("verified_versions") or {}).items())

    by_stage, coverage = [], []
    for sid, st in vs.items():
        tcs = [i for i in items.values() if i["type"] in st["types"]]
        c = Counter(latest(t["id"]) for t in tcs)
        by_stage.append({"stage": sid, "name": f"{st['aspice']} {st['name']}", "total": len(tcs),
                         **{k: c[k] for k in ("PASS", "FAIL", "BLOCKED", "NOT_RUN")},
                         "stale": sum(1 for t in tcs if stale(t["id"]))})
        tgt = st.get("verifies")
        if tgt:
            tstage = next(s for s in cfg["stages"] if s["id"] == tgt)
            targets = [i for i in items.values() if i["type"] in tstage["types"] and cfg["types"].get(i["type"], {}).get("prefix")]
            linked = passed = failed = 0
            for t in targets:
                tl = [x for x in verifies.get(t["id"], []) if items[x]["type"] in st["types"]]
                if tl:
                    linked += 1
                    rs = [latest(x) for x in tl]
                    if "FAIL" in rs:
                        failed += 1
                    elif all(r == "PASS" for r in rs) and not any(stale(x) for x in tl):
                        passed += 1
            coverage.append({"stage": sid, "target": tgt, "target_name": f"{tstage['aspice']} {tstage['name']}", "total": len(targets),
                             "linked": linked, "passed": passed, "failed": failed})
    # sessions + daily trend
    sess = sessions(repo, cfg)
    days = defaultdict(lambda: Counter())
    for r in repo.con.execute("SELECT status, executed_at FROM session_cases WHERE executed_at IS NOT NULL"):
        d = dt.datetime.fromtimestamp(r["executed_at"]).date().isoformat()
        days[d][r["status"]] += 1
    trend = []
    cum = Counter()
    for d in sorted(days):
        cum.update(days[d])
        trend.append({"date": d, "executed": sum(days[d].values()) - days[d]["NOT_RUN"], "pass": days[d]["PASS"],
                      "fail": days[d]["FAIL"], "cum_executed": sum(cum.values()) - cum["NOT_RUN"], "cum_pass": cum["PASS"]})
    # failures, by domain, flaky
    open_fail, dom = [], Counter()
    defects = {}
    for r in repo.con.execute("SELECT tc, defect, session, actual FROM session_cases WHERE status IN ('FAIL','BLOCKED') ORDER BY executed_at"):
        defects[r["tc"]] = (r["defect"], r["session"], r["actual"])
    for t in items.values():
        if not t["type"].endswith("_TC"):
            continue
        st = latest(t["id"])
        if st in ("FAIL", "BLOCKED"):
            res = t["attrs"].get("result") or {}
            d, s, note = defects.get(t["id"], ("", res.get("run_id", ""), res.get("note", "")))
            open_fail.append({"tc": t["id"], "title": t["title"], "status": st, "domain": t["domain"], "defect": d, "session": s,
                              "note": note,
                              "verifies": [l["dst"] for l in links if l["src"] == t["id"] and l["role"] == "verifies"]})
            if st == "FAIL":
                dom[t["domain"]] += 1
    flaky = []
    for t in items.values():
        hist = [h.get("status") for h in (t["attrs"].get("results") or [])][-6:]
        if "PASS" in hist and "FAIL" in hist:
            flips = sum(1 for a, b in zip(hist, hist[1:]) if a != b)
            if flips >= 2:
                flaky.append({"tc": t["id"], "title": t["title"], "history": hist})
    total = sum(s["total"] for s in by_stage)
    executed = sum(s["total"] - s["NOT_RUN"] for s in by_stage)
    passed = sum(s["PASS"] for s in by_stage)
    return {"by_stage": by_stage, "coverage": coverage, "sessions": sess[:30], "trend": trend[-90:],
            "open_failures": sorted(open_fail, key=lambda x: (x["status"] != "FAIL", x["tc"]))[:200],
            "fail_by_domain": dom.most_common(), "flaky": flaky[:50],
            "domain_names": cfg.get("domain_names", {}),
            "kpi": {"tcs": total, "executed": executed, "pass": passed,
                    "exec_rate": round(executed * 100 / total, 1) if total else 0,
                    "pass_rate": round(passed * 100 / executed, 1) if executed else None,
                    "fail": sum(s["FAIL"] for s in by_stage), "blocked": sum(s["BLOCKED"] for s in by_stage),
                    "stale": sum(s["stale"] for s in by_stage),
                    "active_sessions": sum(1 for s in sess if s["status"] in ("planned", "in_progress")),
                    "evidence": repo.con.execute("SELECT COUNT(*) FROM evidence").fetchone()[0]}}


# ---------------------------------------------------------------- report (15-52 verification results)
def export_xlsx(repo, cfg, sid: str, out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    s = session(repo, cfg, sid)
    wb = Workbook(); ws = wb.active; ws.title = "요약"
    thin = Side(style="thin", color="C9D1DB"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = f"시험 결과 보고서 (15-52 검증 결과) — {s['id']} {s['name']}"; ws["A1"].font = Font(bold=True, size=14)
    c = s["counts"]
    info = [("검증 단계", s["stage_name"]), ("빌드 / SW 버전", s["build"]), ("HW 샘플", s["sample"]), ("시험 환경", s["environment"]),
            ("담당", s["owner"]), ("계획 기간", f"{s['planned_start']} ~ {s['planned_end']}"),
            ("시작", time.strftime("%Y-%m-%d %H:%M", time.localtime(s["started"])) if s["started"] else "-"),
            ("종료", time.strftime("%Y-%m-%d %H:%M", time.localtime(s["closed"])) if s["closed"] else "-"),
            ("상태", SESSION_STATES.get(s["status"], s["status"])), ("판정", s["verdict"] or "-"),
            ("TC", c["total"]), ("PASS", c["PASS"]), ("FAIL", c["FAIL"]), ("BLOCKED", c["BLOCKED"]), ("미실행", c["NOT_RUN"]),
            ("해당 없음", c["N_A"]), ("합격률", f"{c['pass_rate']}%" if c["pass_rate"] is not None else "-"), ("설명", s["description"])]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(i, 1, k).font = Font(bold=True); ws.cell(i, 2, v)
    ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 60
    cs = wb.create_sheet("TC 결과")
    hdr = ["TC", "시험명", "검증 대상", "사전 조건", "기대 결과", "합격 기준", "결과", "실제 결과", "결함", "실행자", "실행 시각", "TC 버전", "증빙"]
    for j, h in enumerate(hdr, 1):
        x = cs.cell(1, j, h); x.font = Font(bold=True, color="FFFFFF"); x.fill = PatternFill("solid", fgColor="1F5FBF"); x.border = bd
    colr = {"PASS": "E3F4E7", "FAIL": "FDF3F2", "BLOCKED": "FFF3DC"}
    for i, r in enumerate(s["cases"], 2):
        vals = [r["tc"], r["title"], ", ".join(r["verifies"]), r["pre_condition"], r["expected"], r["pass_criteria"], r["status"],
                r["actual"], r["defect"], r["executed_by"],
                time.strftime("%Y-%m-%d %H:%M", time.localtime(r["executed_at"])) if r["executed_at"] else "", r["tc_version"] or "",
                "\n".join(f"{e['filename']} (sha256 {e['sha256'][:12]}…)" for e in r["evidence"])]
        for j, v in enumerate(vals, 1):
            x = cs.cell(i, j, v); x.border = bd; x.alignment = Alignment(wrap_text=True, vertical="top")
        if r["status"] in colr:
            cs.cell(i, 7).fill = PatternFill("solid", fgColor=colr[r["status"]])
    for j, w in enumerate([20, 30, 22, 26, 26, 26, 10, 34, 14, 12, 16, 8, 40], 1):
        cs.column_dimensions[chr(64 + j)].width = w
    cs.freeze_panes = "B2"; cs.auto_filter.ref = f"A1:M{len(s['cases']) + 1}"
    es = wb.create_sheet("증빙")
    for j, h in enumerate(["TC", "파일", "크기(B)", "SHA-256", "메모", "올린 사람", "시각"], 1):
        es.cell(1, j, h).font = Font(bold=True)
    allev = s["session_evidence"] + [e for r in s["cases"] for e in r["evidence"]]
    for i, e in enumerate(allev, 2):
        for j, v in enumerate([e["tc"] or "(세션)", e["filename"], e["size"], e["sha256"], e["note"], e["by"],
                               time.strftime("%Y-%m-%d %H:%M", time.localtime(e["at"]))], 1):
            es.cell(i, j, v)
    for j, w in enumerate([20, 36, 10, 66, 30, 12, 16], 1):
        es.column_dimensions[chr(64 + j)].width = w
    wb.save(out)
    return out
