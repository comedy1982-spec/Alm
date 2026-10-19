"""Revision change analysis: what changed between two revisions of a requirements document, and what it means.

Input is the import summary of a revision (the details of new / modified / deleted / restored items — the same
data the import preview shows). For every change a local rule pass works out *what kind* of change it is:

  numeric        a limit or value changed (120 ms → 100 ms), with the direction (stricter / looser) when the
                 sentence says which way the limit binds (이내·이하·최대 vs 이상·최소)
  obligation     mandatory ↔ recommended ("해야 한다" ↔ "바람직하다 / 할 수 있다")
  negation       a prohibition added or removed ("않아야", "금지")
  asil           the ASIL of the requirement changed (table column or inline)
  verification   verification method changed
  scope          substantive wording added / removed
  editorial      wording only (typo, punctuation, spacing) — same numbers, same obligation
  added/removed  new requirement / requirement withdrawn

and how far it reaches in the trace (items to re-review, test cases to re-run, ASIL items touched), its severity
and the follow-up it calls for. When an LLM is configured (alm.json "llm") the model adds a plain-language
explanation per change, an overall summary and questions worth asking the customer; without one the local
analysis is the report.
"""
from __future__ import annotations

import difflib
import json
import re
import time
from collections import Counter
from pathlib import Path

from .classify import quantities

SCHEMA = """
CREATE TABLE IF NOT EXISTS change_reports (
    import_id TEXT PRIMARY KEY, doc TEXT NOT NULL, rev TEXT NOT NULL, at REAL NOT NULL, by TEXT NOT NULL DEFAULT '',
    report TEXT NOT NULL
);
"""

MUST = re.compile(r"(야\s*(한다|하며|하고|함|합니다|하나)|\bshall\b|\bmust\b)", re.I)
MAY = re.compile(r"(바람직|권장|권고|할\s*수\s*있다|해도\s*된다|\bshould\b|\bmay\b)", re.I)
NEG = re.compile(r"(않아야|하지\s*않|금지|불가|해서는\s*안|shall\s+not|must\s+not)", re.I)
UPPER = re.compile(r"(이내|이하|미만|최대|넘지|초과하지|within|at\s+most|max)", re.I)
LOWER = re.compile(r"(이상|최소|초과|넘어야|at\s+least|min)", re.I)
KIND_LABEL = {"numeric": "수치 변경", "obligation": "의무 수준 변경", "negation": "금지 조건 변경", "asil": "ASIL 변경",
              "verification": "검증 방법 변경", "scope": "내용 변경", "editorial": "표현 수정", "added": "신규 요구사항",
              "removed": "삭제된 요구사항", "restored": "복원", "attribute": "속성 변경", "table": "표 변경", "figure": "그림 변경"}
SEV_LABEL = {"high": "높음", "medium": "보통", "low": "낮음"}


def init(repo):
    repo.con.executescript(SCHEMA)
    repo.con.commit()


def _fmt(q):
    v = q["value"]
    unit = q["unit"]
    m = re.match(r"^([\d.]+)", q["raw"])
    return f"{m.group(1) if m else v} {unit}".strip()


def _numeric(prev: str, new: str) -> list[dict]:
    a, b = quantities(prev or ""), quantities(new or "")
    out = []
    by_dim_a, by_dim_b = {}, {}
    for q in a:
        by_dim_a.setdefault(q["dim"], []).append(q)
    for q in b:
        by_dim_b.setdefault(q["dim"], []).append(q)
    for dim in sorted(set(by_dim_a) | set(by_dim_b)):
        xa, xb = by_dim_a.get(dim, []), by_dim_b.get(dim, [])
        for i in range(max(len(xa), len(xb))):
            qa = xa[i] if i < len(xa) else None
            qb = xb[i] if i < len(xb) else None
            if qa and qb and abs(qa["value"] - qb["value"]) < 1e-12:
                continue
            d = {"dim": dim, "from": _fmt(qa) if qa else "", "to": _fmt(qb) if qb else ""}
            if qa and qb and qa["value"]:
                d["delta_pct"] = round((qb["value"] - qa["value"]) * 100 / abs(qa["value"]), 1)
                pos = (new or "").find(qb["raw"])
                ctx = (new or "")[max(0, pos - 4):pos + len(qb["raw"]) + 12]
                after = (new or "")[pos + len(qb["raw"]):pos + len(qb["raw"]) + 24]
                if re.search(r"(하면|되면|않으면|않거나|경우|때|시\s)", after):
                    d["direction"] = "threshold"      # a trigger condition, not a limit the product must meet
                elif UPPER.search(ctx):
                    d["direction"] = "stricter" if qb["value"] < qa["value"] else "looser"
                elif LOWER.search(ctx):
                    d["direction"] = "stricter" if qb["value"] > qa["value"] else "looser"
                else:
                    d["direction"] = "changed"
            else:
                d["direction"] = "added" if qb else "removed"
            out.append(d)
    return out


def _words(s):
    return re.findall(r"[\w.%°µμ/+-]+", s or "")


def _token_diff(prev: str, new: str) -> tuple[list[str], list[str], float]:
    a, b = _words(prev), _words(new)
    sm = difflib.SequenceMatcher(a=a, b=b, autojunk=False)
    added, removed = [], []
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op in ("replace", "delete"):
            removed.append(" ".join(a[i1:i2]))
        if op in ("replace", "insert"):
            added.append(" ".join(b[j1:j2]))
    return added, removed, sm.ratio()


def _asil(attrs: dict) -> str:
    a = {str(k).replace(" ", "").lower(): v for k, v in (attrs or {}).items()}
    v = a.get("asil", "")
    return str(v).replace("ASIL", "").strip()


def classify(d: dict) -> dict:
    ch = d["change"]
    prev, new = d.get("prev_text", "") or "", d.get("text", "") or ""
    kinds, nums, notes = [], [], []
    if d.get("type") in ("TBL", "FIG"):
        # requirement rows inside tables are analysed as their own items; the table itself is context only
        added, removed, ratio = _token_diff(prev, new)
        d["_added"], d["_removed"] = added, removed
        return {"kinds": ["table" if d["type"] == "TBL" else "figure"], "numbers": [], "notes": []}
    if ch == "new":
        kinds.append("added")
        nums = [{"dim": q["dim"], "from": "", "to": _fmt(q), "direction": "added"} for q in quantities(new)]
    elif ch == "deleted" and d.get("still_in_doc"):
        # the ID is still in the document but the sentence is no longer a requirement (e.g. shall -> should)
        new = d["still_in_doc"]
        d["text"] = new
        kinds.append("obligation")
        notes.append("필수 요구사항 → 권고·정보 (의무 표현 삭제)" if MAY.search(new) or not MUST.search(new) else "요구사항 형식 변경")
        nums = _numeric(prev, new)
        if nums:
            kinds.append("numeric")
    elif ch == "deleted":
        kinds.append("removed")
    elif ch == "restored":
        kinds.append("restored")
    else:
        nums = _numeric(prev, new)
        if nums:
            kinds.append("numeric")
        mp, mn = bool(MUST.search(prev)), bool(MUST.search(new))
        yp, yn = bool(MAY.search(prev)), bool(MAY.search(new))
        if mp != mn or yp != yn:
            kinds.append("obligation")
            notes.append("필수 → 권고" if mp and not mn else "권고 → 필수" if mn and not mp else "의무 표현 변경")
        if bool(NEG.search(prev)) != bool(NEG.search(new)):
            kinds.append("negation")
            notes.append("금지 조건 추가" if NEG.search(new) else "금지 조건 삭제")
        ab, aa = _asil(d.get("attrs_before")), _asil(d.get("attrs_after"))
        if (ab or aa) and ab != aa:
            kinds.append("asil")
            notes.append(f"ASIL {ab or '—'} → {aa or '—'}")
        vb = (d.get("attrs_before") or {}).get("verification_method") or next((v for k, v in (d.get("attrs_before") or {}).items() if "검증" in k), "")
        va = (d.get("attrs_after") or {}).get("verification_method") or next((v for k, v in (d.get("attrs_after") or {}).items() if "검증" in k), "")
        if vb != va and (vb or va):
            kinds.append("verification")
            notes.append(f"검증 방법 {vb or '—'} → {va or '—'}")
        other = {k: (v, (d.get("attrs_after") or {}).get(k)) for k, v in (d.get("attrs_before") or {}).items()
                 if (d.get("attrs_after") or {}).get(k) != v and k.lower() != "asil" and "검증" not in k and k != "verification_method"}
        if other:
            kinds.append("attribute")
            notes += [f"{k}: {a or '—'} → {b or '—'}" for k, (a, b) in list(other.items())[:3]]
        added, removed, ratio = _token_diff(prev, new)
        d["_added"], d["_removed"], d["_ratio"] = added, removed, ratio
        if "numeric" in kinds and sum(len(x) for x in added if not re.fullmatch(r"[\d.,%°µμ]+\w{0,3}", x)) > 8:
            kinds.append("scope")
        if prev != new and not ({"numeric", "obligation", "negation"} & set(kinds)):
            if ratio >= 0.9 and sum(len(x) for x in added + removed) <= 12:
                kinds.append("editorial")
            else:
                kinds.append("scope")
        elif prev != new and ratio < 0.75:
            kinds.append("scope")
    return {"kinds": kinds or ["scope"], "numbers": nums, "notes": notes}


def _impact(repo, cfg, iid: str) -> dict:
    from . import sup
    try:
        imp = sup.impact(repo, cfg, [iid])
    except Exception:
        return {"review": [], "retest": [], "asil": [], "by_stage": {}}
    return {"review": imp["review"], "retest": imp["retest"], "asil": [x for x in imp["safety_items"] if x != iid],
            "by_stage": {k: v for k, v in imp["by_stage"].items()}}


def _severity(kinds: set, item_asil: str, imp: dict, d: dict) -> str:
    safety = (item_asil and item_asil not in ("QM", "")) or bool(imp["asil"])
    if "asil" in kinds or "negation" in kinds or "obligation" in kinds:
        return "high"
    if "removed" in kinds and (d.get("dependents") or imp["review"] or imp["retest"]):
        return "high"
    nums = d.get("_nums") or []
    if "numeric" in kinds and (any(n.get("direction") == "stricter" for n in nums) or
                               (safety and any(n.get("direction") in ("threshold", "changed") for n in nums))):
        return "high"
    if kinds & {"table", "figure"}:
        return "low"
    if kinds & {"numeric", "scope", "added", "removed", "verification", "attribute", "restored", "table", "figure"}:
        return "medium"
    return "low"


def _actions(kinds: set, sev: str, d: dict, imp: dict) -> list[str]:
    out = []
    if "added" in kinds:
        out.append("SYS.2 로 전개해 시스템 요구사항·검증 기준을 만드세요")
    if "removed" in kinds:
        out.append("하위 요구사항·TC 를 폐기 또는 재할당할지 결정하세요" if (d.get("dependents") or imp["review"]) else "하위 연결 없음 — 삭제 확인만 하면 됩니다")
    if imp["review"]:
        out.append(f"하위 항목 {len(imp['review'])}건 재검토" + ("" if "removed" in kinds else " (Suspect 표시됨)"))
    if imp["retest"]:
        out.append(f"관련 TC {len(imp['retest'])}건 폐기·수정 검토" if "removed" in kinds else
                   f"TC {len(imp['retest'])}건 재시험 — 테스트 세션 조건 '대상이 바뀐 TC'")
    if "numeric" in kinds and any(n.get("direction") == "stricter" for n in d.get("_nums", [])):
        out.append("강화된 한계값을 현재 설계·실측 데이터로 만족하는지 확인하세요")
    if "asil" in kinds:
        out.append("안전 분석(HARA·FMEDA·안전 요구) 갱신과 안전 담당 검토가 필요합니다")
    if "obligation" in kinds:
        out.append("의무 수준 변경 의도를 고객에게 확인하고 준수 여부 회신을 갱신하세요")
    if sev == "high":
        out.append("변경 요청(SUP.10)으로 영향 분석·CCB 결정을 남기세요")
    if "editorial" in kinds and len(kinds) == 1:
        out.append("표현만 바뀜 — 승인 상태로 되돌리기 전 확인만 하면 됩니다")
    return out


def _explain(d: dict, c: dict) -> str:
    return " · ".join(_explain_parts(d, c, True))


def _explain_parts(d: dict, c: dict, with_numbers: bool) -> list[str]:
    parts = []
    for n in (c["numbers"] if with_numbers else [n for n in c["numbers"] if not (n["from"] and n["to"])]):
        if n["from"] and n["to"]:
            dirw = {"stricter": "강화", "looser": "완화", "changed": "변경", "threshold": "판정 임계값"}.get(n.get("direction"), "")
            parts.append(f"{n['from']} → {n['to']}" + (f" ({n['delta_pct']:+g}%, {dirw})" if "delta_pct" in n else ""))
        elif n["to"]:
            parts.append(f"수치 추가 {n['to']}")
        elif n["from"]:
            parts.append(f"수치 삭제 {n['from']}")
    parts += c["notes"]
    if ("table" in c["kinds"] or "figure" in c["kinds"]) and d.get("_added"):
        parts.append("표 추가 내용: " + " ".join(x for x in d["_added"] if x)[:160])
    if "scope" in c["kinds"] and d.get("_added"):
        parts.append("추가: " + ", ".join(x for x in d["_added"] if x)[:80])
    if "scope" in c["kinds"] and d.get("_removed"):
        parts.append("삭제: " + ", ".join(x for x in d["_removed"] if x)[:80])
    if "editorial" in c["kinds"]:
        parts.append("표현만 수정")
    if "added" in c["kinds"]:
        parts.append("새 요구사항")
    if "removed" in c["kinds"]:
        parts.append("요구사항 삭제" + (f" — 하위 {len(d.get('dependents') or [])}건이 근거를 잃음" if d.get("dependents") else ""))
    return parts


def analyze_local(repo, cfg, summary: dict) -> dict:
    items = {i["id"]: i for i in repo.items(include_deleted=True)}
    rows = []
    for d in summary.get("details", []):
        if d.get("type") in ("HDG", "INFO") and d["change"] != "modified":
            continue
        c = classify(d)
        d["_nums"] = c["numbers"]
        kinds = set(c["kinds"])
        exists = d["id"] in items and d["change"] != "new"
        imp = _impact(repo, cfg, d["id"]) if exists else {"review": [], "retest": [], "asil": [], "by_stage": {}}
        it = items.get(d["id"]) or {}
        asil = _asil(d.get("attrs_after") or d.get("attrs_before") or {}) or str((it.get("attrs") or {}).get("asil") or "")
        sev = _severity(kinds, asil, imp, d)
        if d.get("type") == "INFO" and sev == "medium":
            sev = "low"
        rows.append({"id": d["id"], "native_id": d.get("native_id") or (it.get("attrs") or {}).get("native_id", ""),
                     "change": d["change"], "type": d.get("type", ""), "section": d.get("section", ""),
                     "prev_text": d.get("prev_text", ""), "text": d.get("text", ""), "kinds": c["kinds"],
                     "kind_labels": [KIND_LABEL.get(k, k) for k in c["kinds"]], "numbers": c["numbers"], "notes": c["notes"],
                     "asil": asil, "severity": sev, "impact": {"review": len(imp["review"]), "retest": len(imp["retest"]),
                                                               "asil_items": imp["asil"][:20], "by_stage": imp["by_stage"],
                                                               "review_ids": imp["review"][:50], "retest_ids": imp["retest"][:50]},
                     "dependents": d.get("dependents") or [], "was_approved": d.get("was_approved", False),
                     "explain": _explain(d, c), "detail": " · ".join(_explain_parts(d, c, False)),
                     "actions": _actions(kinds, sev, d, imp), "ai": ""})
    order = {"high": 0, "medium": 1, "low": 2}
    rows.sort(key=lambda r: (order[r["severity"]], r["section"], r["id"]))
    kc = Counter(k for r in rows for k in r["kinds"])
    sc = Counter(r["severity"] for r in rows)
    retest = sorted({t for r in rows for t in r["impact"]["retest_ids"]})
    review = sorted({t for r in rows for t in r["impact"]["review_ids"]})
    stricter = [r for r in rows if any(n.get("direction") == "stricter" for n in r["numbers"])]
    lines = [f"{summary.get('doc', '')} rev {summary.get('rev', '')}: 신규 {len(summary.get('new', []))} · 변경 {len(summary.get('modified', []))} · "
             f"삭제 {len(summary.get('deleted', []))}건 (영향도 높음 {sc['high']}, 보통 {sc['medium']}, 낮음 {sc['low']})"]
    if stricter:
        lines.append(f"한계값이 강화된 요구사항 {len(stricter)}건: " + ", ".join((r["native_id"] or r["id"]) for r in stricter[:6]))
    if kc["asil"]:
        lines.append(f"ASIL 이 바뀐 요구사항 {kc['asil']}건 — 안전 분석 갱신 필요")
    if kc["obligation"]:
        lines.append(f"필수/권고 수준이 바뀐 요구사항 {kc['obligation']}건 — 고객 의도 확인 권장")
    if kc["removed"]:
        lines.append(f"삭제된 요구사항 {kc['removed']}건 — 하위 산출물 정리 필요")
    if kc["added"]:
        lines.append(f"신규 요구사항 {kc['added']}건 — SYS.2 전개 필요")
    if kc["editorial"]:
        lines.append(f"표현만 바뀐 항목 {kc['editorial']}건 (영향 낮음)")
    lines.append(f"영향: 하위 재검토 {len(review)}건 · 재시험 TC {len(retest)}건")
    return {"doc": summary.get("doc", ""), "rev": summary.get("rev", ""), "import_id": summary.get("import_id", ""),
            "dry_run": bool(summary.get("dry_run")), "at": time.time(), "mode": "local", "summary": lines, "questions": [],
            "counts": {"severity": dict(sc), "kinds": dict(kc), "review": len(review), "retest": len(retest),
                       "new": len(summary.get("new", [])), "modified": len(summary.get("modified", [])),
                       "deleted": len(summary.get("deleted", []))},
            "retest_ids": retest, "review_ids": review, "rows": rows, "kind_labels": KIND_LABEL, "sev_labels": SEV_LABEL}


SYSTEM = """너는 자동차 ECU 요구사항 엔지니어다. 고객 요구사양서의 개정 전/후 차이를 검토해 개발팀에 알려 준다.
입력: 변경 목록(JSON). 각 항목에 id, 원문 ID, 이전 문장, 새 문장, 규칙 분석(kinds, numbers, notes), 영향(재검토·재시험 수, ASIL)이 있다.
반드시 아래 JSON 만 출력한다 (코드블록·설명 금지):
{"summary": ["개정 전체 요약 3~6줄, 가장 중요한 변경부터"],
 "items": {"<id>": {"explain": "무엇이 어떻게 바뀌었고 개발에 어떤 의미인지 한두 문장", "risk": "high|medium|low", "question": "고객에게 확인할 점(없으면 빈 문자열)"}},
 "questions": ["고객에게 보낼 확인 질의 (없으면 빈 배열)"]}
규칙: 입력에 없는 사실을 지어내지 않는다. 수치는 입력 그대로 쓴다. risk 는 규칙 분석보다 낮추지 말고, 근거가 있을 때만 높인다."""


def analyze(repo, cfg, summary: dict, use_ai: bool = True) -> dict:
    return enrich(cfg, analyze_local(repo, cfg, summary), use_ai)


def enrich(cfg, rep: dict, use_ai: bool = True) -> dict:
    """Add the LLM explanation to a local report. Needs no database access (run it outside the write lock)."""
    from .assistant import status
    st = status(cfg)
    if not use_ai or st["mode"] != "llm" or not rep["rows"]:
        rep["ai_status"] = "AI 미연결 — 규칙 기반 분석" if st["mode"] != "llm" else ""
        return rep
    from .llm import chat_call
    pri = [r for r in rep["rows"] if r["severity"] != "low"][:60] + [r for r in rep["rows"] if r["severity"] == "low"][:20]
    payload = [{"id": r["id"], "native_id": r["native_id"], "change": r["change"], "section": r["section"][:60],
                "before": (r["prev_text"] or "")[:400], "after": (r["text"] or "")[:400], "kinds": r["kinds"],
                "numbers": r["numbers"], "notes": r["notes"], "asil": r["asil"],
                "impact": {"review": r["impact"]["review"], "retest": r["impact"]["retest"], "asil_items": len(r["impact"]["asil_items"])}}
               for r in pri]
    try:
        raw = chat_call(cfg, SYSTEM, [{"role": "user", "content": json.dumps({"doc": rep["doc"], "rev": rep["rev"], "changes": payload},
                                                                                ensure_ascii=False)}], max_tokens=4096, timeout=180)
        m = re.search(r"\{.*\}", raw, re.S)
        j = json.loads(m.group(0)) if m else {}
        order = {"high": 0, "medium": 1, "low": 2}
        for r in rep["rows"]:
            x = (j.get("items") or {}).get(r["id"]) or {}
            if x.get("explain"):
                r["ai"] = str(x["explain"])[:600]
            if x.get("question"):
                r["question"] = str(x["question"])[:300]
            if x.get("risk") in order and order[x["risk"]] < order[r["severity"]]:
                r["severity"] = x["risk"]; r["raised_by_ai"] = True
        if j.get("summary"):
            rep["summary"] = [str(s)[:300] for s in j["summary"]][:8] + rep["summary"][-1:]
        rep["questions"] = [str(q)[:300] for q in (j.get("questions") or [])][:10]
        rep["mode"] = "llm"; rep["model"] = st["model"]; rep["ai_status"] = ""
        rep["rows"].sort(key=lambda r: (order[r["severity"]], r["section"], r["id"]))
        rep["counts"]["severity"] = dict(Counter(r["severity"] for r in rep["rows"]))
    except Exception as e:     # keep the local analysis, say why the AI part is missing
        rep["ai_status"] = f"AI 호출 실패 — 규칙 기반 분석만 표시 ({e})"[:200]
    return rep


def save(repo, rep: dict, user: str) -> dict:
    init(repo)
    if not rep.get("import_id") or rep.get("dry_run"):
        raise ValueError("저장은 실제로 가져온 개정에만 할 수 있습니다")
    repo.con.execute("INSERT INTO change_reports VALUES(?,?,?,?,?,?) ON CONFLICT(import_id) DO UPDATE SET report=excluded.report, at=excluded.at",
                     (rep["import_id"], rep["doc"], rep["rev"], time.time(), user, json.dumps(rep, ensure_ascii=False)))
    repo.con.commit()
    return {"saved": rep["import_id"]}


def reports(repo) -> list[dict]:
    init(repo)
    out = []
    for r in repo.con.execute("SELECT import_id, doc, rev, at, by, report FROM change_reports ORDER BY at DESC"):
        rep = json.loads(r["report"])
        out.append({"import_id": r["import_id"], "doc": r["doc"], "rev": r["rev"], "at": r["at"], "by": r["by"],
                    "mode": rep.get("mode"), "counts": rep.get("counts", {}), "headline": (rep.get("summary") or [""])[0]})
    return out


def report(repo, import_id: str) -> dict:
    init(repo)
    r = repo.con.execute("SELECT report FROM change_reports WHERE import_id=?", (import_id,)).fetchone()
    if not r:
        raise KeyError(f"변경 분석 {import_id} 없음")
    return json.loads(r["report"])


def to_cr(repo, cfg, import_id: str, ids: list[str], user: dict, title: str = "") -> dict:
    """Open a SUP.10 change request for the selected changes, with the impact analysis already saved."""
    from . import sup
    rep = report(repo, import_id)
    rows = {r["id"]: r for r in rep["rows"]}
    pick = [i for i in ids if i in rows and repo.get(i)]
    if not pick:
        raise ValueError("변경 요청에 넣을 항목을 고르세요 (기존 항목만 가능)")
    desc = "\n".join(f"- {rows[i]['native_id'] or i}: {rows[i]['ai'] or rows[i]['explain']}" for i in pick)
    cr = sup.create(repo, cfg, "CR", {
        "title": title or f"{rep['doc']} rev {rep['rev']} 고객 요구 변경 반영 ({len(pick)}건)",
        "items": [{"id": i, "role": "scope"} for i in pick],
        "fields": {"description": f"고객 요구사양서 {rep['doc']} rev {rep['rev']} 변경 반영\n{desc}"[:4000], "reason": "고객 요구 변경",
                   "priority": "높음" if any(rows[i]["severity"] == "high" for i in pick) else "보통",
                   "target_release": ""}}, user)
    try:
        cr = sup.save_impact(repo, cfg, cr["id"], user)
    except ValueError:
        pass
    return cr


def export_xlsx(rep: dict, out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "변경 분석"
    ws["A1"] = f"요구사양서 변경 분석 — {rep['doc']} rev {rep['rev']} ({'AI ' + rep.get('model', '') if rep.get('mode') == 'llm' else '규칙 기반'})"
    ws["A1"].font = Font(bold=True, size=13)
    r0 = 3
    for i, line in enumerate(rep["summary"], r0):
        ws.cell(i, 1, "• " + line)
    r0 += len(rep["summary"]) + 1
    if rep.get("questions"):
        ws.cell(r0, 1, "고객 확인 질의").font = Font(bold=True)
        for i, q in enumerate(rep["questions"], r0 + 1):
            ws.cell(i, 1, f"Q{i - r0}. {q}")
        r0 += len(rep["questions"]) + 2
    hdr = ["영향도", "ID", "원문 ID", "구분", "변경 유형", "이전", "변경 후", "분석", "AI 설명", "재검토", "재시험", "ASIL", "권장 조치"]
    fill = {"high": "FDF3F2", "medium": "FFF3DC", "low": "EEF1F5"}
    for j, h in enumerate(hdr, 1):
        x = ws.cell(r0, j, h); x.font = Font(bold=True, color="FFFFFF"); x.fill = PatternFill("solid", fgColor="1F5FBF")
    for i, r in enumerate(rep["rows"], r0 + 1):
        vals = [SEV_LABEL[r["severity"]], r["id"], r["native_id"], {"new": "신규", "modified": "변경", "deleted": "삭제", "restored": "복원"}.get(r["change"], r["change"]),
                ", ".join(r["kind_labels"]), r["prev_text"], r["text"], r["explain"], r.get("ai", ""), r["impact"]["review"], r["impact"]["retest"],
                r["asil"], "\n".join(r["actions"])]
        for j, v in enumerate(vals, 1):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=fill[r["severity"]])
    for j, w in enumerate([8, 18, 14, 7, 18, 46, 46, 40, 46, 8, 8, 6, 40], 1):
        ws.column_dimensions[chr(64 + j)].width = w
    wb.save(out)
    return out
