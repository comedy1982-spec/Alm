"""Exports: Excel (items, links, trace, every table as its own sheet), Polarion-style JSON,
Markdown per document (docs-as-code), and a static read-only HTML snapshot of the web UI."""
from __future__ import annotations

import base64
import json
import time
import re
from collections import defaultdict
from pathlib import Path

from .db import Repo

WEB = Path(__file__).parent / "web"


def _links_by(repo: Repo):
    up, down = defaultdict(list), defaultdict(list)
    for l in repo.links():
        up[l["src"]].append(l)
        down[l["dst"]].append(l)
    return up, down


def export_excel(repo: Repo, cfg: dict, out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    up, down = _links_by(repo)
    head = ["ID", "Type", "ASPICE", "Polarion Type", "Domain", "Status", "Ver", "Document", "Rev", "Section",
            "Text", "Native ID", "Fields", "Lint", "Upstream", "Downstream", "Source"]
    ws.append(head)
    for it in repo.items(include_deleted=True):
        if it["type"] in ("HDG",):
            continue
        t = cfg["types"].get(it["type"], {})
        ws.append([it["id"], it["type"], t.get("aspice", ""), t.get("polarion", ""), it["domain"], it["status"],
                   it["version"], it["doc"], it["doc_rev"], it["section"], it["text"][:32000],
                   it["attrs"].get("native_id", ""),
                   json.dumps(it["attrs"].get("fields", {}), ensure_ascii=False) if it["attrs"].get("fields") else "",
                   "\n".join(f"[{f['sev']}] {f['msg']}" for f in it["lint"]),
                   "\n".join(f"{l['role']} {l['dst']}{' (suspect)' if l['suspect'] else ''}" for l in up[it["id"]]),
                   "\n".join(f"{l['src']} {l['role']}" for l in down[it["id"]]),
                   it["source_loc"]])
    widths = [16, 8, 7, 18, 7, 9, 5, 18, 14, 28, 80, 14, 30, 40, 28, 28, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    wl = wb.create_sheet("Links")
    wl.append(["Source", "Role", "Target", "State", "Suspect", "Note"])
    for l in repo.links():
        wl.append([l["src"], l["role"], l["dst"], l["state"], "Y" if l["suspect"] else "", l["note"]])

    # every extracted table as its own sheet, merged cells preserved
    used = set()
    for it in repo.items(type_="TBL"):
        grid = it["attrs"].get("grid") or []
        if not grid:
            continue
        name = re.sub(r"[\[\]\*\?/\\:]", "", it["id"])[:31]
        if name in used:
            continue
        used.add(name)
        tws = wb.create_sheet(name)
        tws.append([f"{it['id']}  {it['attrs'].get('caption', '')}  ({it['doc']} {it['source_loc']})"])
        tws["A1"].font = Font(bold=True)
        for r_i, row in enumerate(grid):
            for c_i, cell in enumerate(row):
                if cell.get("skip"):
                    continue
                tws.cell(row=r_i + 2, column=c_i + 1, value=cell.get("t", ""))
                rs, cs = cell.get("rs", 1), cell.get("cs", 1)
                if rs > 1 or cs > 1:
                    tws.merge_cells(start_row=r_i + 2, start_column=c_i + 1, end_row=r_i + 1 + rs, end_column=c_i + cs)
                if r_i == 0:
                    tws.cell(row=2, column=c_i + 1).font = Font(bold=True)
        for c_i in range(len(grid[0])):
            tws.column_dimensions[get_column_letter(c_i + 1)].width = 22
    out = Path(out)
    wb.save(out)
    return out


def export_json(repo: Repo, cfg: dict, out: Path) -> Path:
    """Polarion-style work item JSON (type names and link roles from the strategy schema)."""
    up, _ = _links_by(repo)
    wis = []
    for it in repo.items():
        t = cfg["types"].get(it["type"], {})
        wis.append({
            "id": it["id"], "type": t.get("polarion", it["type"]), "title": it["title"],
            "description": it["text"], "status": it["status"], "document": it["doc"],
            "customFields": {"aspiceProcess": t.get("aspice", ""), "domain": it["domain"],
                             "nativeId": it["attrs"].get("native_id", ""), "section": it["section"],
                             "sourceLocation": it["source_loc"], "revision": it["doc_rev"],
                             **{f"field_{k}": v for k, v in it["attrs"].get("fields", {}).items()}},
            "linkedWorkItems": [{"role": l["role"], "target": l["dst"], "suspect": bool(l["suspect"])}
                                for l in up[it["id"]] if l["state"] == "confirmed"],
        })
    Path(out).write_text(json.dumps({"project": cfg["project"], "workItems": wis}, ensure_ascii=False, indent=2),
                         encoding="utf-8")
    return Path(out)


def _md_table(grid) -> str:
    if not grid:
        return ""
    rows = [[c.get("t", "").replace("\n", "<br>").replace("|", "\\|") for c in r] for r in grid]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def export_markdown(repo: Repo, cfg: dict, outdir: Path) -> list[Path]:
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    up, _ = _links_by(repo)
    paths = []
    for d in repo.docs():
        lines = [f"# {d['doc']}  (rev {d['rev']})", ""]
        for it in repo.items(doc=d["doc"]):
            typ = it["type"]
            if typ == "HDG":
                lvl = min(6, int(it["attrs"].get("level", 1)) + 1)
                lines += [f"{'#' * lvl} {it['text']}", ""]
            elif typ == "INFO":
                lines += [it["text"], ""]
            elif typ == "TBL":
                cap = it["attrs"].get("caption") or ""
                lines += [f"<a id=\"{it['id']}\"></a>**{it['id']}** {cap}", "", _md_table(it["attrs"].get("grid")), ""]
            elif typ == "FIG":
                f = it["attrs"].get("file")
                cap = it["attrs"].get("caption") or it["id"]
                if f:
                    lines += [f"![{cap}](../.alm/assets/{f})", f"*{it['id']} — {cap}*", ""]
                if it["attrs"].get("diagram_code"):
                    lines += ["```mermaid", it["attrs"]["diagram_code"], "```", ""]
            else:
                links = ", ".join(f"{l['role']} {l['dst']}" for l in up[it["id"]] if l["state"] == "confirmed")
                lines += [f"- **[{it['id']}]** `{typ}` `{it['status']}` {it['text']}" + (f"  \n  ↳ {links}" if links else ""), ""]
        p = outdir / (re.sub(r"[^\w\-가-힣]+", "_", d["doc"]) + ".md")
        p.write_text("\n".join(lines), encoding="utf-8")
        paths.append(p)
    return paths


def state_payload(repo: Repo, cfg: dict) -> dict:
    return {
        "project": cfg["project"],
        "types": cfg["types"],
        "stages": cfg["stages"],
        "stage_counts": __import__("ecualm.stages", fromlist=["stage_counts"]).stage_counts(repo, cfg),
        "domain_names": cfg["domain_names"],
        "link_rules": cfg["link_rules"],
        "signal_classes": cfg.get("signal_classes", {}),
        "docs": repo.docs(),
        "items": repo.items(include_deleted=True),
        "links": repo.links(),
        "imports": repo.imports(),
        "baselines": repo.baselines(),
    }


def export_static_html(repo: Repo, cfg: dict, out: Path) -> Path:
    from .audit import run_audit
    state = state_payload(repo, cfg)
    assets = {}
    for it in state["items"]:
        f = it["attrs"].get("file")
        if f and f not in assets and (repo.assets / f).exists():
            ext = f.rsplit(".", 1)[-1].replace("jpg", "jpeg").replace("svg", "svg+xml")
            assets[f] = f"data:image/{ext};base64," + base64.b64encode((repo.assets / f).read_bytes()).decode()
    for it in state["items"]:
        for h in [it]:
            h["history"] = repo.history(it["id"])
    from .stages import stage_doc
    stagedocs = {s["id"]: {"pam": stage_doc(repo, cfg, s["id"], "pam"), "source": stage_doc(repo, cfg, s["id"], "source")}
                 for s in cfg["stages"]}
    payload = {"state": state, "audit": run_audit(repo, cfg), "assets": assets, "stagedocs": stagedocs}
    html = (WEB / "index.html").read_text(encoding="utf-8")
    inject = "<script>window.__STATIC__=" + json.dumps(payload, ensure_ascii=False).replace("</", "<\\/") + ";</script>"
    html = html.replace("<!--STATIC-->", inject)
    Path(out).write_text(html, encoding="utf-8")
    return Path(out)


def export_stage_docx(repo: Repo, cfg: dict, sid: str, out: Path, mode: str = "pam") -> Path:
    """PAM 4.0 structured stage document (SRS / SWRS / ...) as a Word file."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from .stages import stage_doc

    sd = stage_doc(repo, cfg, sid, mode)
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Malgun Gothic"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    st.font.size = Pt(10)
    d.add_paragraph(f"{sd['title']}", style="Title")
    d.add_paragraph(f"프로젝트 {cfg['project']} · ASPICE PAM 4.0 · 생성 {time.strftime('%Y-%m-%d %H:%M')} · 원천 문서: {', '.join(sd['docs']) or '-'}")

    def item_block(it):
        p = d.add_paragraph()
        r = p.add_run(f"{it['id']}  ")
        r.bold = True
        r.font.color.rgb = RGBColor(0x1F, 0x5F, 0xBF)
        p.add_run(f"[{it['status']}]  ").italic = True
        p.add_run(it["text"])
        a = it["attrs"]
        meta = []
        if a.get("native_id"):
            meta.append(f"원문 ID {a['native_id']}")
        if a.get("asil") or (a.get("fields") or {}).get("ASIL"):
            meta.append(f"ASIL {a.get('asil') or a['fields']['ASIL']}")
        if a.get("task_cycle"):
            meta.append(f"Task_Cycle {a['task_cycle']}")
        if a.get("target_layer"):
            meta.append(f"Target_Layer {a['target_layer']}")
        if a.get("discipline"):
            meta.append(f"Discipline {a['discipline']}")
        ups = ", ".join(f"{l['role']} {l['dst']}" for l in it.get("_up", []))
        if ups:
            meta.append(f"추적 {ups}")
        if meta:
            q = d.add_paragraph("   " + " · ".join(meta))
            q.runs[0].font.size = Pt(8.5)
        if a.get("spec"):
            sp = a["spec"]
            t = d.add_table(rows=0, cols=2)
            t.style = "Table Grid"
            for k, lab in (("pre_condition", "Pre-Condition"), ("trigger", "Trigger / Input"), ("timing", "Processing / Timing"), ("output", "Output / Post-Condition")):
                row = t.add_row().cells
                row[0].text = lab
                row[1].text = sp.get(k) or "-"
        if a.get("verification_criteria"):
            v = d.add_paragraph("   검증 기준: " + a["verification_criteria"])
            v.runs[0].font.size = Pt(8.5)

    for sec in sd["sections"]:
        d.add_heading(f"{sec['number']} {sec['title']}".strip(), 1)
        if sec["kind"] == "text":
            for t in sec["text"]:
                d.add_paragraph(t)
        elif sec["kind"] == "group":
            for sub in sec["sub"]:
                d.add_heading(f"{sub['number']} {sub['title']}", 2)
                for it in sub["items"]:
                    item_block(it)
        elif sec["kind"] == "source":
            for it in sec["items"]:
                if it["type"] == "HDG":
                    d.add_heading(it["text"], min(3, int(it["attrs"].get("level", 1)) + 1))
                elif it["type"] == "INFO":
                    d.add_paragraph(it["text"])
                elif it["type"] == "TBL":
                    grid = it["attrs"].get("grid") or []
                    if grid:
                        t = d.add_table(rows=len(grid), cols=len(grid[0]))
                        t.style = "Table Grid"
                        for ri, row in enumerate(grid):
                            for ci, c in enumerate(row):
                                if not c.get("skip"):
                                    t.cell(ri, ci).text = c.get("t", "")
                elif it["type"] == "FIG":
                    f = it["attrs"].get("file")
                    if f and (repo.assets / f).exists() and f.rsplit(".", 1)[-1] in ("png", "jpg", "jpeg", "gif"):
                        from docx.shared import Cm
                        d.add_picture(str(repo.assets / f), width=Cm(15))
                    d.add_paragraph(f"{it['id']} {it['attrs'].get('caption', '')}")
                else:
                    item_block(it)
        elif sec["kind"] == "verification":
            t = d.add_table(rows=1, cols=4)
            t.style = "Table Grid"
            for i, h in enumerate(("ID", "검증 기준", "방법", "TC")):
                t.cell(0, i).text = h
            for r in sec["rows"]:
                c = t.add_row().cells
                c[0].text, c[1].text, c[2].text, c[3].text = r["id"], r["criteria"], r["method"], ", ".join(r["tc"])
        elif sec["kind"] == "trace":
            t = d.add_table(rows=1, cols=3)
            t.style = "Table Grid"
            for i, h in enumerate(("ID", "상위 (Upstream)", "하위 (Downstream)")):
                t.cell(0, i).text = h
            for r in sec["rows"]:
                c = t.add_row().cells
                c[0].text, c[1].text, c[2].text = r["id"], "\n".join(r["up"]), "\n".join(r["down"])
    out = Path(out)
    d.save(out)
    return out
