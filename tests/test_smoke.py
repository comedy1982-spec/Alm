"""End-to-end smoke test:  python -m pytest tests  (or python tests/test_smoke.py)"""
from pathlib import Path
import tempfile

from ecualm import ops
from ecualm.audit import run_audit
from ecualm.config import load_config, write_default_config
from ecualm.db import Repo
from ecualm.importer import import_document

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def test_revision_flow():
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        write_default_config(root)
        repo, cfg = Repo(root), load_config(root)

        a = import_document(repo, cfg, SAMPLES / "SRS_BLDC_revA.docx", doc="SRS", rev="A")
        assert a["by_type"]["SYS2"] == 7 and a["by_type"]["HSI"] == 5 and a["by_type"]["FIG"] == 2

        timeout_req = next(i for i in repo.items(type_="SWE1") if i["attrs"].get("native_id") == "REQ-COM-002")
        tc = ops.create_item(repo, cfg, "SWE6_TC", "300ms 경과 후 Timeout 판정 확인", domain="COM", doc="SRS")
        ops.link(repo, cfg, tc["id"], "verifies", timeout_req["id"])
        ops.edit_item(repo, cfg, timeout_req["id"], status="approved")

        b = import_document(repo, cfg, SAMPLES / "SRS_BLDC_revB.docx", doc="SRS", rev="B")
        assert timeout_req["id"] in b["modified"]            # same ID kept, change detected
        assert len(b["new"]) == 1 and len(b["deleted"]) == 1
        assert repo.get(tc["id"])["status"] != "deleted"     # manual items survive re-import
        it = repo.get(timeout_req["id"])
        assert it["status"] == "draft" and "500ms" in it["text"]

        res = run_audit(repo, cfg)
        rules = {(f["rule"], f["item"]) for f in res["findings"]}
        assert ("suspect", tc["id"]) in rules
        assert ("consistency", tc["id"]) in rules


def test_pdf_tables_and_figures():
    from ecualm.parse_pdf import parse_pdf
    blocks = parse_pdf(SAMPLES / "SRS_BLDC_revA.pdf")
    tables = [b for b in blocks if b.kind == "table"]
    figs = [b for b in blocks if b.kind == "figure"]
    assert len(tables) == 2 and len(figs) == 2
    pin = tables[1]
    assert len(pin.grid) == 6                               # table split across pages stitched
    assert pin.grid[1][3]["rs"] == 3                        # vertical merge preserved
    vec = [b for b in parse_pdf(SAMPLES / "vector_figure_example.pdf") if b.kind == "figure"]
    assert vec and vec[0].style == "vector" and vec[0].image




def test_stage_flow():
    from ecualm.stages import derive, stage_doc
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        write_default_config(root)
        repo, cfg = Repo(root), load_config(root)
        import_document(repo, cfg, SAMPLES / "SRS_BLDC_revA.docx", doc="CRS", rev="1.0", mode="crs")
        assert derive(repo, cfg, "SYS1", "SYS2")["created"]
        sd = stage_doc(repo, cfg, "SYS2")
        kinds = [s["kind"] for s in sd["sections"]]
        assert kinds == ["text", "text", "text", "group", "verification", "trace"]
        r = derive(repo, cfg, "SYS2", "SWE1")
        assert r["created"] and r["skipped"]              # HW-only requirements stay at SYS.2
        assert derive(repo, cfg, "SYS2", "SWE1")["created"] == []   # idempotent
        derive(repo, cfg, "SWE1", "SWE6"); derive(repo, cfg, "SWE1", "SWE2"); derive(repo, cfg, "SYS2", "SYS3")
        swe1 = repo.items(type_="SWE1")[0]
        assert "spec" in swe1["attrs"] and swe1["attrs"]["derived_from"].startswith("SYS2-")
        # a CRS revision must not delete derived items, and must mark derived links suspect
        b = import_document(repo, cfg, SAMPLES / "SRS_BLDC_revB.docx", doc="CRS", rev="1.1", mode="crs")
        assert all(not d.startswith(("SYS2", "SWE1", "SWE6")) for d in b["deleted"])
        res = run_audit(repo, cfg)
        assert any(f["rule"] == "suspect" and f["item"].startswith("SYS2") for f in res["findings"])
        assert not any(f["rule"] == "duplicate" for f in res["findings"])


def test_hwe_architecture_flow():
    """HW branch: SYS.2 → HWE.1 → HWE.2 blocks (+ canvas interfaces) → HWE.3/HWE.4, results and audit."""
    from ecualm import arch
    from ecualm.stages import derive
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        write_default_config(root)
        repo, cfg = Repo(root), load_config(root)
        import_document(repo, cfg, SAMPLES / "SRS_BLDC_revA.docx", doc="CRS", rev="1.0", mode="crs")
        derive(repo, cfg, "SYS1", "SYS2"); derive(repo, cfg, "SYS2", "SYS3")
        hw = derive(repo, cfg, "SYS2", "HWE1")
        assert hw["created"] and hw["skipped"]                      # SW-only requirements stay out of HWE.1
        assert all(not repo.get(i)["text"].lstrip().startswith("SW") for i in hw["created"])
        derive(repo, cfg, "SYS3", "HWE2"); derive(repo, cfg, "HWE1", "HWE2")
        blocks = repo.items(type_="HWE2")
        assert blocks and all("기구" not in b["title"] for b in blocks)
        assert derive(repo, cfg, "HWE2", "HWE3")["created"] and derive(repo, cfg, "HWE1", "HWE4")["created"]

        # moving a block is layout, not content: no new version, no suspect
        b0 = blocks[0]
        arch.set_layout(repo, {b0["id"]: {"x": 123, "y": 45}})
        assert repo.get(b0["id"])["version"] == b0["version"]
        assert repo.get(b0["id"])["attrs"]["canvas_layout"] == {"x": 123, "y": 45}

        # an interface line is an item that connects both blocks; a block change makes it suspect
        a, b = blocks[0], blocks[1]
        f = arch.create_interface(repo, cfg, "HWE2", a["id"], b["id"], "PWM x6", "pwm", "CH0, CH1")
        assert f["attrs"]["signals"] == ["CH0", "CH1"]
        assert {(l["role"], l["dst"]) for l in repo.links(f["id"]) if l["src"] == f["id"]} == {("connects", a["id"]), ("connects", b["id"])}

        # test result: PASS recorded with versions; editing the verified block makes the PASS stale
        tc_id = next(l["src"] for l in repo.links(a["id"]) if l["role"] == "verifies" and l["dst"] == a["id"])
        r = arch.ingest_results(repo, cfg, {"results": [{"tc": tc_id, "status": "pass", "run_id": "R1"}, {"tc": "NOPE", "status": "FAIL"}]})
        assert r["updated"] == [tc_id] and r["unknown"] == ["NOPE"]
        assert repo.get(tc_id)["version"] == 1 and repo.get(tc_id)["attrs"]["result"]["status"] == "PASS"
        ops.edit_item(repo, cfg, a["id"], attrs={"hw": {"part_number": "X1", "interfaces": "", "safety_mechanism": ""}})
        assert repo.get(a["id"])["title"] == a["title"]                 # a block keeps its name
        rules = {(x["rule"], x["item"]) for x in run_audit(repo, cfg)["findings"]}
        assert ("result_stale", tc_id) in rules
        arch.ingest_results(repo, cfg, [{"tc": tc_id, "status": "FAIL", "note": "over limit"}])
        assert ("test_fail", tc_id) in {(x["rule"], x["item"]) for x in run_audit(repo, cfg)["findings"]}

        # sample HW architecture merges into the derived blocks by name, and the canvas round-trips as JSON
        n_blocks = len(repo.items(type_="HWE2"))
        res = arch.load_sample(repo, cfg)
        assert res["interfaces"] == 10 and res["results"] == 3
        assert len(repo.items(type_="HWE2")) == n_blocks + len(res["created"])
        exp = arch.export_arch(repo, cfg, "HWE2")
        again = arch.import_arch(repo, cfg, exp)
        assert again["created"] == [] and again["interfaces"] == 0       # idempotent
        art = arch.artifact_json(repo, cfg, next(b["id"] for b in repo.items(type_="HWE2") if b["title"] == "과전류 하드웨어 차단 경로 (Shut-off Path)"))
        assert art["safety_classification"] == "B" and art["traceability"]["interfaces"]
        assert any(v["status"] == "PASS" for v in art["traceability"]["verified_by_downstream"])


def test_config_merge_adds_new_stages():
    """An alm.json written by an older version still gets the HWE stages and rules."""
    import json
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        old = {"project": "OLD", "stages": [{"id": "SYS2", "aspice": "SYS.2", "name": "SRS", "short": "SRS", "types": ["SYS2"],
                                             "next": ["SWE1"], "side": "spec"}],
               "link_rules": [{"src": "SWE1", "role": "refines", "dst": ["SYS2"]}]}
        (root / "alm.json").write_text(json.dumps(old), encoding="utf-8")
        cfg = load_config(root)
        ids = [s["id"] for s in cfg["stages"]]
        assert "HWE2" in ids and "HWE1" in next(s for s in cfg["stages"] if s["id"] == "SYS2")["next"]
        assert any(r["src"] == "HWE2" for r in cfg["link_rules"])


def test_assistant_local_and_llm(monkeypatch=None):
    """Chat panel back-end: offline commands work without a provider; with a provider the context carries
    the selected item and invalid action proposals are dropped."""
    from ecualm import assistant
    from ecualm.stages import derive
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        write_default_config(root)
        repo, cfg = Repo(root), load_config(root)
        import_document(repo, cfg, SAMPLES / "SRS_BLDC_revA.docx", doc="CRS", rev="1.0", mode="crs")
        derive(repo, cfg, "SYS1", "SYS2")
        assert assistant.status(cfg)["mode"] == "local"
        r = assistant.chat(repo, cfg, [{"role": "user", "content": "SYS.2 → SWE.1 전개"}])
        assert r["mode"] == "local" and r["actions"][0] == {"type": "derive", "from": "SYS2", "to": "SWE1", "label": "SYS2 → SWE1 전개"}
        sys2 = repo.items(type_="SYS2")[0]
        r = assistant.chat(repo, cfg, [{"role": "user", "content": f"{sys2['id']} 설명"}])
        assert sys2["id"] in r["reply"] and r["actions"][0]["type"] == "open"
        assert "감사 결과" in assistant.chat(repo, cfg, [{"role": "user", "content": "감사"}])["reply"]

        # LLM path with a stubbed transport
        cfg["llm"].update(provider="openai_compatible", base_url="http://stub", model="stub")
        seen = {}
        def fake(cfg_, system, messages, **kw):
            seen["ctx"] = messages[-1]["content"]
            return ('{"reply":"ok","actions":[{"type":"edit","id":"%s","text":"새 문장"},'
                    '{"type":"link","src":"NOPE","role":"refines","dst":"X"},{"type":"derive","from":"SYS2","to":"BAD"}]}' % sys2["id"])
        orig = assistant.chat_call
        assistant.chat_call = fake
        try:
            r = assistant.chat(repo, cfg, [{"role": "user", "content": "선택한 요구사항을 고쳐줘"}], {"selected": sys2["id"], "stage": "SYS2"})
        finally:
            assistant.chat_call = orig
        assert r["mode"] == "llm" and [a["type"] for a in r["actions"]] == ["edit"]
        assert f"[아이템 {sys2['id']}]" in seen["ctx"] and "SYS.2" in seen["ctx"]
        # a failing provider falls back to local answers instead of erroring
        def boom(*a, **k):
            raise RuntimeError("no network")
        assistant.chat_call = boom
        try:
            r = assistant.chat(repo, cfg, [{"role": "user", "content": "현황"}])
        finally:
            assistant.chat_call = orig
        assert r["mode"] == "local" and "AI 호출 실패" in r["reply"]


def _http(port, path, body=None, opener=None, csrf=True):
    import json, urllib.request, urllib.error
    h = {"content-type": "application/json"}
    if csrf:
        h["X-ECUALM"] = "1"
    req = urllib.request.Request(f"http://127.0.0.1:{port}{path}", None if body is None else json.dumps(body).encode(), h)
    try:
        r = opener.open(req)
        return r.status, json.loads(r.read() or b"{}")
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read() or b"{}")


def test_login_roles_and_plan():
    import http.cookiejar, threading, urllib.request
    from ecualm.server import _Server, make_handler
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        write_default_config(root)
        srv = _Server(("127.0.0.1", 0), make_handler(root)); port = srv.server_address[1]
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        mk = lambda: urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
        adm, eng, anon = mk(), mk(), mk()
        try:
            assert _http(port, "/api/state", opener=anon)[0] == 401
            assert _http(port, "/api/auth/state", opener=anon)[1]["setup"] is True
            assert _http(port, "/api/auth/setup", {"username": "pm", "password": "Passw0rd!"}, adm)[0] == 200
            assert _http(port, "/api/auth/setup", {"username": "x", "password": "Passw0rd!"}, anon)[0] == 409   # only once
            assert _http(port, "/api/state", opener=adm)[1]["me"]["role"] == "admin"
            assert _http(port, "/api/plan/template", {"start": "2026-11-02"}, adm, csrf=False)[0] == 403       # CSRF header required
            assert _http(port, "/api/plan/template", {"start": "2026-11-02"}, adm)[1]["created"] > 30
            plan = _http(port, "/api/plan", opener=adm)[1]
            t = next(x for x in plan["tasks"] if x["wbs"] == "4.1")
            assert not t["flags"], t["flags"]                       # template respects finish-to-start
            assert plan["metrics"]["bac"] > 0 and plan["load"]
            # engineer: temp password must be changed, then only own task progress/actuals
            assert _http(port, "/api/users", {"username": "kim", "password": "Temp12345", "role": "engineer"}, adm)[0] == 200
            assert _http(port, "/api/auth/login", {"username": "kim", "password": "nope"}, eng)[0] == 401
            assert _http(port, "/api/auth/login", {"username": "kim", "password": "Temp12345"}, eng)[1]["user"]["must_change"]
            assert _http(port, "/api/plan/task", {"id": t["id"], "progress": 10}, eng)[0] == 403
            assert _http(port, "/api/auth/password", {"old": "Temp12345", "new": "Kim-12345"}, eng)[0] == 200
            assert _http(port, "/api/plan/task", {"id": t["id"], "progress": 10}, eng)[0] == 403      # not the owner yet
            assert _http(port, "/api/plan/task", {"id": t["id"], "owner": "kim"}, adm)[0] == 200
            r = _http(port, "/api/plan/task", {"id": t["id"], "progress": 40, "effort_actual": 12}, eng)
            assert r[0] == 200 and r[1]["status"] == "in_progress"
            assert _http(port, "/api/plan/task", {"id": t["id"], "start": "2027-01-04"}, eng)[0] == 403   # dates are the PM's
            assert _http(port, "/api/baseline", {"name": "b"}, eng)[0] == 403
            # baseline then slip -> flags; cycle rejected
            assert _http(port, "/api/plan/baseline", {"name": "BL1"}, adm)[0] == 200
            nxt = next(x for x in plan["tasks"] if x["wbs"] == "4.2")
            import datetime as _dt
            later = (_dt.date.fromisoformat(t["finish"]) + _dt.timedelta(days=14)).isoformat()
            _http(port, "/api/plan/task", {"id": t["id"], "finish": later}, adm)
            t2 = next(x for x in _http(port, "/api/plan", opener=adm)[1]["tasks"] if x["id"] == t["id"])
            assert "slip" in t2["flags"] and t2["slip_days"] == 14
            assert "dep_conflict" in next(x for x in _http(port, "/api/plan", opener=adm)[1]["tasks"] if x["id"] == nxt["id"])["flags"]
            assert _http(port, "/api/plan/task", {"id": t["id"], "deps": [nxt["id"]]}, adm)[0] == 400
            x = adm.open(__import__("urllib.request").request.Request(f"http://127.0.0.1:{port}/api/plan/xlsx"))
            assert x.status == 200 and x.read(2) == b"PK"
            assert _http(port, "/api/auth/logout", {}, eng)[0] == 200 and _http(port, "/api/state", opener=eng)[0] == 401
        finally:
            srv.shutdown()
        from ecualm import schedule
        out = schedule.export_xlsx(Repo(root), load_config(root), root / "plan.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(out)
        assert wb.sheetnames == ["일정·공수 계획", "공수 부하", "계획 변경 이력"]
        assert wb["일정·공수 계획"]["A6"].value == "1"


def test_work_product_tables():
    from ecualm import workproducts as W
    from ecualm.stages import derive
    with tempfile.TemporaryDirectory() as td:
        root = Path(td); write_default_config(root)
        cfg, repo = load_config(root), Repo(root)
        import_document(repo, cfg, SAMPLES / "CRS_eLSD_ECU_rev1.0.docx", doc="CRS", rev="1.0", mode="crs")
        derive(repo, cfg, "SYS1", "SYS2"); derive(repo, cfg, "SYS2", "SWE1")
        t = W.table(repo, cfg, "SYS2")
        keys = {c["key"] for c in t["columns"] if c["required"]}
        assert {"id", "text", "asil", "verification_criteria", "upstream", "tests", "status"} <= keys
        assert len(t["rows"]) > 150 and t["missing_by_column"]["tests"] == len(t["rows"])   # no SYS.5 yet
        carried = [r for r in t["rows"] if r["values"]["asil"]]
        assert carried, "ASIL from the CRS table rows should reach SYS.2"
        # custom column: add, fill (no Suspect downstream), required -> counted as missing
        col = W.add_column(repo, cfg, "SYS2", "릴리스", "enum", ["A", "B"], required=True, user="t")
        row = next(r for r in t["rows"] if r["values"]["downstream"])
        W.set_cell(repo, cfg, row["id"], col["key"], "B", "t")
        assert not any(l["suspect"] for l in repo.links() if l["dst"] == row["id"])
        try:
            W.set_cell(repo, cfg, row["id"], col["key"], "Z", "t"); assert False
        except ValueError:
            pass
        t2 = W.table(repo, cfg, "SYS2")
        assert t2["missing_by_column"][col["key"]] == len(t2["rows"]) - 1
        # requirement content change marks downstream Suspect; nested attr path works for SWE.1 spec
        W.set_cell(repo, cfg, row["id"], "asil", "C", "t")
        assert any(l["suspect"] for l in repo.links() if l["dst"] == row["id"])
        swe = W.table(repo, cfg, "SWE1")["rows"][0]
        W.set_cell(repo, cfg, swe["id"], "timing", "10 ms", "t")
        assert repo.get(swe["id"])["attrs"]["spec"]["timing"] == "10 ms"
        assert repo.get(swe["id"])["attrs"]["spec"].get("output") is not None    # sibling keys kept
        # audit: approved item with missing required columns -> warn
        ops.edit_item(repo, cfg, row["id"], status="approved", user="t")
        f = [x for x in run_audit(repo, cfg)["findings"] if x["rule"] == "wp_required" and x["item"] == row["id"]]
        assert f and f[0]["sev"] == "warn"
        out = W.export_xlsx(repo, cfg, "SYS2", root / "t.xlsx")
        from openpyxl import load_workbook
        assert load_workbook(out)["SYS2"]["A4"].value == "ID *"
        W.delete_column(repo, "SYS2", col["key"])
        assert col["key"] not in {c["key"] for c in W.table(repo, cfg, "SYS2")["columns"]}


def test_test_sessions_and_evidence():
    from ecualm import testruns as T
    from ecualm.stages import derive
    with tempfile.TemporaryDirectory() as td:
        root = Path(td); write_default_config(root)
        cfg, repo = load_config(root), Repo(root); T.init(repo)
        import_document(repo, cfg, SAMPLES / "CRS_eLSD_ECU_rev1.0.docx", doc="CRS", rev="1.0", mode="crs")
        derive(repo, cfg, "SYS1", "SYS2"); derive(repo, cfg, "SYS2", "SWE1"); derive(repo, cfg, "SWE1", "SWE6")
        eng = {"username": "kim", "name": "김", "role": "engineer"}
        pm = {"username": "pm", "name": "PM", "role": "manager"}
        s = T.create_session(repo, cfg, {"name": "r1", "stage": "SWE6", "q": "MOT", "build": "0.1"}, eng)
        tcs = [c["tc"] for c in s["cases"]]
        assert s["counts"]["total"] == len(tcs) > 5 and s["status"] == "planned"
        T.record(repo, cfg, s["id"], {"tc": tcs[0], "status": "PASS"}, eng)
        assert repo.get(tcs[0])["attrs"]["result"]["status"] == "PASS"                   # pushed to the TC
        assert T.session(repo, cfg, s["id"])["status"] == "in_progress"
        try:
            T.record(repo, cfg, s["id"], {"tc": tcs[1], "status": "FAIL"}, eng); assert False
        except ValueError:
            pass
        T.record(repo, cfg, s["id"], {"tc": tcs[1], "status": "FAIL", "actual": "150ms", "defect": "J-1"}, eng)
        T.record(repo, cfg, s["id"], {"tc": tcs[2], "status": "N_A"}, eng)
        e = T.add_evidence(repo, s["id"], tcs[1], "wave.png", b"\x89PNG data", "scope", eng)
        meta, data = T.get_evidence(repo, e["id"])
        assert data == b"\x89PNG data" and meta["sha256"] == __import__("hashlib").sha256(data).hexdigest()
        assert any(f["rule"] == "test_fail" and f["item"] == tcs[1] for f in run_audit(repo, cfg)["findings"])
        d = T.dashboard(repo, cfg)
        assert d["kpi"]["fail"] == 1 and d["open_failures"][0]["defect"] == "J-1"
        cov = next(c for c in d["coverage"] if c["stage"] == "SWE6")
        assert cov["failed"] >= 1 and cov["linked"] > 0
        # retest session picks the failed TC only
        r2 = T.candidates(repo, cfg, "SWE6", only="failed")
        assert [c["id"] for c in r2] == [tcs[1]]
        done = T.update_session(repo, cfg, s["id"], {"action": "complete"}, eng)
        assert done["verdict"] == "FAIL" and done["status"] == "completed"
        for bad in ({"tc": tcs[3], "status": "PASS"},):
            try:
                T.record(repo, cfg, s["id"], bad, eng); assert False
            except ValueError:
                pass
        try:
            T.delete_evidence(repo, e["id"], eng); assert False                  # completed -> locked
        except ValueError:
            pass
        try:
            T.update_session(repo, cfg, s["id"], {"action": "reopen"}, eng); assert False
        except PermissionError:
            pass
        assert T.update_session(repo, cfg, s["id"], {"action": "reopen"}, pm)["status"] == "in_progress"
        out = T.export_xlsx(repo, cfg, s["id"], root / "r.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(out)
        assert wb.sheetnames == ["요약", "TC 결과", "증빙"] and wb["증빙"]["D2"].value == e["sha256"]


def test_sup_processes():
    from ecualm import sup as U, testruns as T
    from ecualm.stages import derive
    with tempfile.TemporaryDirectory() as td:
        root = Path(td); write_default_config(root)
        cfg, repo = load_config(root), Repo(root); T.init(repo); U.init(repo)
        import_document(repo, cfg, SAMPLES / "CRS_eLSD_ECU_rev1.0.docx", doc="CRS", rev="1.0", mode="crs")
        derive(repo, cfg, "SYS1", "SYS2"); derive(repo, cfg, "SYS2", "SWE1"); derive(repo, cfg, "SWE1", "SWE6")
        eng = {"username": "kim", "name": "김", "role": "engineer"}
        eng2 = {"username": "lee", "name": "이", "role": "engineer"}
        pm = {"username": "pm", "name": "PM", "role": "manager"}
        s = T.create_session(repo, cfg, {"name": "r1", "stage": "SWE6", "q": "COM"}, eng)
        tc = s["cases"][0]["tc"]
        T.record(repo, cfg, s["id"], {"tc": tc, "status": "FAIL", "actual": "x"}, eng)
        # SUP.9: problem from failed test, guarded workflow
        p = U.problem_from_test(repo, cfg, s["id"], tc, eng)
        assert {x["role"] for x in p["items"]} == {"found_by", "verify", "affects"}
        assert T.session(repo, cfg, s["id"])["cases"][0]["defect"] == p["id"]
        U.transition(repo, cfg, p["id"], "analyzing", "", eng)
        try:
            U.transition(repo, cfg, p["id"], "resolving", "", eng); assert False
        except ValueError:
            pass
        U.update(repo, cfg, p["id"], {"fields": {"root_cause_cat": "코드", "root_cause": "rc", "corrective": "fix", "verification": "TS"}}, eng)
        U.transition(repo, cfg, p["id"], "resolving", "", eng)
        U.transition(repo, cfg, p["id"], "verifying", "", eng)
        try:
            U.transition(repo, cfg, p["id"], "closed", "", eng); assert False      # verify TC still FAIL
        except ValueError:
            pass
        # SUP.10: CR from the problem, impact analysis, CCB by manager only, close after retest PASS
        cr = U.cr_from_problem(repo, cfg, p["id"], eng)
        assert p["id"] in cr["refs"] and cr["id"] in U.get(repo, p["id"])["refs"]
        U.transition(repo, cfg, cr["id"], "analyzing", "", eng)
        try:
            U.transition(repo, cfg, cr["id"], "review", "", eng); assert False
        except ValueError:
            pass
        cr = U.save_impact(repo, cfg, cr["id"], eng)
        imp = cr["fields"]["impact"]
        assert imp["rows"][0]["action"] == "modify" and tc in imp["retest"]
        U.transition(repo, cfg, cr["id"], "review", "", eng)
        U.update(repo, cfg, cr["id"], {"fields": {"decision": "ok"}}, pm)
        try:
            U.transition(repo, cfg, cr["id"], "approved", "", eng); assert False
        except ValueError:
            pass
        U.transition(repo, cfg, cr["id"], "approved", "", pm)
        U.transition(repo, cfg, cr["id"], "implementing", "", eng); U.transition(repo, cfg, cr["id"], "verifying", "", eng)
        U.update(repo, cfg, cr["id"], {"fields": {"verification": "retest"}}, eng)
        s2 = T.create_session(repo, cfg, {"name": "retest", "stage": "SWE6", "tcs": imp["retest"]}, eng)
        for t in imp["retest"]:
            T.record(repo, cfg, s2["id"], {"tc": t, "status": "PASS"}, eng)
        U.transition(repo, cfg, cr["id"], "closed", "요구사항 문구 변경 없음 (코드 수정)", eng)
        U.transition(repo, cfg, p["id"], "closed", "", eng)
        # SUP.1: checklist auto evaluation, findings, independence of verification
        qa = U.create(repo, cfg, "QA", {"title": "SWE1", "fields": {"check_type": "산출물 점검", "target": "SWE1"}}, pm)
        qa = U.evaluate_checklist(repo, cfg, qa["id"], pm)
        assert any(c["result"] == "NOK" for c in qa["fields"]["checklist"])
        qf = U.findings_from_checklist(repo, cfg, qa["id"], pm)["created"][0]
        U.update(repo, cfg, qf, {"fields": {"corrective": "c"}}, eng); U.transition(repo, cfg, qf, "resolved", "", eng)
        try:
            U.transition(repo, cfg, qf, "verified", "", eng); assert False
        except ValueError:
            pass
        U.transition(repo, cfg, qf, "verified", "", pm)
        try:
            U.transition(repo, cfg, qa["id"], "done", "", pm); assert False     # summary + manual item missing
        except ValueError:
            pass
        # SUP.8: release checks, baseline, backup/restore
        rel = U.create(repo, cfg, "REL", {"title": "B", "fields": {"sw_version": "0.4.1", "git_commit": "abc", "sessions": s2["id"]}}, pm)
        assert {c["key"] for c in rel["checks"]} >= {"audit_errors", "open_cr", "commit"}
        rel = U.transition(repo, cfg, rel["id"], "released", "데모", pm)
        assert repo.baseline(rel["fields"]["item_baseline"]) is not None
        assert U.baseline_diff(repo, rel["fields"]["item_baseline"], None)["rows"] == []
        z = U.backup(repo, root / "b.zip")
        dst = root / "restored"
        U.restore(dst, z)
        assert Repo(dst).get(tc) is not None
        U.summary(repo, cfg)
        assert any(f["rule"] == "cr_without_impact" for f in U.audit_findings(repo, cfg)) is False


def test_revision_change_analysis():
    import json as _json
    from ecualm import changes as C, llm, sup as U
    from ecualm.stages import derive
    sys_path = SAMPLES / "CRS_eLSD_ECU_rev1.1.docx"
    if not sys_path.exists():
        import subprocess, sys
        subprocess.run([sys.executable, str(SAMPLES / "make_elsd_rev.py")], check=True, capture_output=True)
    with tempfile.TemporaryDirectory() as td:
        root = Path(td); write_default_config(root)
        cfg, repo = load_config(root), Repo(root); U.init(repo)
        import_document(repo, cfg, SAMPLES / "CRS_eLSD_ECU_rev1.0.docx", doc="CRS", rev="1.0", mode="crs")
        derive(repo, cfg, "SYS1", "SYS2"); derive(repo, cfg, "SYS2", "SWE1"); derive(repo, cfg, "SWE1", "SWE6")
        prev = import_document(repo, cfg, sys_path, doc="CRS", rev="1.1", mode="crs", dry_run=True)
        rep = C.analyze(repo, cfg, prev)
        by = {r["native_id"]: r for r in rep["rows"] if r["native_id"]}
        assert by["REQ-CLU-001"]["numbers"][0]["direction"] == "stricter" and by["REQ-CLU-001"]["severity"] == "high"
        assert by["REQ-YAW-003"]["numbers"][0]["direction"] == "threshold"
        assert by["REQ-MODE-014"]["numbers"][0]["direction"] == "looser" and by["REQ-MODE-014"]["severity"] == "medium"
        assert "asil" in by["REQ-THERM-003"]["kinds"]                      # table ASIL column only
        assert "obligation" in by["REQ-LOCK-010"]["kinds"]                  # shall -> should (ID still in the doc)
        assert by["REQ-LOCK-012"]["kinds"] == ["removed"]
        assert "negation" in by["REQ-DIAG-003"]["kinds"] and "obligation" not in by["REQ-DIAG-003"]["kinds"]
        assert by["REQ-ENV-001"]["kinds"] == ["editorial"] and by["REQ-ENV-001"]["severity"] == "low"
        assert by["REQ-SEC-007"]["change"] == "new"
        assert rep["mode"] == "local" and rep["counts"]["retest"] > 0
        # LLM enrichment (stubbed): explanation merged, risk can only go up, failure keeps local report
        cfg["llm"].update(provider="openai_compatible", base_url="http://stub", model="stub")
        orig = llm.chat_call
        llm.chat_call = lambda *a, **k: _json.dumps({"summary": ["요약"], "questions": ["Q1"],
                                                     "items": {by["REQ-ENV-001"]["id"]: {"explain": "표현", "risk": "medium"}}})
        try:
            rep2 = C.analyze(repo, cfg, prev)
        finally:
            llm.chat_call = orig
        r2 = next(r for r in rep2["rows"] if r["native_id"] == "REQ-ENV-001")
        assert rep2["mode"] == "llm" and r2["ai"] == "표현" and r2["severity"] == "medium" and rep2["questions"] == ["Q1"]
        llm.chat_call = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("down"))
        try:
            rep3 = C.analyze(repo, cfg, prev)
        finally:
            llm.chat_call = orig
        assert rep3["mode"] == "local" and "AI 호출 실패" in rep3["ai_status"]
        cfg["llm"]["provider"] = "none"
        # real import -> save report -> change request with impact
        real = import_document(repo, cfg, sys_path, doc="CRS", rev="1.1", mode="crs")
        rep = C.analyze(repo, cfg, real)
        C.save(repo, rep, "t")
        assert C.reports(repo)[0]["import_id"] == real["import_id"]
        pick = [r["id"] for r in rep["rows"] if r["severity"] == "high" and r["change"] == "modified"]
        cr = C.to_cr(repo, cfg, real["import_id"], pick, {"username": "t", "name": "t", "role": "engineer"})
        assert cr["fields"]["reason"] == "고객 요구 변경" and cr["fields"]["impact"]["scope"] == pick
        out = C.export_xlsx(rep, root / "c.xlsx")
        assert out.exists()


def test_source_side_by_side():
    from ecualm import sources as SR
    with tempfile.TemporaryDirectory() as td:
        root = Path(td); write_default_config(root)
        cfg, repo = load_config(root), Repo(root); SR.init(repo)
        a = import_document(repo, cfg, SAMPLES / "CRS_eLSD_ECU_rev1.0.pdf", doc="E", rev="1.0", mode="crs")
        b = import_document(repo, cfg, SAMPLES / "CRS_eLSD_ECU_rev1.1.pdf", doc="E", rev="1.1", mode="crs")
        assert SR.source_file(repo, a["import_id"]).name == "CRS_eLSD_ECU_rev1.0.pdf"      # each revision kept
        m = SR.meta(repo, b["import_id"])
        assert m["kind"] == "pdf" and len(m["pages"]) > 40
        La, Lb = SR.locate(repo, a["import_id"]), SR.locate(repo, b["import_id"])
        reqs = [i for i in Lb["items"] if i["type"] == "CRS"]
        assert sum(1 for i in reqs if i["id"] in Lb["locs"]) >= len(reqs) - 1
        clu = next(i for i in Lb["items"] if i["native_id"] == "REQ-CLU-001")
        old = next(i for i in La["items"] if i["id"] == clu["id"])
        assert "100ms" in clu["text"] and "120ms" in old["text"]                       # text as of each revision
        assert clu["change"] == "modified" and clu["id"] in La["locs"] and clu["id"] in Lb["locs"]
        assert not any(i["native_id"] == "REQ-LOCK-012" for i in Lb["items"])          # withdrawn in 1.1
        assert SR.page_png(repo, b["import_id"], 0)[:4] == b"\x89PNG"
        h = SR.html(repo, b["import_id"])
        assert "data-items" in h["html"]


if __name__ == "__main__":
    test_assistant_local_and_llm()
    test_revision_flow()
    test_pdf_tables_and_figures()
    test_stage_flow()
    test_hwe_architecture_flow()
    test_config_merge_adds_new_stages()
    test_login_roles_and_plan()
    test_work_product_tables()
    test_test_sessions_and_evidence()
    test_sup_processes()
    test_revision_change_analysis()
    test_source_side_by_side()
    print("OK")
